"""In-memory readers and the shared Program import pipeline."""

import csv
import datetime as dt
import io
import math
import re
from collections.abc import Callable
from bisect import bisect_right
from dataclasses import dataclass, field, replace
from typing import TYPE_CHECKING, Any, Iterable, Protocol

from fastapi import BackgroundTasks
from openpyxl import Workbook, load_workbook

from agent.ProgramState import (
    MAX_PROGRAM_LOAD_KG,
    MAX_PROGRAM_LOAD_PCT_E1RM,
    MIN_PROGRAM_LOAD_KG,
    MIN_PROGRAM_LOAD_PCT_E1RM,
)
from agent.program_prescription import (
    DEFAULT_EXERCISE_REST_SECONDS,
    DEFAULT_TARGET_REPS_MAX,
    DEFAULT_TARGET_REPS_MIN,
    DEFAULT_TARGET_RIR,
    DEFAULT_TARGET_SETS,
    MAX_EXERCISES_PER_DAY,
    MAX_PROGRAM_DAYS,
    MAX_PROGRAM_WEEKS,
    MAX_PROGRAM_NOTES_LENGTH,
    MAX_REPS,
    MAX_TEMPO_LENGTH,
    MAX_TARGET_RIR,
    MIN_REPS,
    MIN_TARGET_RIR,
)
from service import coach_program_drafts
from service.assignments import authorized_player_ledger
from service.program_import_constants import (
    MAX_PROGRAM_IMPORT_COLUMNS,
    MAX_PROGRAM_IMPORT_DAY_NAME_LENGTH,
    MAX_PROGRAM_IMPORT_EXERCISE_NAME_LENGTH,
    MAX_PROGRAM_IMPORT_FILE_BYTES,
    MAX_PROGRAM_IMPORT_TEXT_CHARS,
    MAX_PROGRAM_IMPORT_CELLS,
    MAX_PROGRAM_IMPORT_REST_SECONDS,
    MAX_PROGRAM_IMPORT_ROWS,
    MAX_PROGRAM_IMPORT_SETS,
    MAX_PROGRAM_IMPORT_TABS,
    MAX_PROGRAM_IMPORT_TAB_PROBE_ROWS,
    MAX_PROGRAM_IMPORT_GRID_COLUMNS,
    MAX_PROGRAM_IMPORT_GRID_ROWS,
    MAX_PROGRAM_IMPORT_GRID_RECTANGLE,
    MAX_PROGRAM_IMPORT_HEADER_SCAN_ROWS,
    MIN_PROGRAM_IMPORT_LAYOUT_CONFIDENCE,
)
from service.text_normalization import normalize_digits
from service import analytics, program_import_ai

if TYPE_CHECKING:
    from svc.llm import InferenceScope

TEMPLATE_COLUMNS = (
    "week", "day", "day_name", "order", "exercise", "sets", "reps", "rir", "rpe",
    "load_kg", "load_pct_e1rm", "rest_sec", "tempo", "notes",
)
_REP_RANGE = re.compile(r"^(\d+)\s*-\s*(\d+)$")
_LOAD_PLACEHOLDERS = frozenset({"", "—", "-", "n/a", "none"})
_PER_SET_REP_LIST = re.compile(r"^(?:\d+\s*[x×]\s*)?(\d+(?:\s*[,/]\s*\d+)+)$", re.IGNORECASE)
_EFFORT_VALUE = re.compile(r"^(?:RIR\s*)?(\d+(?:\.\d+)?)$", re.IGNORECASE)
_EFFORT_RANGE = re.compile(r"^(?:RIR\s*)?(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)$", re.IGNORECASE)
_RPE_VALUE = re.compile(r"^(?:RPE\s*|@\s*)?(\d+(?:\.\d+)?)$", re.IGNORECASE)
_RPE_RANGE = re.compile(r"^(?:RPE\s*|@\s*)?(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)$", re.IGNORECASE)
_PRESCRIPTION_SOURCE_FIELDS = frozenset({"sets", "reps", "prescription", "reps_by_set"})
_SOURCE_HEADER_FIELDS = {
    "day": "day", "day number": "day", "training day": "day", "يوم": "day", "اليوم": "day",
    "session": "day_name", "session name": "day_name", "day name": "day_name",
    "اسم اليوم": "day_name", "اسم الحصة": "day_name", "الحصة": "day_name",
    "workout": "day_name", "workout name": "day_name",
    "exercise": "exercise", "movement": "exercise", "تمرين": "exercise", "التمرين": "exercise",
    "exercise as written": "exercise", "lift": "exercise",
    "slot": "slot", "week": "week", "week number": "week",
    "training week": "week", "set": "sets", "sets": "sets", "مجموعات": "sets", "جولات": "sets",
    "سيتات": "sets",
    "rpe": "rpe", "rir": "rir", "reps": "reps", "rep": "reps",
    "تكرارات": "reps", "عدات": "reps", "عدّات": "reps",
    "راحة": "rest_seconds", "rir/rpe": "effort", "rpe/rir": "effort",
    "وزن": "load_kg", "الوزن": "load_kg", "وزن/كجم": "load_kg", "الوزن/كجم": "load_kg",
    "وزن (كجم)": "load_kg", "الوزن (كجم)": "load_kg", "ملاحظات": "notes", "ملاحظة": "notes",
    "ملاحظات المدرب": "notes",
    "prescription": "prescription", "reps by set": "reps_by_set",
    "rep scheme": "reps_by_set",
}
_UNSUPPORTED_PRESCRIPTION_PATTERNS = (
    (r"\bAMRAP\b|\btop\s+set\b|\bback[ -]?off\b|\bsuperset\b", "prescription"),
    (r"\d+\s*[x×]\s*\d+(?:\s*[,/]\s*\d+)+", "reps"),
    (r"\d+(?:\.\d+)?\s*%", "load_pct_e1rm"),
    (r"weekly\s+(?:progression|increase|increase reps|add reps)|[+-]\s*\d+(?:\.\d+)?\s*(?:kg|lb)\s*/\s*(?:week|wk)", "notes"),
)


class ProgramImportError(Exception):
    def __init__(
        self,
        code: str,
        message: str,
        status_code: int = 400,
        message_params: dict[str, Any] | None = None,
    ):
        super().__init__(message)
        self.code = code
        self.status_code = status_code
        self.message_params = message_params or {}


@dataclass(frozen=True)
class Approximation:
    code: str
    column: str
    original: str = ""


ImportValue = int | float | str | bool | dt.date | dt.datetime | dt.time | None


@dataclass(frozen=True)
class ProgramImportRow:
    """Reader-neutral row. Values stay typed where possible; source text is retained for review."""

    source_row: int
    day: ImportValue
    day_name: ImportValue
    order: ImportValue
    exercise: ImportValue
    sets: ImportValue
    reps_min: ImportValue
    reps_max: ImportValue
    rir: ImportValue
    rpe: ImportValue
    rest_seconds: ImportValue
    tempo: ImportValue
    notes: ImportValue
    original_text: str
    load_kg: ImportValue = None
    load_pct_e1rm: ImportValue = None
    load: ImportValue = None
    reps_original: ImportValue = None
    approximation_markers: tuple[Approximation, ...] = ()
    confirmed_exercise_id: str | None = None
    week: str | None = None
    week_order: ImportValue = None
    suggestion_terms: tuple[str, ...] = ()
    input_exercise_name: str | None = None


@dataclass(frozen=True)
class ReaderOutput:
    detected_tabs: list[str]
    selected_tab: str | None
    requires_tab_choice: bool
    rows: list[ProgramImportRow] = field(default_factory=list)
    detected_weeks: list[str] = field(default_factory=list)
    selected_week: str | None = None
    weeks_not_imported: list[str] = field(default_factory=list)
    confirm_layout: bool = False
    layout_days: list[dict[str, Any]] = field(default_factory=list)
    layout_confidence: float | None = None
    import_mode: str = "template"


@dataclass(frozen=True)
class ReaderDetection:
    detected_tabs: list[str]
    selected_tab: str | None
    requires_tab_choice: bool
    is_template: bool


@dataclass(frozen=True)
class FreeFormCellGrid:
    grid: dict[str, Any] | None
    tabs: list[str]
    selected_tab: str | None


@dataclass(frozen=True)
class FreeFormGridRows:
    columns_by_row: dict[int, dict[int, tuple[str, int | None]]]
    header_rows: frozenset[int]
    exercise_data_rows: tuple[int, ...]
    exercise_data_columns_by_schema: dict[tuple[tuple[int, str, int | None], ...], frozenset[int]]
    section_title_rows: frozenset[int]


@dataclass(frozen=True)
class FreeFormGridContext:
    cells_by_row: dict[int, dict[int, str]]
    columns: dict[int, tuple[str, int | None]]
    header_row: int | None
    week_labels: list[str]
    title_week_labels: list[str]
    explicit_week_labels: list[str]
    implicit_week_columns: dict[int, int]
    implicit_week_count: int
    day_names: dict[int, str]
    day_numbers: dict[str, int]
    row_facts: FreeFormGridRows
    preamble_text: str = ""
    layout_ambiguous: bool = False


@dataclass(frozen=True)
class FreeFormSourceRow:
    fields: dict[str, str]
    text: str


@dataclass(frozen=True)
class TextSourceMarkerRow:
    original_text: str
    notes: str = ""


@dataclass(frozen=True)
class FreeFormRowValues:
    day: Any
    day_name: str | None
    rir: Any
    rpe: Any
    reps_original: Any
    approximation_markers: tuple[Approximation, ...]


@dataclass(frozen=True)
class ProgramImportRequest:
    coach_account_id: str
    assignment_id: str
    payload: bytes
    filename: str
    text: str | None = None
    selected_tab: str | None = None
    reader: "ProgramImportReader | None" = None
    client: analytics.ClientContext = analytics.UNKNOWN_CLIENT
    background_tasks: BackgroundTasks | None = None


class ProgramImportReader(Protocol):
    def read(self, payload: bytes, filename: str, selected_tab: str | None) -> ReaderOutput:
        """Return canonical rows; all validation and resolution happens after this boundary."""


class TemplateProgramImportReader:
    """Translate the MAYOS XLSX/CSV columns to the reader-neutral row shape."""

    def read(self, payload: bytes, filename: str, selected_tab: str | None) -> ReaderOutput:
        extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
        if extension == "csv":
            if selected_tab:
                raise ProgramImportError("program_import.tab_invalid.v1", "CSV files have one sheet.")
            return _read_csv(payload)
        if extension == "xlsx":
            return _read_xlsx(payload, selected_tab)
        raise ProgramImportError("program_import.file_type.v1", "Upload an XLSX, CSV, or TXT file.")


class FreeFormProgramImportReader:
    """Translate one bounded, coach-selected sheet through the hosted coach model."""

    def __init__(self, inference_scope: "InferenceScope"):
        self.inference_scope = inference_scope

    def read(self, prepared: FreeFormCellGrid) -> ReaderOutput:
        if prepared.grid is None:
            return ReaderOutput(prepared.tabs, prepared.selected_tab, True, import_mode="freeform")
        reply = _read_import_translation(lambda: program_import_ai.translate_sheet(
            prepared.grid,
            self.inference_scope,
            implicit_week_count=_grid_implicit_week_count(prepared.grid),
        ))
        return _freeform_reader_output(reply, prepared.tabs, prepared.selected_tab, prepared.grid)


class PlainTextProgramImportReader:
    """Translate coach-written program text through the hosted coach model."""

    def __init__(self, inference_scope: "InferenceScope"):
        self.inference_scope = inference_scope

    def read(self, text: str) -> ReaderOutput:
        reply = _read_import_translation(
            lambda: program_import_ai.translate_text(text, self.inference_scope)
        )
        return _text_reader_output(reply, text)


def _read_import_translation(translate: Callable[[], Any]) -> Any:
    try:
        return translate()
    except program_import_ai.ProgramImportAIError as error:
        raise ProgramImportError(
            "program_import.interpretation_failed.v1", str(error), 422
        ) from None
    except (program_import_ai.ProgramImportProviderUnavailable, TimeoutError):
        raise ProgramImportError(
            program_import_ai.PROVIDER_BUSY_MESSAGE_CODE,
            "The import assistant is busy with another request. Try again in a minute.",
            503,
        ) from None


def import_program_sheet(db: Any, request: ProgramImportRequest) -> dict[str, Any] | None:
    authorized = authorized_player_ledger(db, request.coach_account_id, request.assignment_id)
    if authorized is None:
        return None
    if len(request.payload) > MAX_PROGRAM_IMPORT_FILE_BYTES:
        raise ProgramImportError("program_import.file_too_large.v1", "Trim the file to 1 MB or less.", 413)
    text_mode = request.text is not None or request.filename.lower().endswith(".txt")
    text = None
    if text_mode:
        gate = program_import_ai.resolve_text_enable_gate()
        if not gate.enabled:
            raise ProgramImportError(
                "program_import.text_disabled.v1",
                "Plain-text import is unavailable.",
                503,
            )
        text = _normalised_import_text(request)
    ledger, _ = authorized
    with ledger:
        draft_exists = ledger.get_program_draft(request.assignment_id) is not None
    if text_mode:
        source, rows = _run_text_import(db, request, text)
    else:
        if request.reader is not None:
            source = request.reader.read(request.payload, request.filename, request.selected_tab)
        else:
            detection = _detect_program_reader(
                request.payload, request.filename, request.selected_tab
            )
            if detection.requires_tab_choice:
                source = ReaderOutput(
                    detection.detected_tabs,
                    detection.selected_tab,
                    True,
                    import_mode="freeform" if not detection.is_template else "template",
                )
                return _build_import_result(source, draft_exists, [])
            source = (
                TemplateProgramImportReader().read(
                    request.payload,
                    request.filename,
                    detection.selected_tab
                    if request.filename.lower().endswith(".xlsx")
                    else None,
                )
                if detection.is_template
                else None
            )
        if source is not None and source.requires_tab_choice:
            return _build_import_result(source, draft_exists, [])
        if source is None:
            freeform_result = _read_freeform_program(db, request)
            if freeform_result is None:
                return None
            source, rows = freeform_result
        else:
            rows = validate_and_resolve_rows(db, request.coach_account_id, source.rows)
    if text_mode and db.get_active_assignment_for_coach(
        request.coach_account_id, request.assignment_id
    ) is None:
        return None
    return _build_import_result(source, draft_exists, rows)


def _normalised_import_text(request: ProgramImportRequest) -> str:
    if request.text is not None:
        text = request.text
    else:
        try:
            text = request.payload.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ProgramImportError(
                "program_import.text_encoding.v1", "Text files must use UTF-8 encoding.", 422
            ) from None
    text = text.lstrip("\ufeff").replace("\r\n", "\n").replace("\r", "\n")
    if len(text) > MAX_PROGRAM_IMPORT_TEXT_CHARS:
        raise ProgramImportError(
            "program_import.text_too_large.v1",
            f"Shorten the program to {MAX_PROGRAM_IMPORT_TEXT_CHARS:,} characters or fewer.",
            422,
        )
    if not text.strip():
        raise ProgramImportError(
            "program_import.text_empty.v1", "Enter or upload program text.", 422
        )
    return text


def _read_freeform_program(
    db: Any, request: ProgramImportRequest
) -> tuple[ReaderOutput, list[dict[str, Any]]] | None:
    gate = program_import_ai.resolve_enable_gate()
    if not gate.enabled:
        raise ProgramImportError(
            "program_import.freeform_disabled.v1",
            "Free-form import is unavailable. Use the MAYOS template.",
            503,
        )
    source, rows = _run_freeform_import(db, request)
    if db.get_active_assignment_for_coach(
        request.coach_account_id, request.assignment_id
    ) is None:
        return None
    return source, rows


def _run_freeform_import(
    db: Any, request: ProgramImportRequest
) -> tuple[ReaderOutput, list[dict[str, Any]]]:
    prepared = _freeform_cell_grid(
        request.payload, request.filename, request.selected_tab
    )
    if prepared.grid is None:
        return ReaderOutput(
            prepared.tabs, prepared.selected_tab, True, import_mode="freeform"
        ), []
    from svc.llm import InferenceScope, inference_turn

    scope = InferenceScope(
        account_id=request.coach_account_id,
        role="coach",
        purpose="coach_program_import",
        store=db,
        client=request.client,
    )
    with inference_turn(scope, background_tasks=request.background_tasks):
        source = FreeFormProgramImportReader(scope).read(prepared)
        rows = validate_and_resolve_rows(db, request.coach_account_id, source.rows)
        _add_model_suggestions(db, rows, scope)
    return source, rows


def _detect_program_reader(
    payload: bytes, filename: str, selected_tab: str | None
) -> ReaderDetection:
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if extension == "csv":
        if selected_tab:
            raise ProgramImportError("program_import.tab_invalid.v1", "CSV files have one sheet.")
        header = _csv_header(payload)
        return ReaderDetection(["CSV"], "CSV", False, _is_template_header(header))
    if extension != "xlsx":
        raise ProgramImportError("program_import.file_type.v1", "Upload an XLSX, CSV, or TXT file.")
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        return _detect_xlsx_reader(workbook, selected_tab)
    except ProgramImportError:
        raise
    except Exception as error:
        raise ProgramImportError(
            "program_import.file_invalid.v1", "The XLSX file could not be read."
        ) from error
    finally:
        if workbook is not None:
            workbook.close()


def _csv_header(payload: bytes) -> list[str]:
    try:
        table = csv.reader(io.StringIO(payload.decode("utf-8-sig")))
        header = _first_non_empty_header(table)
    except (UnicodeDecodeError, csv.Error) as error:
        raise ProgramImportError(
            "program_import.file_invalid.v1", "The CSV file could not be read."
        ) from error
    if not header and not payload.strip():
        raise ProgramImportError("program_import.sheet_empty.v1", "The CSV sheet is empty.")
    return header


def _detect_xlsx_reader(workbook: Any, selected_tab: str | None) -> ReaderDetection:
    if len(workbook.worksheets) > MAX_PROGRAM_IMPORT_TABS:
        raise ProgramImportError(
            "program_import.too_many_tabs.v1", "Trim the workbook to 20 sheets or fewer.", 413
        )
    tabs = [sheet.title for sheet in workbook.worksheets if _sheet_has_values(sheet)]
    if not tabs:
        raise ProgramImportError("program_import.sheet_empty.v1", "The workbook has no non-empty sheets.")
    choice_tabs = [tab for tab in tabs if tab.casefold() != "example"]
    if len(choice_tabs) > 1 and selected_tab is None:
        return ReaderDetection(tabs, None, True, False)
    selected = selected_tab or ("Program" if "Program" in tabs else (choice_tabs or tabs)[0])
    if selected not in tabs:
        raise ProgramImportError(
            "program_import.tab_invalid.v1", "Choose a non-empty sheet from this workbook."
        )
    sheet = workbook[selected]
    header = _first_non_empty_header(
        sheet.iter_rows(
            max_row=MAX_PROGRAM_IMPORT_HEADER_SCAN_ROWS,
            max_col=MAX_PROGRAM_IMPORT_GRID_COLUMNS,
            values_only=True,
        )
    )
    if not header:
        if (
            (sheet.max_row or 0) > MAX_PROGRAM_IMPORT_GRID_ROWS
            or (sheet.max_column or 0) > MAX_PROGRAM_IMPORT_GRID_COLUMNS
        ):
            return ReaderDetection(tabs, selected, False, False)
        raise ProgramImportError("program_import.sheet_empty.v1", "The selected sheet is empty.")
    return ReaderDetection(tabs, selected, False, _is_template_header(header))


def _first_non_empty_header(table: Iterable[Iterable[Any]]) -> list[str]:
    for row_number, values in enumerate(table, start=1):
        if row_number > MAX_PROGRAM_IMPORT_HEADER_SCAN_ROWS:
            break
        header = [_cell_text(value).casefold() for value in values]
        if any(header):
            return header
    return []


def _is_template_header(header: list[str]) -> bool:
    required = {"day", "exercise", "sets", "reps"}
    return required.issubset(set(header))


def _add_model_suggestions(
    db: Any, rows: list[dict[str, Any]], scope: "InferenceScope"
) -> None:
    names = list(dict.fromkeys(
        row["exercise_name"]
        for row in rows
        if row["exercise_id"] is None and row["exercise_name"]
    ))
    if not names:
        return
    search_terms = {
        row["exercise_name"]: row["suggestion_search_terms"]
        for row in rows
        if row["exercise_name"] in names and row["suggestion_search_terms"]
    }
    try:
        suggested = program_import_ai.suggest_exercises(db, names, scope, search_terms)
    except (program_import_ai.ProgramImportAIError, TimeoutError):
        for row in rows:
            if row["exercise_id"] is None:
                row["warnings"].append(_row_warning(
                    "program_import.suggestions_unavailable.v1",
                    row["source_row"],
                    "exercise",
                    "Automated suggestions are unavailable; search the library or create a Coach exercise.",
                ))
        return
    for row in rows:
        model_matches = suggested.get(row["exercise_name"], [])
        row["suggestions"] = _merge_suggestions(row["suggestions"], model_matches)


def _merge_suggestions(
    deterministic: list[dict[str, str]], model_suggestions: list[dict[str, str]]
) -> list[dict[str, str]]:
    merged: dict[str, dict[str, str]] = {}
    for candidate in (*deterministic, *model_suggestions):
        merged.setdefault(candidate["exercise_id"], candidate)
    return list(merged.values())[:3]


def create_imported_program_draft(
    db: Any,
    coach_account_id: str,
    assignment_id: str,
    program_name: str | None,
    rows: list[dict[str, Any]],
    *,
    replace: bool = False,
) -> dict[str, Any] | None:
    authorized = authorized_player_ledger(db, coach_account_id, assignment_id)
    if authorized is None:
        return None
    ledger, _ = authorized
    # authorized_player_ledger opens the ledger; close it here before the draft service reopens it.
    with ledger:
        pass
    source_rows = [_confirmed_row(row) for row in rows]
    reviewed_rows = validate_and_resolve_rows(db, coach_account_id, source_rows)
    weeks = _confirmed_rows_by_week(db, coach_account_id, reviewed_rows)
    days = [day for week in weeks for day in week["days"]]
    draft = {
        "program_name": (program_name or "").strip() or "Imported program",
        "split_type": "custom",
        "weekly_frequency": max((len(week["days"]) for week in weeks), default=1),
        "instructions": "",
        "days": days,
        "weeks": weeks,
    }
    from svc.schemas import CoachProgramDraftIn

    validated = CoachProgramDraftIn.model_validate(draft).model_dump()
    return _save_imported_draft(db, coach_account_id, assignment_id, validated, replace)


def _save_imported_draft(
    db: Any,
    coach_account_id: str,
    assignment_id: str,
    draft: dict[str, Any],
    replace: bool,
) -> dict[str, Any] | None:
    try:
        if replace:
            try:
                return coach_program_drafts.replace_program_draft(
                    db, coach_account_id, assignment_id, draft, preserve_source_template=False
                )
            except coach_program_drafts.ProgramDraftNotFound:
                return coach_program_drafts.create_program_draft(db, coach_account_id, assignment_id, draft)
        return coach_program_drafts.create_program_draft(db, coach_account_id, assignment_id, draft)
    except coach_program_drafts.ProgramDraftAlreadyExists as error:
        raise ProgramImportError("assignment.program_draft_exists.v1", str(error), 409) from error


def template_csv_bytes() -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(TEMPLATE_COLUMNS)
    writer.writerow((1, 1, "Example: Push", 1, "Bench Press", 3, "6-8", 2, "", "", "", 120, "2-0-1", "Pause briefly at the chest"))
    return output.getvalue().encode("utf-8-sig")


def template_xlsx_bytes() -> bytes:
    workbook = Workbook()
    program = workbook.active
    program.title = "Program"
    program.append(TEMPLATE_COLUMNS)
    example = workbook.create_sheet("Example")
    example.append(TEMPLATE_COLUMNS)
    example.append((1, 1, "Push", 1, "Bench Press", 3, "6-8", 2, "", "", "", 120, "2-0-1", "Pause briefly at the chest"))
    for worksheet in (program, example):
        for column in ("G", "H", "I"):
            worksheet.column_dimensions[column].number_format = "@"
            for row_number in range(1, MAX_PROGRAM_IMPORT_ROWS + 2):
                worksheet[f"{column}{row_number}"].number_format = "@"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _build_import_result(source: ReaderOutput, draft_exists: bool, rows: list[dict[str, Any]]) -> dict[str, Any]:
    public_rows = [
        {key: value for key, value in row.items() if key != "suggestion_search_terms"}
        for row in rows
    ]
    rows_by_week = {
        week: [row for row in public_rows if row["week"] == week]
        for week in source.detected_weeks
    }
    return {
        "detected_tabs": source.detected_tabs,
        "selected_tab": source.selected_tab,
        "requires_tab_choice": source.requires_tab_choice,
        "detected_weeks": source.detected_weeks,
        "selected_week": source.selected_week,
        "weeks_not_imported": source.weeks_not_imported,
        "confirm_layout": source.confirm_layout,
        "import_mode": source.import_mode,
        "layout_days": source.layout_days,
        "layout_confidence": source.layout_confidence,
        "draft_exists": draft_exists,
        **program_import_capabilities(),
        "rows": public_rows,
        "rows_by_week": rows_by_week,
        "errors": [error for row in rows for error in row["errors"]],
        "unresolved_names": _unresolved_names(rows),
    }


def program_import_capabilities() -> dict[str, bool]:
    return {
        "freeform_available": program_import_ai.resolve_enable_gate().enabled,
        "text_available": program_import_ai.resolve_text_enable_gate().enabled,
    }


def _run_text_import(
    db: Any, request: ProgramImportRequest, text: str
) -> tuple[ReaderOutput, list[dict[str, Any]]]:
    from svc.llm import InferenceScope, inference_turn

    scope = InferenceScope(
        account_id=request.coach_account_id,
        role="coach",
        purpose="coach_program_import",
        store=db,
        client=request.client,
    )
    with inference_turn(scope, background_tasks=request.background_tasks):
        source = PlainTextProgramImportReader(scope).read(text)
        rows = validate_and_resolve_rows(
            db, request.coach_account_id, source.rows, text_import=True
        )
        _add_model_suggestions(db, rows, scope)
    return source, rows


def _text_reader_output(
    reply: program_import_ai.TextProgramImportReply, text: str
) -> ReaderOutput:
    context = _empty_freeform_context()
    heading_scan = _text_heading_occurrences(reply.rows, text)
    inferred_weeks = _text_repeated_block_weeks(reply.rows, text, heading_scan)
    initially_inferred_weeks = inferred_weeks
    identical_blocks = _text_has_strong_identical_blocks(reply.rows, text, inferred_weeks, heading_scan)
    model_repeated_days = _text_model_assigns_repeated_headings_as_days(
        reply, heading_scan
    )
    ambiguous_model_week_reading = identical_blocks and model_repeated_days
    if ambiguous_model_week_reading:
        inferred_weeks = {}
    structural_weeks = bool(inferred_weeks) and not _text_has_explicit_week_titles(text)
    weeks = _resolved_week_labels(reply.layout.weeks, context)
    model_weeks = list(dict.fromkeys(
        label for week in reply.layout.weeks if (label := _canonical_week_label(week))
    ))
    if not weeks and len(model_weeks) == 1 and (_week_number(model_weeks[0]) or 1) != 1:
        weeks = model_weeks
    if inferred_weeks:
        inferred_labels = {f"Week {week}" for week in inferred_weeks.values()}
        model_labels = {label for row in reply.rows if row.source_row not in inferred_weeks
                        if (label := _canonical_week_label(row.week))}
        weeks = sorted(inferred_labels | model_labels, key=lambda label: (_week_number(label) or MAX_PROGRAM_WEEKS + 1, label))
    selected_week = _default_import_week(weeks)
    day_assignments, day_conflict = _text_day_number_assignments(
        reply.rows, inferred_weeks, structural_weeks, heading_scan
    )
    rows = _text_program_rows(reply, text, weeks, inferred_weeks, day_assignments)
    week_conflict = _text_week_inference_conflicts(reply.rows, inferred_weeks)
    progression_blocks = _text_has_strong_progression_blocks(
        reply.rows, text, initially_inferred_weeks, heading_scan
    )
    model_confirmation = _text_layout_confirmation(reply, context, text, heading_scan)
    same_name_day_read = progression_blocks and _text_repeated_heading_read_as_days(
        reply, heading_scan, initially_inferred_weeks, text
    )
    progression_agrees = same_name_day_read or (
        progression_blocks and _text_repeated_block_assignments_agree(reply.rows, day_assignments, inferred_weeks)
    )
    confirm_layout = (
        model_confirmation
        or _has_unplaced_text_prescription_lines(text, reply.rows)
        or (week_conflict and not progression_blocks)
        or (
            _text_has_explicit_week_titles(text)
            and len(weeks) > 1
            and any(
                _week_number(_canonical_week_label(row.week)) != inferred_weeks.get(row.source_row)
                for row in reply.rows
            )
        )
        or (day_conflict and not same_name_day_read)
        or _text_headingless_layout_needs_confirmation(text, reply.rows, heading_scan)
        or identical_blocks
        or ambiguous_model_week_reading
    )
    if progression_agrees:
        confirm_layout = False
    return ReaderOutput(
        [], None, False, _renumber_out_of_range_orders(rows),
        detected_weeks=weeks,
        selected_week=selected_week,
        weeks_not_imported=[week for week in weeks if week != selected_week],
        confirm_layout=confirm_layout,
        layout_days=_text_layout_days(reply),
        layout_confidence=reply.layout.confidence,
        import_mode="text",
    )


def _text_program_rows(
    reply: program_import_ai.TextProgramImportReply,
    text: str,
    weeks: list[str],
    inferred_weeks: dict[int, int],
    day_numbers: dict[tuple[str | None, int], int],
) -> list[ProgramImportRow]:
    context = _empty_freeform_context()
    selected_week = _default_import_week(weeks)
    week_lookup = {_week_identity(week): week for week in weeks}
    rows = []
    for model_row in reply.rows:
        source_text = _line_text(text, model_row.source_row)
        if _text_heading_without_prescription(source_text, model_row.day_name):
            continue
        exercise = _text_exercise_name(source_text, model_row.exercise)
        if not exercise:
            continue
        cleaned_day_name = _clean_day_name(model_row.day_name) or model_row.day_name
        row = model_row.model_copy(update={"exercise": exercise, "day_name": cleaned_day_name})
        inferred_week = inferred_weeks.get(model_row.source_row)
        week_label = f"Week {inferred_week}" if inferred_week else (
            _matched_row_week(model_row.week, week_lookup, selected_week) if model_row.week else selected_week
        )
        canonical = _program_import_row(
            row,
            week_label,
            FreeFormSourceRow({}, source_text),
            context,
        )
        canonical = replace(canonical, input_exercise_name=exercise)
        canonical = replace(canonical, load_kg=model_row.load_kg)
        if load_text := (
            _text_load_value(source_text)
            if model_row.load_kg in (None, "")
            else None
        ):
            canonical = replace(canonical, notes=load_text)
        hold = _text_hold_range(source_text)
        rep_list = _text_per_set_rep_list(source_text)
        if hold is not None and hold[1] <= MAX_REPS:
            canonical = replace(canonical, reps_min=hold[0], reps_max=hold[1])
        elif hold is not None:
            canonical = replace(
                canonical,
                reps_min=DEFAULT_TARGET_REPS_MIN,
                reps_max=DEFAULT_TARGET_REPS_MAX,
            )
        elif rep_list:
            canonical = replace(canonical, reps_min=min(rep_list), reps_max=max(rep_list))
        elif rep_bounds := _text_prescribed_rep_bounds(source_text):
            canonical = replace(canonical, reps_min=rep_bounds[0], reps_max=rep_bounds[1])
        source_markers = _text_source_approximation_markers(source_text)
        if model_row.load_kg not in (None, ""):
            recognized_load = _text_load_value(source_text)
            source_markers = [
                marker for marker in source_markers
                if not (
                    marker.original == recognized_load
                    and marker.code in {
                        "program_import.load_preserved_as_note.v1",
                        "program_import.unsupported_preserved_as_note.v1",
                    }
                )
            ]
        if hold is not None and hold[1] > MAX_REPS:
            duration = f"{hold[0]} sec hold" if hold[0] == hold[1] else f"{hold[0]}-{hold[1]} sec hold"
            source_markers.append(Approximation(
                "program_import.reps_approximated.v1", "reps", duration
            ))
        source_has_effort_approximation = any(
            marker.code == "program_import.effort_approximated.v1" for marker in source_markers
        )
        markers = [
            marker for marker in canonical.approximation_markers
            if not (
                (
                    model_row.load_kg not in (None, "")
                    and marker.code == "program_import.load_preserved_as_note.v1"
                )
                or (
                    marker.code == "program_import.effort_approximated.v1"
                    and (
                        re.search(r"\bAMRAP\b", source_text, re.IGNORECASE)
                        or not source_has_effort_approximation
                    )
                )
            )
        ]
        markers.extend(source_markers)
        canonical = replace(
            canonical,
            approximation_markers=tuple(_dedupe_approximation_markers(markers)),
        )
        inferred_day = _text_day_number(row, day_numbers, inferred_week)
        parsed_day = _parse_integer(row.day)
        if inferred_day is not None and (
            parsed_day is None or not 1 <= parsed_day <= MAX_PROGRAM_DAYS
            or parsed_day != inferred_day
        ):
            canonical = replace(canonical, day=inferred_day)
        rows.append(canonical)
    return rows


def _text_is_day_heading(source_text: str, day_name: Any) -> bool:
    name = _text_heading_label(_cell_text(day_name))
    if not name:
        return False
    return _day_name_key(_text_heading_label(source_text)) == _day_name_key(name)


def _text_heading_label(value: str) -> str:
    label = re.sub(r"^\s*\[[^\]\n]+\]\s*", "", value)
    label = re.sub(r"^\s*[•▪◦]\s*", "", label)
    label = re.sub(r"^\s*#+\s*", "", label)
    label = re.sub(r"[*_`~]", "", label).strip()
    label = re.sub(r"\s*\((?:week|wk)\s*\d+\)\s*$", "", label, flags=re.IGNORECASE)
    label = re.sub(r"\s+day\s*$", "", label, flags=re.IGNORECASE)
    return _clean_day_name(label)


def _text_heading_without_prescription(value: str, day_name: Any = None) -> bool:
    if _TEXT_PRESCRIPTION_PATTERN.search(value) or re.search(
        r"\d+\s*(?:sets?|مجموعات|جولات)?\s*[x×]\s*\d+", normalize_digits(value), re.IGNORECASE
    ):
        return False
    label = _text_heading_label(value)
    explicit = bool(re.match(r"\s*#+\s*|\s*\*\*", value))
    explicit = explicit or bool(re.search(r"\s+day\s*$|\(\s*(?:week|wk)\s*\d+\s*\)\s*$", value, re.IGNORECASE))
    return bool(label and (explicit or _text_is_day_heading(value, day_name)))


def _text_repeated_block_weeks(
    model_rows: list[Any],
    text: str,
    heading_scan: list[tuple[int, str, list[Any]]] | None = None,
) -> dict[int, int]:
    week_titles: list[tuple[int, int]] = []
    for line_number, line in enumerate(text.splitlines(), 1):
        number = _text_standalone_week_number(line, start_title=(line_number == 1))
        if number is not None:
            week_titles.append((line_number, number))
    if week_titles:
        inferred = {}
        for row in model_rows:
            title = next((entry for entry in reversed(week_titles) if entry[0] < row.source_row), None)
            if title:
                inferred[row.source_row] = title[1]
        return inferred
    return _text_implicit_block_weeks(model_rows, text, heading_scan)


def _text_has_explicit_week_titles(text: str) -> bool:
    return any(
        _text_standalone_week_number(line, start_title=(line_number == 1)) is not None
        for line_number, line in enumerate(text.splitlines(), 1)
    )


def _text_implicit_block_weeks(
    model_rows: list[Any], text: str, heading_scan: list[tuple[int, str, list[Any]]] | None = None
) -> dict[int, int]:
    occurrences = heading_scan if heading_scan is not None else _text_heading_occurrences(model_rows, text)
    if len(occurrences) < 2:
        return {}
    block_size = _text_repeated_heading_block_size(occurrences)
    if block_size is None or not _text_blocks_correspond(occurrences, block_size, text):
        return {}
    return {
        row.source_row: index // block_size + 1
        for index, occurrence in enumerate(occurrences)
        for row in occurrence[2]
    }


def _text_repeated_heading_block_size(occurrences: list[tuple[int, str, list[Any]]]) -> int | None:
    return next((size for size in range(1, len(occurrences) // 2 + 1)
                 if len(occurrences) % size == 0 and all(
                     occurrences[index][1] == occurrences[index % size][1]
                     for index in range(size, len(occurrences))
                 )), None)


def _text_heading_occurrences(model_rows: list[Any], text: str) -> list[tuple[int, str, list[Any]]]:
    day_names = {
        _day_name_key(name): name
        for row in model_rows
        if (name := _clean_day_name(row.day_name))
    }
    headings = []
    for line_number, line in enumerate(text.splitlines(), 1):
        name = next((key for key, value in day_names.items() if _text_is_day_heading(line, value)), None)
        if name:
            headings.append((line_number, name, []))
    for row in model_rows:
        heading = next((entry for entry in reversed(headings) if entry[0] < row.source_row), None)
        if heading:
            heading[2].append(row)
    return [heading for heading in headings if heading[2]]


def _text_blocks_correspond(occurrences: list[tuple[int, str, list[Any]]], block_size: int, text: str) -> bool:
    return _text_block_match_ratio(occurrences, block_size, text) >= 0.6


def _text_block_match_ratio(occurrences: list[tuple[int, str, list[Any]]], block_size: int, text: str) -> float:
    first_block = occurrences[:block_size]
    compared = occurrences[block_size:]
    total_rows = matched_rows = 0
    for index, current in enumerate(compared):
        reference = first_block[index % block_size]
        if len(reference[2]) != len(current[2]):
            return 0.0
        for first_row, current_row in zip(reference[2], current[2], strict=True):
            total_rows += 1
            first_name = _text_exercise_name(_line_text(text, first_row.source_row), first_row.exercise)
            current_name = _text_exercise_name(_line_text(text, current_row.source_row), current_row.exercise)
            matched_rows += _text_exercise_names_equivalent(first_name, current_name)
    return matched_rows / total_rows if total_rows else 0.0


def _text_exercise_names_equivalent(first: str, second: str) -> bool:
    def tokens(name: str) -> set[str]:
        return {token[:-1] if len(token) > 4 and token.endswith("s") else token
                for token in re.findall(r"\w+", name.casefold())}

    first_tokens, second_tokens = tokens(first), tokens(second)
    if not first_tokens or not second_tokens:
        return False
    overlap = len(first_tokens & second_tokens)
    shared_movement_tokens = (first_tokens & second_tokens) - {
        "barbell", "dumbbell", "machine", "cable", "smith", "band", "plate", "bodyweight",
        "incline", "decline", "seated", "standing", "single", "front", "rear", "side",
    }
    return (
        overlap / len(first_tokens | second_tokens) >= 0.5
        or overlap == min(len(first_tokens), len(second_tokens))
        or bool(shared_movement_tokens)
    )


def _text_standalone_week_number(line: str, *, start_title: bool) -> int | None:
    normalized = normalize_digits(line).strip()
    pattern = r"(?:week|wk|الأسبوع|اسبوع)\s*(\d+|الأول|الأولى|التاني|الثاني|التالت|الثالث)"
    match = re.fullmatch(pattern, normalized, re.IGNORECASE)
    if not match and start_title:
        match = re.match(rf"^\s*{pattern}(?=\s*[-—:：]|\s+.+$)", normalized, re.IGNORECASE)
    if not match:
        return None
    token = match.group(1).casefold()
    ordinal = {"الأول": 1, "الأولى": 1, "التاني": 2, "الثاني": 2, "التالت": 3, "الثالث": 3}
    return int(token) if token.isdigit() else ordinal.get(token)


def _text_week_inference_conflicts(rows: list[Any], inferred: dict[int, int]) -> bool:
    return any(
        row.week and row.source_row in inferred
        and _week_number(_canonical_week_label(row.week)) != inferred[row.source_row]
        for row in rows
    )


def _text_repeated_block_assignments_agree(
    rows: list[Any], day_numbers: dict[tuple[str | None, int], int], inferred: dict[int, int]
) -> bool:
    if not rows or any(
        row.source_row not in inferred
        or _week_number(_canonical_week_label(row.week)) != inferred[row.source_row]
        for row in rows
    ):
        return False
    return all(
        _parse_integer(row.day) == _text_day_number(row, day_numbers, inferred[row.source_row])
        for row in rows
    )


def _text_repeated_heading_read_as_days(
    reply: Any, headings: list[tuple[int, str, list[Any]]], inferred: dict[int, int], text: str
) -> bool:
    if _text_has_explicit_week_titles(text) or len(set(inferred.values())) < 2:
        return False
    block_size = _text_repeated_heading_block_size(headings)
    if not block_size or len(reply.layout.weeks) > 1 or len(reply.layout.days) != len(headings):
        return False
    names = [_day_name_key(day.name) for day in reply.layout.days]
    if not names or len(set(names)) != 1 or len({heading[1] for heading in headings}) != 1:
        return False
    if any(_parse_integer(day.day) != index for index, day in enumerate(reply.layout.days, 1)):
        return False
    layout_weeks = {_week_identity(day.week) for day in reply.layout.days if day.week}
    row_weeks = {_week_identity(row.week) for row in reply.rows if row.week}
    if len(layout_weeks | row_weeks) > 1:
        return False
    return all(
        {_parse_integer(row.day) for row in heading[2]} == {index}
        for index, heading in enumerate(headings, 1)
    )


def _text_model_assigns_repeated_headings_as_days(
    reply: Any, headings: list[tuple[int, str, list[Any]]]
) -> bool:
    if len(headings) < 2 or len(reply.layout.days) != len(headings):
        return False
    if any(_parse_integer(day.day) != index for index, day in enumerate(reply.layout.days, 1)):
        return False
    if any(_day_name_key(_clean_day_name(day.name)) != name
           for day, (_, name, _) in zip(reply.layout.days, headings, strict=True)):
        return False
    week_ids = {_week_identity(week) for week in reply.layout.weeks if week}
    week_ids.update(_week_identity(day.week) for day in reply.layout.days if day.week)
    week_ids.update(_week_identity(row.week) for row in reply.rows if row.week)
    if len(week_ids) > 1:
        return False
    return all(
        {_parse_integer(row.day) for row in heading[2]} == {index}
        for index, heading in enumerate(headings, 1)
    )


def _text_repeated_heading_blocks(
    rows: list[Any], text: str, inferred: dict[int, int],
    headings: list[tuple[int, str, list[Any]]],
) -> tuple[list[tuple[int, str, list[Any]]], int] | None:
    if not inferred or _text_has_explicit_week_titles(text):
        return None
    block_size = _text_repeated_heading_block_size(headings)
    if block_size is None or _text_block_match_ratio(headings, block_size, text) < 0.75:
        return None
    return headings, block_size


def _text_has_strong_progression_blocks(
    rows: list[Any], text: str, inferred: dict[int, int], headings: list[tuple[int, str, list[Any]]]
) -> bool:
    repeated = _text_repeated_heading_blocks(rows, text, inferred, headings)
    return bool(repeated and _text_blocks_show_progression(*repeated, text))


def _text_has_strong_identical_blocks(
    rows: list[Any], text: str, inferred: dict[int, int], headings: list[tuple[int, str, list[Any]]]
) -> bool:
    repeated = _text_repeated_heading_blocks(rows, text, inferred, headings)
    return bool(repeated and not _text_blocks_show_progression(*repeated, text))


def _text_blocks_show_progression(
    headings: list[tuple[int, str, list[Any]]], block_size: int, text: str
) -> bool:
    first_block = headings[:block_size]
    for index, current in enumerate(headings[block_size:]):
        reference = first_block[index % block_size]
        signatures = (
            (_text_prescription_signature(_line_text(text, first_row.source_row)),
             _text_prescription_signature(_line_text(text, current_row.source_row)))
            for first_row, current_row in zip(reference[2], current[2], strict=True)
        )
        if any(first != second for first, second in signatures):
            return True
    return False


def _text_prescription_signature(source_text: str) -> str:
    normalized = normalize_digits(source_text).casefold()
    match = _TEXT_PRESCRIPTION_PATTERN.search(normalized)
    if match:
        return re.sub(r"\s+", "", match.group(0).replace("×", "x"))
    if hold := _text_hold_range(source_text):
        return f"hold:{hold[0]}-{hold[1]}"
    if rep_list := _text_per_set_rep_list(source_text):
        return "reps:" + "/".join(map(str, rep_list))
    return ""


def _text_day_number_assignments(
    model_rows: list[Any], inferred_weeks: dict[int, int], structural_weeks: bool,
    heading_scan: list[tuple[int, str, list[Any]]],
) -> tuple[dict[tuple[str | None, int], int], bool]:
    rows = sorted(model_rows, key=lambda row: row.source_row)
    headings = [(line, name) for line, name, _ in heading_scan]

    occurrences: dict[tuple[str | None, int], list[Any]] = {}
    fallback_occurrence: dict[str | None, int] = {}
    fallback_name: dict[str | None, str] = {}
    row_occurrences = {}
    for row in rows:
        week = f"number:{inferred_weeks[row.source_row]}" if row.source_row in inferred_weeks else (
            _week_identity(row.week) if row.week else None
        )
        heading = next((entry for entry in reversed(headings) if entry[0] < row.source_row), None)
        if heading:
            occurrence = heading[0]
        else:
            key = _day_name_key(_clean_day_name(row.day_name))
            if key != fallback_name.get(week):
                fallback_occurrence[week] = fallback_occurrence.get(week, 0) + 1
                fallback_name[week] = key
            occurrence = -fallback_occurrence.get(week, 1)
        occurrence_key = (week, occurrence)
        row_occurrences[row.source_row] = occurrence_key
        occurrences.setdefault(occurrence_key, []).append(row)

    ordered_occurrences: dict[str | None, list[tuple[str | None, int]]] = {}
    for occurrence_key in sorted(occurrences, key=lambda key: min(row.source_row for row in occurrences[key])):
        ordered_occurrences.setdefault(occurrence_key[0], []).append(occurrence_key)

    assignments = {}
    conflicts = False
    for week, week_occurrences in ordered_occurrences.items():
        model_days = {}
        for occurrence_key in week_occurrences:
            day_values = {
                day for row in occurrences[occurrence_key]
                if (day := _parse_integer(row.day)) is not None and 1 <= day <= MAX_PROGRAM_DAYS
            }
            model_days[occurrence_key] = day_values
        day_counts = {}
        for day_values in model_days.values():
            if len(day_values) == 1:
                day = next(iter(day_values))
                day_counts[day] = day_counts.get(day, 0) + 1
        for index, occurrence_key in enumerate(week_occurrences, start=1):
            day_values = model_days[occurrence_key]
            model_day = next(iter(day_values)) if len(day_values) == 1 else None
            if model_day is not None and day_counts[model_day] == 1 and model_day != index:
                conflicts = True
            if structural_weeks and model_day is not None and model_day != index:
                resolved_day = index
            elif model_day is not None and day_counts[model_day] == 1:
                resolved_day = model_day
            else:
                resolved_day = index
            for row in occurrences[occurrence_key]:
                assignments[(week, row.source_row)] = resolved_day
    return assignments, conflicts


def _text_day_number(
    row: Any, numbers: dict[tuple[str | None, int], int], inferred_week: int | None
) -> int | None:
    week = f"number:{inferred_week}" if inferred_week else _week_identity(row.week) if row.week else None
    return numbers.get((week, row.source_row))


def _text_layout_days(reply: Any) -> list[dict[str, Any]]:
    layout_days = []
    day_counts: dict[str | None, set[int]] = {}
    day_positions: dict[str | None, int] = {}
    for day in reply.layout.days:
        week = _week_identity(day.week) if getattr(day, "week", None) else None
        parsed = _parse_integer(day.day)
        day_positions[week] = day_positions.get(week, 0) + 1
        used = day_counts.setdefault(week, set())
        if parsed is None or not 1 <= parsed <= MAX_PROGRAM_DAYS or parsed in used:
            parsed = day_positions[week]
        used.add(parsed)
        layout_days.append({
            "day": parsed,
            "name": _clean_day_name(day.name) or day.name,
        })
    return layout_days


_TEXT_EXERCISE_BOUNDARY = re.compile(
    r"\s+(?=(?:\d+\s*(?:sets?|مجموعات|جولات)(?:\s*[x×]|\s+(?:of\s+)?\d+)|"
    r"\d+\s*[x×]\s*(?:\d+|AMRAP)|\d+\s*(?:reps?|عدة|عدات)\b|"
    r"(?:\d+\s*[/,]\s*)+\d+\s*(?:reps?|عدة|عدات)\b|"
    r"x\s*\d+(?:\s*[,/]\s*\d+)+))",
    re.IGNORECASE,
)


def _text_exercise_name(source_text: str, model_exercise: Any) -> str:
    model_name = _cell_text(model_exercise).strip(" *`_")
    normal_source = _normalise_exercise_source(source_text)
    if model_name and _normalise_exercise_source(model_name) in normal_source:
        return model_name[:MAX_PROGRAM_IMPORT_EXERCISE_NAME_LENGTH]
    line = re.sub(r"^\s*(?:[-*•]+\s*)?", "", source_text.strip())
    line = re.sub(r"^\s*(?:\d+[.)]|[A-Z]\d+[.)])\s*", "", line, flags=re.IGNORECASE)
    line = re.sub(r"^[^\w]+", "", line)
    line = re.sub(r"^\s*\d+\s*[x×]\s*\d+\s+", "", line)
    line = re.sub(r"\s*[:：]\s*", " ", line, count=1)
    match = _TEXT_EXERCISE_BOUNDARY.search(line)
    if match:
        line = line[:match.start()]
    line = re.sub(r"\s*[—–:-]\s*$", "", line)
    line = re.split(r"\s+[—–]\s+", line, maxsplit=1)[0].strip(" *`_")
    candidate = line.strip(" *`_")
    return candidate[:MAX_PROGRAM_IMPORT_EXERCISE_NAME_LENGTH]


def _normalise_exercise_source(value: str) -> str:
    return " ".join(re.sub(r"[*_`~]", "", value).split())


def _text_per_set_rep_list(source_text: str) -> list[int]:
    normal = normalize_digits(source_text)
    token_matches = list(re.finditer(r"\d+\s*(?:sets?|مجموعات|جولات)?\s*[x×]\s*(\d+(?:\s*/\s*\d+)+)", normal, re.IGNORECASE))
    reps_matches = list(re.finditer(r"\b(\d+(?:\s*/\s*\d+)+)\s*(?:reps?|عدة|عدات)\b", normal, re.IGNORECASE))
    # Multiple prescription tokens are ambiguous (e.g. a top set and backoff).
    if token_matches:
        if len(token_matches) != 1:
            return []
        match = token_matches[0]
    elif len(reps_matches) == 1:
        match = reps_matches[0]
    else:
        return []
    return [int(value) for value in re.findall(r"\d+", match.group(1))]


def _text_prescribed_rep_bounds(source_text: str) -> tuple[int, int] | None:
    normal = normalize_digits(source_text)
    if _text_timed_non_hold_prescription(source_text):
        return None
    prescriptions = list(re.finditer(
        r"\d+\s*(?:sets?|مجموعات|جولات)?\s*[x×]\s*(\d+)(?:\s*[-–]\s*(\d+))?\b",
        normal, re.IGNORECASE,
    ))
    if not prescriptions:
        prescriptions = list(re.finditer(r"\b(\d+)(?:\s*[-–]\s*(\d+))?\s*(?:reps?|عدة|عدات)\b", normal, re.IGNORECASE))
    if len(prescriptions) != 1:
        return None
    first = int(prescriptions[0].group(1))
    second = int(prescriptions[0].group(2) or first)
    return min(first, second), max(first, second)


def _text_timed_non_hold_prescription(source_text: str) -> bool:
    normal = normalize_digits(source_text)
    return (
        _text_hold_range(source_text) is None
        and bool(re.search(r"\d+\s*[x×]\s*\d+(?:\s*[-–]\s*\d+)?\s*(?:sec(?:ond)?s?|s|min(?:ute)?s?)\b", normal, re.IGNORECASE))
    )


def _text_hold_range(source_text: str) -> tuple[int, int] | None:
    normal = normalize_digits(source_text)
    hold_cue = re.search(r"\b(?:hold|plank|wall\s+sit|dead\s+hang|l[ -]?sit|hollow\s+hold|isometric)\b", normal, re.IGNORECASE)
    if not hold_cue:
        return None
    match = re.search(
        r"\d+\s*[x×]\s*(\d+)(?:\s*[-–]\s*(\d+))?\s*(s|sec(?:ond)?s?|min(?:ute)?s?)\b|"
        r"\b(\d+)(?:\s*[-–]\s*(\d+))?\s*(s|sec(?:ond)?s?|min(?:ute)?s?)\s*hold\b",
        normal, re.IGNORECASE,
    )
    if not match:
        return None
    first, second, unit = (match.group(1), match.group(2), match.group(3)) if match.group(1) else (match.group(4), match.group(5), match.group(6))
    low, high = int(first), int(second or first)
    if unit.casefold().startswith("min"):
        low, high = low * 60, high * 60
    return min(low, high), max(low, high)


def _text_source_approximation_markers(source_text: str) -> list[Approximation]:
    normal = normalize_digits(source_text)
    source_fields = {}
    if rep_list := _text_per_set_rep_list(source_text):
        source_fields["reps_by_set"] = "/".join(map(str, rep_list))
    elif re.search(r"\bAMRAP\b", normal, re.IGNORECASE):
        source_fields["reps"] = "AMRAP"
    markers = _source_approximation_markers(FreeFormSourceRow(source_fields, source_text))
    if source_fields.get("reps") and _is_unsupported_rep_prescription(source_fields["reps"]):
        markers.append(Approximation("program_import.reps_approximated.v1", "reps", "AMRAP"))
    row = TextSourceMarkerRow(normal)
    markers.extend(_detect_unsupported_prescriptions(row, markers, normal))
    if load := _text_load_value(source_text):
        markers.append(Approximation(
            "program_import.unsupported_preserved_as_note.v1", "prescription", load
        ))
        markers.append(Approximation("program_import.load_preserved_as_note.v1", "load_kg", load))
    if _text_timed_non_hold_prescription(source_text):
        duration = re.search(
            r"\d+\s*[x×]\s*\d+(?:\s*[-–]\s*\d+)?\s*(?:sec(?:ond)?s?|s|min(?:ute)?s?)\b",
            normal, re.IGNORECASE,
        )
        if duration:
            markers.append(Approximation(
                "program_import.unsupported_preserved_as_note.v1", "prescription", duration.group(0)
            ))
    effort = re.search(
        r"\b(?:to|until)\s+(?:near[ -]?)?failure\b|\bnear[ -]failure\b|"
        r"\bRIR\s*\d+(?:\.\d+)?\s*[-–]\s*\d+(?:\.\d+)?|"
        r"\b\d+(?:\.\d+)?\s*[-–]\s*\d+(?:\.\d+)?\s*in reserve\b|"
        r"\bRPE\s*\d+(?:\.\d+)?\s*[-–]\s*\d+(?:\.\d+)?|"
        r"حتى\s+قرب\s+الفشل|حتى\s+الفشل|قرب\s+الفشل|لحد\s+الفشل",
        normal, re.IGNORECASE,
    )
    if effort and not re.search(r"\bAMRAP\b", normal, re.IGNORECASE):
        markers.append(Approximation("program_import.effort_approximated.v1", "rir", effort.group(0)))
    if _text_hold_range(source_text) is not None:
        markers.append(Approximation("program_import.effort_approximated.v1", "rir", "hold duration"))
    return _dedupe_approximation_markers(markers)


def _text_load_value(source_text: str) -> str | None:
    normal = normalize_digits(source_text)
    load = re.search(
        r"(?:\d+(?:\.\d+)?\s*%\s*1RM|bodyweight\s*\+\s*\d+(?:\.\d+)?\s*(?:kg|lb)?|\+\s*\d+(?:\.\d+)?\s*(?:kg|lb)|\b\d+(?:\.\d+)?\s*kg\b)",
        normal, re.IGNORECASE,
    )
    return load.group(0) if load else None


def _text_layout_confirmation(
    reply: program_import_ai.TextProgramImportReply, context: FreeFormGridContext,
    text: str, heading_scan: list[tuple[int, str, list[Any]]],
) -> bool:
    if _text_day_headings_cover_rows(heading_scan, reply.rows, require_distinct_names=True):
        return context.layout_ambiguous or _has_unplaced_exercise_rows(context, reply.rows)
    return _needs_layout_confirmation(
        context,
        reply.rows,
        reply.layout.confirm_layout,
        reply.layout.confidence,
    )


def _text_day_headings_cover_rows(
    heading_scan: list[tuple[int, str, list[Any]]], rows: list[Any], *, require_distinct_names: bool
) -> bool:
    headings = [(line, name) for line, name, _ in heading_scan]
    if not headings or not rows:
        return False
    for row in rows:
        expected = _day_name_key(_clean_day_name(row.day_name))
        preceding = next(
            (name for line, name in reversed(headings) if line < row.source_row),
            None,
        )
        if not expected or preceding != expected:
            return False
    return not require_distinct_names or not any(_is_generic_day_name(name) for _, name in headings)


def _text_headingless_layout_needs_confirmation(
    text: str, rows: list[Any], heading_scan: list[tuple[int, str, list[Any]]]
) -> bool:
    if _text_day_headings_cover_rows(heading_scan, rows, require_distinct_names=False):
        return False
    prescription_count = sum(
        bool(_TEXT_PRESCRIPTION_PATTERN.search(line)) for line in text.splitlines()
    )
    has_group_cue = any(
        re.search(r"^\s*[-*_]{3,}|\b(?:alternate|alternative)\s+split\b|"
                  r"\bpair\s+as\s+accessories\b|[?؟]\s*$", line, re.IGNORECASE)
        for line in text.splitlines()
    )
    return has_group_cue or prescription_count > 8


_TEXT_PRESCRIPTION_PATTERN = re.compile(
    r"(?<![\w])(?:[0-9٠-٩۰-۹]+)\s*[xX×]\s*"
    r"(?:[0-9٠-٩۰-۹]+)(?:\s*[-–]\s*[0-9٠-٩۰-۹]+)?"
)


def _has_unplaced_text_prescription_lines(text: str, model_rows: list[Any]) -> bool:
    placed_lines = {row.source_row for row in model_rows}
    return any(
        line.strip() and line_number not in placed_lines and _TEXT_PRESCRIPTION_PATTERN.search(line)
        for line_number, line in enumerate(text.split("\n"), 1)
    )


def _empty_freeform_context() -> FreeFormGridContext:
    return FreeFormGridContext(
        cells_by_row={},
        columns={},
        header_row=None,
        week_labels=[],
        title_week_labels=[],
        explicit_week_labels=[],
        implicit_week_columns={},
        implicit_week_count=0,
        day_names={},
        day_numbers={},
        row_facts=FreeFormGridRows(
            columns_by_row={},
            header_rows=frozenset(),
            exercise_data_rows=(),
            exercise_data_columns_by_schema={},
            section_title_rows=frozenset(),
        ),
    )


def _line_text(text: str, line_number: int) -> str:
    lines = text.split("\n")
    return lines[line_number - 1] if 0 < line_number <= len(lines) else ""


def _grid_implicit_week_count(grid: dict[str, Any]) -> int:
    _header_row, columns = _source_columns(_grid_cells_by_row(grid))
    return max(_implicit_week_block_columns(columns).values(), default=0)


def _freeform_grid_context(
    grid: dict[str, Any], layout_days: list[Any]
) -> FreeFormGridContext:
    cells_by_row = _grid_cells_by_row(grid)
    header_row, columns = _source_columns(cells_by_row)
    row_facts = _freeform_grid_rows(cells_by_row, columns)
    layout_day_numbers = _layout_day_numbers(layout_days) if layout_days else None
    day_names, day_numbers = _source_day_labels(
        cells_by_row, header_row, columns, layout_day_numbers, row_facts
    )
    _add_layout_day_labels(layout_days, day_names, day_numbers)
    implicit_week_columns = _implicit_week_block_columns(columns)
    implicit_week_count = max(implicit_week_columns.values(), default=0)
    title_weeks, explicit_weeks = _source_week_labels(
        cells_by_row, header_row, columns, row_facts
    )
    title_weeks.extend(_week_labels_in_cell(_cell_text(grid.get("sheet", ""))))
    title_weeks = list(dict.fromkeys(title_weeks))
    weeks = title_weeks + explicit_weeks
    weeks.extend(_week_number_labels(implicit_week_count))
    weeks = list(dict.fromkeys(weeks))
    return FreeFormGridContext(
        cells_by_row=cells_by_row,
        columns=columns,
        header_row=header_row,
        week_labels=weeks,
        title_week_labels=title_weeks,
        explicit_week_labels=explicit_weeks,
        implicit_week_columns=implicit_week_columns,
        implicit_week_count=implicit_week_count,
        day_names=day_names,
        day_numbers=day_numbers,
        row_facts=row_facts,
        preamble_text=_sheet_preamble_text(cells_by_row, header_row),
        layout_ambiguous=_grid_layout_is_ambiguous(columns) or implicit_week_count > 1,
    )


def _grid_cells_by_row(grid: dict[str, Any]) -> dict[int, dict[int, str]]:
    rows: dict[int, dict[int, str]] = {}
    for cell in grid.get("cells", []):
        match = re.fullmatch(r"([A-Z]+)(\d+)", cell["address"])
        if match:
            row_number = int(match.group(2))
            column_number = _column_number(match.group(1))
            rows.setdefault(row_number, {})[column_number] = cell["value"]
    return rows


def _source_columns(
    cells_by_row: dict[int, dict[int, str]],
) -> tuple[int | None, dict[int, tuple[str, int | None]]]:
    best_row = None
    best_columns: dict[int, tuple[str, int | None]] = {}
    for row_number, cells in cells_by_row.items():
        columns = _source_header_columns(cells)
        fields = {field_name for field_name, _week in columns.values()}
        if "exercise" in fields and fields & {"day", "day_name"} and len(columns) > len(best_columns):
            best_row, best_columns = row_number, columns
    return best_row, best_columns


def _source_header_columns(cells: dict[int, str]) -> dict[int, tuple[str, int | None]]:
    return {
        column: field
        for column, cell_value in cells.items()
        if (field := _source_header_field(cell_value)) is not None
    }


def _source_header_field(header_value: str) -> tuple[str, int | None] | None:
    header = re.sub(r"\s*/\s*", "/", " ".join(header_value.casefold().split()))
    week_prefix = re.match(r"^(?:week|wk|w|الأسبوع|اسبوع)\s*(\d+)\s+(.+)$", header)
    week_number = int(week_prefix.group(1)) if week_prefix else None
    header = week_prefix.group(2) if week_prefix else header
    for header_part in (header, *header.split("/")):
        field_name = _source_field_name(header_part.strip())
        if field_name:
            return field_name, week_number
    return None


def _source_field_name(header: str) -> str | None:
    if header in _SOURCE_HEADER_FIELDS:
        return _SOURCE_HEADER_FIELDS[header]
    if "progression" in header:
        return "progression"
    if re.search(r"\b(?:load|weight)\b.*\b(?:kg|kilograms?)\b", header):
        return "load_kg"
    if "rep" in header and "rpe" in header:
        return "reps_effort"
    if "effort" in header:
        return "effort"
    return None


def _source_day_labels(
    cells_by_row: dict[int, dict[int, str]],
    header_row: int | None,
    columns: dict[int, tuple[str, int | None]],
    layout_day_numbers: set[int] | None,
    row_facts: FreeFormGridRows,
) -> tuple[dict[int, str], dict[str, int]]:
    row_labels = _day_labels_from_columns(cells_by_row, header_row, columns, row_facts)
    title_labels = _day_labels_from_titles(
        cells_by_row, columns, layout_day_numbers, row_facts
    )
    for day_number, day_name in title_labels[0].items():
        row_labels[0].setdefault(day_number, day_name)
    row_labels[1].update({key: number for key, number in title_labels[1].items() if key not in row_labels[1]})
    return row_labels


def _day_labels_from_columns(
    cells_by_row: dict[int, dict[int, str]],
    header_row: int | None,
    columns: dict[int, tuple[str, int | None]],
    row_facts: FreeFormGridRows,
) -> tuple[dict[int, str], dict[str, int]]:
    day_names: dict[int, str] = {}
    day_numbers: dict[str, int] = {}
    for row_number, cells in cells_by_row.items():
        if row_number == header_row:
            continue
        row_columns = _columns_for_source_row(row_facts, row_number, columns)
        day_column = next((column for column, spec in row_columns.items() if spec[0] == "day"), None)
        name_column = next((column for column, spec in row_columns.items() if spec[0] == "day_name"), None)
        day_number = _parse_integer(cells.get(day_column)) if day_column else None
        if day_number is None or not 1 <= day_number <= MAX_PROGRAM_DAYS:
            continue
        name = _clean_day_name(cells.get(name_column)) if name_column else ""
        if name and not _is_generic_day_name(name):
            _store_day_label(day_number, name, day_names, day_numbers)
    return day_names, day_numbers


def _day_labels_from_titles(
    cells_by_row: dict[int, dict[int, str]],
    columns: dict[int, tuple[str, int | None]],
    allowed_day_numbers: set[int] | None,
    row_facts: FreeFormGridRows,
) -> tuple[dict[int, str], dict[str, int]]:
    day_names: dict[int, str] = {}
    day_numbers: dict[str, int] = {}
    for row_number, cells in cells_by_row.items():
        for column, text in cells.items():
            if not _is_title_cell(cells_by_row, row_number, column, columns, row_facts):
                continue
            title = re.match(r"^\s*day\s*(\d+)\s*[-:—–.)]+\s*(.+?)\s*$", text, re.IGNORECASE)
            if title:
                day_number = int(title.group(1))
                if allowed_day_numbers is None or day_number in allowed_day_numbers:
                    _store_day_label(day_number, title.group(2), day_names, day_numbers)
    return day_names, day_numbers


def _add_layout_day_labels(
    layout_days: list[Any], day_names: dict[int, str], day_numbers: dict[str, int]
) -> None:
    for index, layout_day in enumerate(layout_days, start=1):
        day_number = _parse_integer(layout_day.day) or index
        name = _clean_day_name(layout_day.name)
        if 1 <= day_number <= MAX_PROGRAM_DAYS and name and not _is_generic_day_name(name):
            day_names[day_number] = name
            day_numbers[_day_name_key(name)] = day_number


def _layout_day_numbers(layout_days: list[Any]) -> set[int]:
    day_numbers = set()
    for index, layout_day in enumerate(layout_days, start=1):
        day_number = _parse_integer(layout_day.day) or index
        if 1 <= day_number <= MAX_PROGRAM_DAYS:
            day_numbers.add(day_number)
    return day_numbers


def _store_day_label(
    day_number: int,
    name: str,
    day_names: dict[int, str],
    day_numbers: dict[str, int],
) -> None:
    clean_name = _clean_day_name(name)
    if not 1 <= day_number <= MAX_PROGRAM_DAYS or not clean_name:
        return
    day_names.setdefault(day_number, clean_name)
    day_numbers.setdefault(_day_name_key(clean_name), day_number)


def _clean_day_name(value: Any) -> str:
    name = _cell_text(value)
    name = re.sub(r"^\s*day\s*\d+\s*[-:—–.)]+\s*", "", name, flags=re.IGNORECASE)
    name = re.sub(
        r"^\s*اليوم\s+(?:الأول|الأولى|الاول|الاولى|التاني|الثاني|الثانية|التانية|"
        r"التالت|الثالث|الثالثة|الرابع|الرابعة|الخامس|الخامسة|السادس|السادسة|"
        r"السابع|السابعة|الثامن|الثامنة|التاسع|التاسعة|العاشر|العاشرة|\d+)\s*[-:—–.)]+\s*",
        "", name,
    )
    return name.strip()


def _day_name_key(value: Any) -> str:
    return " ".join(_cell_text(value).split()).casefold()


def _is_generic_day_name(value: Any) -> bool:
    return bool(re.fullmatch(r"(?:training\s+)?day\s*\d+", _cell_text(value), re.IGNORECASE))


def _source_week_labels(
    cells_by_row: dict[int, dict[int, str]],
    header_row: int | None,
    columns: dict[int, tuple[str, int | None]],
    row_facts: FreeFormGridRows,
) -> tuple[list[str], list[str]]:
    title_labels = _title_week_labels(cells_by_row, columns, row_facts)
    explicit_labels = _explicit_week_labels(cells_by_row, header_row, columns)
    return title_labels, explicit_labels


def _title_week_labels(
    cells_by_row: dict[int, dict[int, str]],
    columns: dict[int, tuple[str, int | None]],
    row_facts: FreeFormGridRows,
) -> list[str]:
    labels: list[str] = []
    first_exercise_row = min(row_facts.exercise_data_rows, default=None)
    for row_number, cells in cells_by_row.items():
        for column, text in cells.items():
            if not _is_title_cell(cells_by_row, row_number, column, columns, row_facts):
                continue
            # A week-shaped note inside the exercise rows is not a sheet title.
            if (
                first_exercise_row is not None
                and row_number >= first_exercise_row
                and (
                    row_number not in row_facts.section_title_rows
                    or not _is_standalone_week_title(text)
                )
            ):
                continue
            labels.extend(_week_labels_in_cell(text))
    return list(dict.fromkeys(label for label in labels if label))


def _is_standalone_week_title(value: str) -> bool:
    return bool(re.fullmatch(
        r"\s*(?:week|wk|w|الأسبوع|اسبوع)\s*(?:\d+|[a-z])\s*",
        value,
        re.IGNORECASE,
    ))


def _explicit_week_labels(
    cells_by_row: dict[int, dict[int, str]],
    header_row: int | None,
    columns: dict[int, tuple[str, int | None]],
) -> list[str]:
    labels: list[str] = []
    for _field, week_number in columns.values():
        if week_number is not None:
            labels.append(_canonical_week_label(week_number))
    week_columns = [column for column, spec in columns.items() if spec[0] == "week"]
    for row_number, cells in cells_by_row.items():
        if row_number != header_row:
            labels.extend(_canonical_week_label(cells[column]) for column in week_columns if cells.get(column))
    return list(dict.fromkeys(label for label in labels if label))


def _implicit_week_block_columns(
    columns: dict[int, tuple[str, int | None]],
) -> dict[int, int]:
    prescription_fields = {
        "sets", "reps", "effort", "rir", "rpe", "rest_seconds", "tempo",
        "load_kg", "prescription", "reps_by_set", "reps_effort",
    }
    columns_by_field: dict[str, list[int]] = {}
    for column, (field_name, week_number) in columns.items():
        if week_number is None and field_name in prescription_fields:
            columns_by_field.setdefault(field_name, []).append(column)
    sets_count = len(columns_by_field.get("sets", []))
    reps_count = len(columns_by_field.get("reps", []))
    if sets_count < 2 or sets_count != reps_count:
        return {}
    block_count = sets_count
    return {
        column: block_number
        for field_columns in columns_by_field.values()
        if len(field_columns) == block_count
        for block_number, column in enumerate(sorted(field_columns), start=1)
    }


def _exercise_source_columns(
    columns: dict[int, tuple[str, int | None]],
) -> list[int]:
    return [column for column, spec in columns.items() if spec[0] == "exercise"]


def _is_title_cell(
    cells_by_row: dict[int, dict[int, str]],
    row_number: int,
    column: int,
    columns: dict[int, tuple[str, int | None]],
    row_facts: FreeFormGridRows,
) -> bool:
    if not cells_by_row[row_number].get(column):
        return False
    row_columns = _columns_for_source_row(row_facts, row_number, columns)
    exercise_rows = row_facts.exercise_data_rows
    section_title_rows = row_facts.section_title_rows
    data_columns = row_facts.exercise_data_columns_by_schema.get(
        _columns_schema(row_columns), frozenset()
    )
    if row_number in exercise_rows:
        return False
    following_index = bisect_right(exercise_rows, row_number)
    if following_index >= len(exercise_rows):
        return False
    if column in data_columns:
        has_day_axis = any(
            field_name in {"day", "day_name"}
            for field_name, _week in row_columns.values()
        )
        return (
            row_number < exercise_rows[0]
            or row_number in section_title_rows
            or not has_day_axis
        )
    return True


def _columns_for_source_row(
    row_facts: FreeFormGridRows,
    row_number: int,
    columns: dict[int, tuple[str, int | None]],
) -> dict[int, tuple[str, int | None]]:
    return row_facts.columns_by_row.get(row_number, columns)


def _columns_schema(
    columns: dict[int, tuple[str, int | None]],
) -> tuple[tuple[int, str, int | None], ...]:
    return tuple(
        (column, field_name, week_number)
        for column, (field_name, week_number) in sorted(columns.items())
    )


def _freeform_grid_rows(
    cells_by_row: dict[int, dict[int, str]],
    columns: dict[int, tuple[str, int | None]],
) -> FreeFormGridRows:
    active_columns = columns
    columns_by_row: dict[int, dict[int, tuple[str, int | None]]] = {}
    header_rows: set[int] = set()
    exercise_data_rows: list[int] = []
    section_title_rows: set[int] = set()
    data_columns_by_schema: dict[
        tuple[tuple[int, str, int | None], ...], set[int]
    ] = {}
    previous_nonempty_row: int | None = None
    for row_number, cells in sorted(cells_by_row.items()):
        is_header_row = _is_source_header_row(cells)
        if is_header_row:
            header_rows.add(row_number)
            if previous_nonempty_row is not None:
                previous_cells = cells_by_row[previous_nonempty_row]
                if sum(bool(text) for text in previous_cells.values()) == 1:
                    section_title_rows.add(previous_nonempty_row)
            active_columns = _source_header_columns(cells)
            columns_by_row[row_number] = active_columns
            previous_nonempty_row = row_number
            continue
        columns_by_row[row_number] = active_columns
        if active_columns and _is_exercise_data_row(cells, active_columns):
            exercise_data_rows.append(row_number)
            schema = _columns_schema(active_columns)
            data_columns = data_columns_by_schema.setdefault(schema, set())
            data_columns.update(column for column in active_columns if cells.get(column))
        if any(cells.values()):
            previous_nonempty_row = row_number
    return FreeFormGridRows(
        columns_by_row=columns_by_row,
        header_rows=frozenset(header_rows),
        exercise_data_rows=tuple(exercise_data_rows),
        exercise_data_columns_by_schema={
            schema: frozenset(data_columns)
            for schema, data_columns in data_columns_by_schema.items()
        },
        section_title_rows=frozenset(section_title_rows),
    )


def _is_exercise_data_row(
    cells: dict[int, str], columns: dict[int, tuple[str, int | None]]
) -> bool:
    exercise_source_columns = _exercise_source_columns(columns)
    prescription_source_columns = [
        column
        for column, (field_name, _week_number) in columns.items()
        if field_name in _PRESCRIPTION_SOURCE_FIELDS
    ]
    return (
        any(cells.get(column) for column in exercise_source_columns)
        and any(cells.get(column) for column in prescription_source_columns)
    )


def _grid_layout_is_ambiguous(
    columns: dict[int, tuple[str, int | None]],
) -> bool:
    unlabelled_fields = [field_name for field_name, week_number in columns.values() if week_number is None]
    return (
        unlabelled_fields.count("day") > 1
        or unlabelled_fields.count("day_name") > 1
        or unlabelled_fields.count("week") > 1
        or (unlabelled_fields.count("sets") > 1 and unlabelled_fields.count("reps") > 1)
    )


def _week_labels_in_cell(value: str) -> list[str]:
    header_or_title = re.match(
        r"^\s*(?:week|wk|الأسبوع|اسبوع)\s*([0-9]+|[a-z])(?=$|\s|[-—/:])",
        value,
        re.IGNORECASE,
    )
    if header_or_title:
        return [_canonical_week_label(f"Week {header_or_title.group(1)}")]
    column_label = re.search(
        r"\b(?:week|wk|w)\s*(\d+)\s+(?:sets?|reps?|effort|rir|rpe|load)\b",
        value,
        re.IGNORECASE,
    )
    return [_canonical_week_label(column_label.group(1))] if column_label else []


def _freeform_source_row(
    context: FreeFormGridContext, source_row: int, week: str | None
) -> FreeFormSourceRow:
    cells = context.cells_by_row.get(source_row, {})
    fields = _freeform_source_fields(context, source_row, cells, week)
    row_text = " | ".join(cells[column] for column in sorted(cells))
    # A title or note above the header can set a rule for the rows below it
    # (e.g. top sets on compound lifts); every row inherits it so an unsupported
    # construct there is flagged for the coach instead of silently dropped.
    return FreeFormSourceRow(
        fields,
        " | ".join(part for part in (row_text, context.preamble_text) if part),
    )


def _freeform_source_fields(
    context: FreeFormGridContext,
    source_row: int,
    cells: dict[int, str],
    week: str | None,
) -> dict[str, str]:
    fields: dict[str, str] = {}
    week_fields: dict[int, dict[str, str]] = {}
    row_columns = _columns_for_source_row(
        context.row_facts, source_row, context.columns
    )
    for column, text in cells.items():
        spec = row_columns.get(column)
        if spec is None:
            continue
        field_name, week_number = spec
        if week_number is None:
            implicit_week_number = context.implicit_week_columns.get(column)
            if implicit_week_number is None:
                fields[field_name] = text
            else:
                week_fields.setdefault(implicit_week_number, {})[field_name] = text
        else:
            week_fields.setdefault(week_number, {})[field_name] = text
    week_number = _week_number(week)
    if week_number is not None:
        fields.update(week_fields.get(week_number, {}))
    return fields


def _sheet_preamble_text(
    cells_by_row: dict[int, dict[int, str]], header_row: int | None
) -> str:
    if header_row is None:
        return ""
    return " | ".join(
        text
        for row_number in sorted(cells_by_row)
        if row_number < header_row
        for text in cells_by_row[row_number].values()
        if text
    )


def _freeform_reader_output(
    reply: program_import_ai.ProgramImportReply,
    tabs: list[str],
    selected_tab: str | None,
    grid: dict[str, Any],
) -> ReaderOutput:
    context = _freeform_grid_context(grid, reply.layout.days)
    weeks = _resolved_week_labels(reply.layout.weeks, context)
    selected_week = _default_import_week(weeks)
    week_lookup = {_week_identity(week): week for week in weeks}
    return ReaderOutput(
        tabs,
        selected_tab,
        False,
        _freeform_program_rows(reply.rows, week_lookup, selected_week, context),
        detected_weeks=weeks,
        selected_week=selected_week,
        weeks_not_imported=[week for week in weeks if week != selected_week],
        confirm_layout=_needs_layout_confirmation(
            context,
            reply.rows,
            reply.layout.confirm_layout,
            reply.layout.confidence,
        ),
        layout_days=[day.model_dump() for day in reply.layout.days],
        layout_confidence=reply.layout.confidence,
        import_mode="freeform",
    )


def _needs_layout_confirmation(
    context: FreeFormGridContext,
    model_rows: list[Any],
    model_confirm_layout: bool = False,
    layout_confidence: float | None = None,
) -> bool:
    return (
        context.layout_ambiguous
        or _has_unplaced_exercise_rows(context, model_rows)
        or _has_unplaceable_day_assignments(context, model_rows)
        or (
            not context.row_facts.header_rows
            and (
                model_confirm_layout
                or (
                    layout_confidence is not None
                    and layout_confidence < MIN_PROGRAM_IMPORT_LAYOUT_CONFIDENCE
                )
            )
        )
    )


def _has_unplaced_exercise_rows(
    context: FreeFormGridContext, model_rows: list[Any]
) -> bool:
    candidate_rows = set(context.row_facts.exercise_data_rows)
    reported_source_rows = {model_row.source_row for model_row in model_rows}
    return bool(candidate_rows - reported_source_rows)


def _has_unplaceable_day_assignments(
    context: FreeFormGridContext, model_rows: list[Any]
) -> bool:
    day_axis_fields = {"day", "day_name"}
    if not any(field_name in day_axis_fields for field_name, _week in context.columns.values()):
        assigned_sessions = set()
        for model_row in model_rows:
            day_number = _parse_integer(model_row.day)
            day_name = _clean_day_name(model_row.day_name or model_row.day)
            if day_number is not None or day_name:
                assigned_sessions.add((day_number, _day_name_key(day_name) if day_name else ""))
        if len(assigned_sessions) <= 1 and len(context.day_names) <= 1:
            return False
    for model_row in model_rows:
        day_number = _parse_integer(model_row.day)
        day_name = _clean_day_name(model_row.day_name or model_row.day)
        mapped_day = context.day_numbers.get(_day_name_key(day_name)) if day_name else None
        if day_number is None and mapped_day is None:
            return True
        if day_number is not None and not 1 <= day_number <= MAX_PROGRAM_DAYS:
            return True
    return False


def _is_source_header_row(cells: dict[int, str]) -> bool:
    fields = {field_name for field_name, _week in _source_header_columns(cells).values()}
    return "exercise" in fields and bool(fields & _PRESCRIPTION_SOURCE_FIELDS)


def _freeform_program_rows(
    model_rows: list[Any],
    week_lookup: dict[str, str],
    selected_week: str | None,
    context: FreeFormGridContext,
) -> list[ProgramImportRow]:
    repaired_model_rows = _repair_implicit_week_duplicates(model_rows, context)
    ordered_model_rows = _expand_implicit_week_rows(repaired_model_rows, context)
    return _renumber_out_of_range_orders([
        _program_import_row(
            model_row,
            _matched_row_week(model_row.week, week_lookup, selected_week),
            _freeform_source_row(context, model_row.source_row, model_row.week),
            context,
        )
        for model_row in ordered_model_rows
    ])


def _repair_implicit_week_duplicates(
    model_rows: list[Any], context: FreeFormGridContext
) -> list[Any]:
    duplicate_sources = _blank_week_duplicate_sources(model_rows)
    if not duplicate_sources:
        return model_rows
    if context.implicit_week_count <= 1:
        raise _duplicate_model_rows_error()
    repaired = _assign_implicit_weeks(model_rows, duplicate_sources, context.implicit_week_count)
    if _has_duplicate_model_positions(repaired):
        raise _duplicate_model_rows_error()
    return repaired


def _blank_week_duplicate_sources(model_rows: list[Any]) -> set[int]:
    positions: dict[tuple[int, str | None], list[int]] = {}
    for index, model_row in enumerate(model_rows):
        week = (
            model_row.week.strip().casefold()
            if model_row.week and model_row.week.strip()
            else None
        )
        positions.setdefault((model_row.source_row, week), []).append(index)
    duplicate_sources = set()
    for (source_row, week), indexes in positions.items():
        if len(indexes) > 1:
            if week is not None:
                raise _duplicate_model_rows_error()
            duplicate_sources.add(source_row)
    return duplicate_sources


def _assign_implicit_weeks(
    model_rows: list[Any], duplicate_sources: set[int], week_count: int
) -> list[Any]:
    implicit_weeks = _week_number_labels(week_count)
    repaired = list(model_rows)
    for source_row in duplicate_sources:
        indexes = [
            index
            for index, model_row in enumerate(model_rows)
            if model_row.source_row == source_row
            and (not model_row.week or not model_row.week.strip())
        ]
        if len(indexes) > week_count:
            raise _duplicate_model_rows_error()
        for index, week in zip(indexes, implicit_weeks):
            repaired[index] = repaired[index].model_copy(update={"week": week})
    return repaired


def _missing_implicit_week_rows(
    source_rows: list[Any], implicit_weeks: list[str]
) -> list[Any]:
    present_weeks = {
        _week_identity(row.week) for row in source_rows if row.week and row.week.strip()
    }
    template = source_rows[0]
    return [
        _implicit_week_row_without_prescription(template, week)
        for week in implicit_weeks
        if _week_identity(week) not in present_weeks
    ]


def _implicit_week_row_without_prescription(model_row: Any, week: str) -> Any:
    return model_row.model_copy(update={
        "week": week,
        "sets": None,
        "reps_min": None,
        "reps_max": None,
        "rir": None,
        "rpe": None,
    })


def _missing_rows_after_source_groups(
    model_rows: list[Any],
    rows_by_source: dict[int, list[Any]],
    implicit_weeks: list[str],
) -> dict[int, list[Any]]:
    last_index_by_source = {
        model_row.source_row: index for index, model_row in enumerate(model_rows)
    }
    implicit_identities = {_week_identity(week) for week in implicit_weeks}
    return {
        last_index_by_source[source_rows[0].source_row]: _missing_implicit_week_rows(
            source_rows, implicit_weeks
        )
        for source_rows in rows_by_source.values()
        if len(source_rows) > 1
        # Only rows already placed in implicit weeks are completed; model-named weeks stay as reported.
        and all(_week_identity(row.week) in implicit_identities for row in source_rows)
    }


def _expand_single_implicit_week_row(
    model_row: Any, implicit_weeks: list[str]
) -> list[Any]:
    expanded = [model_row.model_copy(update={"week": implicit_weeks[0]})]
    expanded.extend(
        _implicit_week_row_without_prescription(model_row, week)
        for week in implicit_weeks[1:]
    )
    return expanded


def _has_duplicate_model_positions(model_rows: list[Any]) -> bool:
    positions = [
        (
            row.source_row,
            row.week.strip().casefold() if row.week and row.week.strip() else None,
        )
        for row in model_rows
    ]
    return len(set(positions)) != len(positions)


def _duplicate_model_rows_error() -> ProgramImportError:
    return ProgramImportError(
        "program_import.interpretation_failed.v1",
        "The sheet could not be interpreted. Try the MAYOS template.",
        422,
    )


def _renumber_out_of_range_orders(rows: list[ProgramImportRow]) -> list[ProgramImportRow]:
    """Renumber a week's day from 1 when the model's orders for it fall outside 1..max (e.g. zero-based)."""
    rows_by_week_and_day: dict[tuple[str | None, Any], list[int]] = {}
    for index, row in enumerate(rows):
        rows_by_week_and_day.setdefault((row.week, row.day), []).append(index)

    renumbered = list(rows)
    for indexes in rows_by_week_and_day.values():
        orders = [_parse_integer(rows[index].order) for index in indexes]
        if all(order is None or 1 <= order <= MAX_PROGRAM_IMPORT_ROWS for order in orders):
            continue
        ordered = sorted(indexes, key=lambda index: _model_order_sort_key(rows[index], index))
        for order, index in enumerate(ordered, start=1):
            renumbered[index] = replace(rows[index], order=order)
    return renumbered


def _model_order_sort_key(row: ProgramImportRow, index: int) -> tuple[bool, int, int, int]:
    model_order = _parse_integer(row.order)
    return (
        model_order is None,
        model_order if model_order is not None else 0,
        row.source_row,
        index,
    )


def _expand_implicit_week_rows(
    model_rows: list[Any], context: FreeFormGridContext
) -> list[Any]:
    if context.implicit_week_count <= 1:
        return model_rows
    rows_by_source: dict[int, list[Any]] = {}
    for model_row in model_rows:
        rows_by_source.setdefault(model_row.source_row, []).append(model_row)
    implicit_weeks = _week_number_labels(context.implicit_week_count)
    missing_rows_after = _missing_rows_after_source_groups(
        model_rows, rows_by_source, implicit_weeks
    )
    expanded_rows = []
    for index, model_row in enumerate(model_rows):
        source_rows = rows_by_source[model_row.source_row]
        if len(source_rows) == 1:
            expanded_rows.extend(_expand_single_implicit_week_row(model_row, implicit_weeks))
        else:
            expanded_rows.append(model_row)
            expanded_rows.extend(missing_rows_after.get(index, []))
    return expanded_rows


def _week_number_labels(count: int) -> list[str]:
    return [f"Week {number}" for number in range(1, count + 1)]


def _resolved_week_labels(
    model_weeks: list[str], context: FreeFormGridContext
) -> list[str]:
    normalized = list(dict.fromkeys(
        label for week in model_weeks if (label := _canonical_week_label(week))
    ))
    model_week_ids = {_week_identity(week) for week in normalized}
    if context.implicit_week_count > 1:
        return _week_number_labels(context.implicit_week_count)
    if normalized and _weeks_mirror_day_numbers(normalized, context) and len(context.week_labels) < 2:
        normalized = []
    title_weeks = context.title_week_labels
    if model_week_ids:
        title_weeks = [week for week in title_weeks if _week_identity(week) in model_week_ids]
    week_labels = list(dict.fromkeys(normalized + context.explicit_week_labels + title_weeks))
    return week_labels if len(week_labels) > 1 else []


def _weeks_mirror_day_numbers(weeks: list[str], context: FreeFormGridContext) -> bool:
    week_numbers = {_week_number(week) for week in weeks}
    week_numbers.discard(None)
    return bool(context.day_names) and week_numbers == set(context.day_names)


def _canonical_week_label(value: Any) -> str:
    label = _cell_text(value)
    match = re.match(
        r"^\s*(?:week|wk|w|الأسبوع|اسبوع)?\s*(\d+)(?=$|\s|[-—/:])",
        label,
        re.IGNORECASE,
    )
    if match:
        return f"Week {int(match.group(1))}"
    word_week = re.match(
        r"^\s*(?:week|wk)\s*([a-z])(?=$|\s|[-—/:])",
        label,
        re.IGNORECASE,
    )
    return f"Week {word_week.group(1).upper()}" if word_week else label


def _week_identity(value: Any) -> str:
    label = _canonical_week_label(value)
    number = _week_number(label)
    if number is not None:
        return f"number:{number}"
    word_week = re.fullmatch(r"(?:week|wk)\s*([a-z])|([a-z])", label, re.IGNORECASE)
    return f"word:{(word_week.group(1) or word_week.group(2)).casefold()}" if word_week else label.casefold()


def _week_number(value: Any) -> int | None:
    match = re.fullmatch(r"(?:week|wk|w)?\s*(\d+)", _cell_text(value), re.IGNORECASE)
    return int(match.group(1)) if match else None


def _matched_row_week(
    row_week: str | None, week_lookup: dict[str, str], selected_week: str | None
) -> str | None:
    if not week_lookup:
        return None
    return week_lookup.get(_week_identity(row_week)) if row_week else selected_week


def _default_import_week(weeks: list[str]) -> str | None:
    if not weeks:
        return None
    return next(
        (week for week in weeks if re.search(r"\b(?:week|wk)\s*1\b|^1$", week, re.IGNORECASE)),
        weeks[0],
    )


def _program_import_row(
    row: Any,
    week: str | None,
    source: FreeFormSourceRow,
    context: FreeFormGridContext,
) -> ProgramImportRow:
    values = _freeform_row_values(row, source, context)
    sets, reps_min, reps_max = row.sets, row.reps_min, row.reps_max
    rir, rpe = values.rir, values.rpe
    if context.implicit_week_count > 1:
        sets, reps_min, reps_max = _fill_missing_implicit_prescription(row, source)
    return ProgramImportRow(
        source_row=row.source_row,
        day=values.day,
        day_name=values.day_name,
        order=row.order,
        exercise=row.exercise,
        sets=sets,
        reps_min=reps_min,
        reps_max=reps_max,
        reps_original=values.reps_original,
        rir=rir,
        rpe=_freeform_rpe_value(rpe),
        rest_seconds=row.rest_seconds,
        tempo=row.tempo,
        notes=row.notes,
        original_text=row.original_text,
        load_kg=row.load_kg if row.load_kg not in (None, "") else source.fields.get("load_kg"),
        load_pct_e1rm=source.fields.get("load_pct_e1rm"),
        approximation_markers=values.approximation_markers,
        week=week,
        suggestion_terms=tuple(row.exercise_search_terms),
    )


def _freeform_row_values(
    row: Any, source: FreeFormSourceRow, context: FreeFormGridContext
) -> FreeFormRowValues:
    day, day_name = _freeform_day_values(row, source, context)
    rir, rpe = _freeform_effort_values(
        row, source, include_source_rir=context.implicit_week_count > 1
    )
    return FreeFormRowValues(
        day,
        day_name,
        rir,
        rpe,
        _freeform_reps_original(row, source),
        tuple(_freeform_approximation_markers(row, source)),
    )


def _freeform_day_values(
    row: Any, source: FreeFormSourceRow, context: FreeFormGridContext
) -> tuple[Any, str | None]:
    source_day = _parse_integer(source.fields.get("day"))
    day = source_day if source_day is not None else row.day
    model_name = _clean_day_name(row.day_name)
    source_name = _clean_day_name(source.fields.get("day_name"))
    day_name = model_name if model_name and not _is_generic_day_name(model_name) else source_name or model_name
    mapped_day = context.day_numbers.get(_day_name_key(day_name)) if day_name else None
    if mapped_day is not None and (source_day is None or _is_generic_day_name(source_name)):
        day = mapped_day
    day_number = _parse_integer(day)
    if day_number is not None and day_number in context.day_names:
        day_name = context.day_names[day_number]
    elif _is_generic_day_name(day_name) or not day_name:
        day_name = None
    return day, day_name


def _freeform_effort_values(
    row: Any,
    source: FreeFormSourceRow,
    include_source_rir: bool = False,
) -> tuple[Any, Any]:
    rir, rpe = row.rir, row.rpe
    if rpe in (None, ""):
        source_rpe = _source_rpe_value(source.fields)
        if source_rpe is not None:
            _target_rir, effort_errors, _warnings, _notes = _normalise_effort(
                rir, source_rpe, row.source_row
            )
            if not effort_errors:
                rpe = source_rpe
    if include_source_rir and rir in (None, "") and rpe in (None, ""):
        source_rir = _source_rir_value(source.fields)
        if source_rir is not None:
            _target_rir, effort_errors, _warnings, _notes = _normalise_effort(
                source_rir, None, row.source_row
            )
            if not effort_errors:
                rir = source_rir
    return rir, rpe


def _fill_missing_implicit_prescription(
    row: Any,
    source: FreeFormSourceRow,
) -> tuple[Any, Any, Any]:
    sets, reps_min, reps_max = row.sets, row.reps_min, row.reps_max
    if sets in (None, "") and source.fields.get("sets"):
        normalized_sets, set_error, _warning = _normalise_sets(
            source.fields["sets"], row.source_row
        )
        if set_error is None:
            sets = normalized_sets

    raw_reps = (
        source.fields.get("reps_by_set")
        or source.fields.get("reps")
        or source.fields.get("prescription")
    )
    if raw_reps and (reps_min in (None, "") or reps_max in (None, "")):
        source_min, source_max, reps_error, _warning, _note = _normalise_reps(
            raw_reps, None, None, row.source_row
        )
        if reps_error is None:
            if reps_min in (None, ""):
                reps_min = source_min
            if reps_max in (None, ""):
                reps_max = source_max

    return sets, reps_min, reps_max


def _source_rir_value(fields: dict[str, str]) -> str | None:
    number_pattern = r"\d+(?:\.\d+)?(?:\s*-\s*\d+(?:\.\d+)?)?"
    for key in ("rir", "effort"):
        raw = fields.get(key, "")
        match = re.fullmatch(
            rf"\s*(?:RIR\s*)?({number_pattern})(?:\s*RIR)?\s*",
            raw,
            re.IGNORECASE,
        )
        if match:
            return match.group(1)
    return None


def _source_rpe_value(fields: dict[str, str]) -> str | None:
    column_value = fields.get("rpe")
    if column_value:
        if re.fullmatch(r"\d+(?:\.\d+)?", column_value.strip()):
            return f"RPE {column_value}"
        match = re.search(r"(?:RPE\s*|@\s*)(\d+(?:\.\d+)?(?:\s*-\s*\d+(?:\.\d+)?)?)", column_value, re.IGNORECASE)
        if match:
            return f"RPE {match.group(1)}"
    for key in ("effort", "reps_effort"):
        raw = fields.get(key, "")
        match = re.search(r"(?:RPE\s*|@\s*)(\d+(?:\.\d+)?(?:\s*-\s*\d+(?:\.\d+)?)?)", raw, re.IGNORECASE)
        if match:
            return f"RPE {match.group(1)}"
    return None


def _freeform_reps_original(row: Any, source: FreeFormSourceRow) -> Any:
    raw = (
        source.fields.get("reps_by_set")
        or source.fields.get("reps")
        or source.fields.get("prescription")
    )
    if raw and ("/" in raw or "," in raw) and re.search(r"\d+\s*[,/]\s*\d+", raw):
        return raw
    return row.reps_original


def _freeform_approximation_markers(
    row: Any, source: FreeFormSourceRow
) -> list[Approximation]:
    markers = [
        Approximation(_allowed_approximation_code(marker.code), marker.column, marker.original)
        for marker in row.approximation_markers
    ]
    markers.extend(_detect_unsupported_prescriptions(row, markers, source.text))
    markers.extend(_source_approximation_markers(source))
    if row.load_kg not in (None, "") and not source.fields.get("load_kg"):
        markers.append(_load_approximation(row.load_kg))
    return _dedupe_approximation_markers(markers)


def _source_approximation_markers(source: FreeFormSourceRow) -> list[Approximation]:
    markers = []
    load = source.fields.get("load_kg")
    if not _is_load_placeholder(load):
        parsed_load = _parse_load_value(load)
        if parsed_load is None or not _load_in_range(parsed_load, "kg"):
            markers.append(_load_approximation(_cell_text(load)))
    slot = source.fields.get("slot", "")
    if re.fullmatch(r"[A-Z]\d{1,2}", slot.strip(), re.IGNORECASE):
        markers.append(Approximation(
            "program_import.unsupported_preserved_as_note.v1", "prescription", f"superset {slot}"
        ))
    reps_by_set = source.fields.get("reps_by_set", "")
    reps = source.fields.get("reps") or source.fields.get("prescription", "")
    if reps_by_set:
        markers.append(Approximation("program_import.reps_approximated.v1", "reps", reps_by_set))
    elif re.search(r"\b(?:per|by)\s+set\b", reps, re.IGNORECASE):
        markers.append(Approximation("program_import.reps_approximated.v1", "reps", reps))
    elif re.search(r"\d+\s*[,/]\s*\d+.*\b(?:reps?|sets?)\b", reps, re.IGNORECASE):
        markers.append(Approximation("program_import.reps_approximated.v1", "reps", reps))
    return markers


def _load_approximation(value: Any, *, column: str = "load_kg") -> Approximation:
    original = _cell_text(value)
    return Approximation("program_import.load_preserved_as_note.v1", column, original)


def _dedupe_approximation_markers(markers: list[Approximation]) -> list[Approximation]:
    unique: dict[tuple[str, str, str], Approximation] = {}
    for marker in markers:
        unique.setdefault((marker.code, marker.column, marker.original.casefold()), marker)
    return list(unique.values())


def _freeform_rpe_value(value: Any) -> Any:
    text = _cell_text(value)
    if text and re.fullmatch(r"\d+(?:\.\d+)?", text):
        return f"RPE {text}"
    return value


def _detect_unsupported_prescriptions(
    row: Any, existing: list[Approximation], source_text: str = ""
) -> list[Approximation]:
    source = " ".join(
        part for part in (row.original_text, _cell_text(row.notes), source_text) if part
    )
    found = []
    # Dedupe on (code, text): a weekly load progression the model already
    # preserved as a load note is still an unsupported construct to flag.
    existing_unsupported = {
        marker.original.casefold()
        for marker in existing
        if marker.original and marker.code == "program_import.unsupported_preserved_as_note.v1"
    }
    for pattern, column in _UNSUPPORTED_PRESCRIPTION_PATTERNS:
        match = re.search(pattern, source, re.IGNORECASE)
        if match and match.group(0).casefold() not in existing_unsupported:
            found.append(
                Approximation(
                    "program_import.unsupported_preserved_as_note.v1",
                    column,
                    match.group(0),
                )
            )
    return found


def _allowed_approximation_code(code: str) -> str:
    allowed = {
        "program_import.reps_approximated.v1",
        "program_import.effort_approximated.v1",
        "program_import.load_preserved_as_note.v1",
        "program_import.rpe_converted.v1",
        "program_import.unsupported_preserved_as_note.v1",
    }
    return code if code in allowed else "program_import.unsupported_preserved_as_note.v1"


def _read_csv(payload: bytes) -> ReaderOutput:
    try:
        text = payload.decode("utf-8-sig")
        return _rows_from_table(csv.reader(io.StringIO(text)), ["CSV"], "CSV")
    except ProgramImportError:
        raise
    except Exception as error:
        raise ProgramImportError("program_import.file_invalid.v1", "The CSV file could not be read.") from error


def _freeform_cell_grid(
    payload: bytes, filename: str, selected_tab: str | None
) -> FreeFormCellGrid:
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if extension == "csv":
        if selected_tab:
            raise ProgramImportError("program_import.tab_invalid.v1", "CSV files have one sheet.")
        try:
            table = csv.reader(io.StringIO(payload.decode("utf-8-sig")))
            cells = _compact_cell_grid(table)
        except ProgramImportError:
            raise
        except (UnicodeDecodeError, csv.Error) as error:
            raise ProgramImportError("program_import.file_invalid.v1", "The CSV file could not be read.") from error
        if not cells:
            raise ProgramImportError("program_import.sheet_empty.v1", "The CSV sheet is empty.")
        return FreeFormCellGrid({"sheet": "CSV", "cells": cells}, ["CSV"], "CSV")
    if extension != "xlsx":
        raise ProgramImportError("program_import.file_type.v1", "Upload an XLSX or CSV file.")
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        return _xlsx_cell_grid(workbook, selected_tab)
    except ProgramImportError:
        raise
    except Exception as error:
        raise ProgramImportError("program_import.file_invalid.v1", "The XLSX file could not be read.") from error
    finally:
        if workbook is not None:
            workbook.close()


def _xlsx_cell_grid(
    workbook: Any, selected_tab: str | None
) -> FreeFormCellGrid:
    if len(workbook.worksheets) > MAX_PROGRAM_IMPORT_TABS:
        raise ProgramImportError("program_import.too_many_tabs.v1", "Trim the workbook to 20 sheets or fewer.", 413)
    tabs = [sheet.title for sheet in workbook.worksheets if _sheet_has_values(sheet)]
    if not tabs:
        raise ProgramImportError("program_import.sheet_empty.v1", "The workbook has no non-empty sheets.")
    if len(tabs) > 1 and selected_tab is None:
        return FreeFormCellGrid(None, tabs, None)
    selected = selected_tab or tabs[0]
    if selected not in tabs:
        raise ProgramImportError("program_import.tab_invalid.v1", "Choose a non-empty sheet from this workbook.")
    sheet = workbook[selected]
    rows, columns = sheet.max_row or 0, sheet.max_column or 0
    if (
        rows * columns > MAX_PROGRAM_IMPORT_GRID_RECTANGLE
        or rows > MAX_PROGRAM_IMPORT_GRID_ROWS
        or columns > MAX_PROGRAM_IMPORT_GRID_COLUMNS
    ):
        raise _cell_limit_error()
    cells = _compact_cell_grid(
        sheet.iter_rows(max_row=rows, max_col=columns, values_only=True)
    )
    if not cells:
        raise ProgramImportError("program_import.sheet_empty.v1", "The selected sheet is empty.")
    return FreeFormCellGrid({"sheet": selected, "cells": cells}, tabs, selected)


def _compact_cell_grid(table: Iterable[Iterable[Any]]) -> list[dict[str, str]]:
    cells: list[dict[str, str]] = []
    non_empty = 0
    for row_number, raw_row in enumerate(table, start=1):
        if row_number > MAX_PROGRAM_IMPORT_GRID_ROWS:
            raise _cell_limit_error()
        for column_number, raw_value in enumerate(raw_row, start=1):
            if column_number > MAX_PROGRAM_IMPORT_GRID_COLUMNS:
                raise _cell_limit_error()
            text = _cell_text(raw_value)
            if not text:
                continue
            non_empty += 1
            if non_empty > MAX_PROGRAM_IMPORT_CELLS:
                raise _cell_limit_error()
            cells.append({"address": f"{_column_label(column_number)}{row_number}", "value": text})
    return cells


def _column_label(column_number: int) -> str:
    letters = ""
    while column_number:
        column_number, remainder = divmod(column_number - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _column_number(label: str) -> int:
    number = 0
    for letter in label:
        number = number * 26 + ord(letter) - ord("A") + 1
    return number


def _cell_limit_error() -> ProgramImportError:
    return ProgramImportError(
        "program_import.too_many_cells.v1",
        f"Trim the selected sheet to {MAX_PROGRAM_IMPORT_CELLS} non-empty cells or fewer, then upload it again.",
        413,
        {"max_cells": MAX_PROGRAM_IMPORT_CELLS},
    )


def _read_xlsx(payload: bytes, selected_tab: str | None) -> ReaderOutput:
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        return _read_selected_xlsx(workbook, selected_tab)
    except ProgramImportError:
        raise
    except Exception as error:
        raise ProgramImportError("program_import.file_invalid.v1", "The XLSX file could not be read.") from error
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass


def _read_selected_xlsx(workbook: Any, selected_tab: str | None) -> ReaderOutput:
    if len(workbook.worksheets) > MAX_PROGRAM_IMPORT_TABS:
        raise ProgramImportError(
            "program_import.too_many_tabs.v1",
            f"Trim the workbook to {MAX_PROGRAM_IMPORT_TABS} sheets or fewer.",
            413,
        )
    tabs = [sheet.title for sheet in workbook.worksheets if _sheet_has_values(sheet)]
    if not tabs:
        raise ProgramImportError("program_import.sheet_empty.v1", "The workbook has no non-empty sheets.")
    choice_tabs = [tab for tab in tabs if tab.casefold() != "example"]
    if len(choice_tabs) > 1 and selected_tab is None:
        return ReaderOutput(tabs, None, True)
    if selected_tab is None:
        selected = "Program" if "Program" in tabs else (choice_tabs or tabs)[0]
    else:
        selected = selected_tab
    if selected not in tabs:
        raise ProgramImportError("program_import.tab_invalid.v1", "Choose a non-empty sheet from this workbook.")
    sheet = workbook[selected]
    if (sheet.max_row or 0) > MAX_PROGRAM_IMPORT_ROWS + 2:
        raise ProgramImportError("program_import.too_many_rows.v1", "Trim the sheet to 500 exercise rows or fewer.", 413)
    if (sheet.max_column or 0) > MAX_PROGRAM_IMPORT_COLUMNS:
        raise ProgramImportError(
            "program_import.too_many_columns.v1",
            f"Trim the sheet to {MAX_PROGRAM_IMPORT_COLUMNS} columns or fewer.",
            413,
        )
    return _rows_from_table(
        sheet.iter_rows(max_row=MAX_PROGRAM_IMPORT_ROWS + 2, max_col=MAX_PROGRAM_IMPORT_COLUMNS, values_only=True),
        tabs,
        selected,
    )


def _sheet_has_values(sheet: Any) -> bool:
    for row in sheet.iter_rows(
        max_row=MAX_PROGRAM_IMPORT_TAB_PROBE_ROWS,
        max_col=MAX_PROGRAM_IMPORT_COLUMNS,
        values_only=True,
    ):
        if any(value is not None and str(value).strip() for value in row):
            return True
    # Dimensions are metadata. Treat content beyond the bounded probe as possibly non-empty.
    return (sheet.max_row or 0) > MAX_PROGRAM_IMPORT_TAB_PROBE_ROWS or (sheet.max_column or 0) > MAX_PROGRAM_IMPORT_COLUMNS


def _rows_from_table(
    table: Iterable[Iterable[Any]], detected_tabs: list[str], selected_tab: str | None
) -> ReaderOutput:
    columns: dict[str, int] | None = None
    rows: list[ProgramImportRow] = []
    required = {"day", "exercise", "sets"}
    for row_number, values in enumerate(table, start=1):
        if row_number > MAX_PROGRAM_IMPORT_ROWS + 2:
            raise ProgramImportError("program_import.too_many_rows.v1", "Trim the sheet to 500 exercise rows or fewer.", 413)
        cells = list(values)
        if len(cells) > MAX_PROGRAM_IMPORT_COLUMNS:
            raise ProgramImportError(
                "program_import.too_many_columns.v1",
                f"Trim the sheet to {MAX_PROGRAM_IMPORT_COLUMNS} columns or fewer.",
                413,
            )
        if columns is None:
            headers = [_cell_text(cell).casefold() for cell in cells]
            if not any(headers):
                continue
            columns = {name: index for index, name in enumerate(headers) if name in TEMPLATE_COLUMNS}
            missing = required - columns.keys()
            if "reps" not in columns:
                missing.add("reps")
            if missing:
                raise ProgramImportError(
                    "program_import.template_columns.v1",
                    "Required template columns are missing: " + ", ".join(sorted(missing)),
                )
            continue
        named = {
            name: cells[index] if index < len(cells) else None
            for name, index in columns.items()
        }
        if not any(value is not None and str(value).strip() for value in named.values()):
            continue
        rows.append(_template_row(row_number, named, has_week_column="week" in columns))
        if len(rows) > MAX_PROGRAM_IMPORT_ROWS:
            break
    if columns is None:
        raise ProgramImportError("program_import.sheet_empty.v1", "The selected sheet is empty.")
    detected_weeks = _detected_template_weeks(rows) if "week" in columns else []
    return ReaderOutput(
        detected_tabs,
        selected_tab,
        False,
        rows,
        detected_weeks=detected_weeks,
    )


def _detected_template_weeks(rows: list[ProgramImportRow]) -> list[str]:
    week_orders = {
        week_order
        for row in rows
        if (week_order := _parse_integer(row.week_order)) is not None
        and 1 <= week_order <= MAX_PROGRAM_WEEKS
    }
    return [f"Week {week_order}" for week_order in sorted(week_orders)]


def _template_row(
    source_row: int, cells: dict[str, Any], *, has_week_column: bool
) -> ProgramImportRow:
    reps = cells.get("reps")
    reps_min, reps_max, rep_markers = _reader_reps(reps)
    notes = _cell_text(cells.get("notes"))
    markers = list(rep_markers)
    load_kg = cells.get("load_kg")
    load_pct_e1rm = cells.get("load_pct_e1rm")
    raw_values = {key: _cell_text(value) for key, value in cells.items() if value is not None}
    week_order = cells.get("week")
    if has_week_column and not _cell_text(week_order):
        week_order = 1
    return ProgramImportRow(
        source_row=source_row,
        day=cells.get("day"),
        day_name=cells.get("day_name"),
        order=cells.get("order"),
        exercise=cells.get("exercise"),
        sets=cells.get("sets"),
        reps_min=reps_min,
        reps_max=reps_max,
        reps_original=reps,
        rir=cells.get("rir"),
        rpe=cells.get("rpe"),
        rest_seconds=cells.get("rest_sec"),
        tempo=cells.get("tempo"),
        notes=notes,
        original_text=" | ".join(f"{key}: {value}" for key, value in raw_values.items()),
        load_kg=load_kg,
        load_pct_e1rm=load_pct_e1rm,
        approximation_markers=tuple(markers),
        week_order=week_order if has_week_column else None,
    )


def _reader_reps(value: Any) -> tuple[ImportValue, ImportValue, tuple[Approximation, ...]]:
    if isinstance(value, (dt.date, dt.datetime)):
        return value, None, ()
    if isinstance(value, bool):
        return value, None, ()
    if isinstance(value, int) or (isinstance(value, float) and math.isfinite(value) and value.is_integer()):
        number = int(value)
        return number, number, ()
    text = _cell_text(value)
    if not text:
        return None, None, ()
    match = _REP_RANGE.fullmatch(text)
    if match:
        return int(match.group(1)), int(match.group(2)), ()
    if re.fullmatch(r"\d+", text):
        value = int(text)
        return value, value, ()
    list_match = _PER_SET_REP_LIST.fullmatch(text)
    if list_match:
        reps = [int(part) for part in re.split(r"[,/]", list_match.group(1))]
        return min(reps), max(reps), (Approximation("program_import.reps_approximated.v1", "reps", text),)
    return text, None, ()


def validate_and_resolve_rows(
    db: Any,
    coach_account_id: str,
    source_rows: list[ProgramImportRow],
    *,
    text_import: bool = False,
) -> list[dict[str, Any]]:
    """Shared normalization, row validation, row-level checks, and exercise resolution."""
    if len(source_rows) > MAX_PROGRAM_IMPORT_ROWS:
        raise ProgramImportError("program_import.too_many_rows.v1", "Trim the sheet to 500 exercise rows or fewer.", 413)
    output: list[dict[str, Any]] = []
    order_by_day: dict[tuple[str, str | int], int] = {}
    for source in source_rows:
        row = _normalise_row(source, order_by_day, text_import=text_import)
        resolution = _resolve_exercise(db, coach_account_id, row["exercise_name"], source.confirmed_exercise_id)
        resolution_errors = resolution.pop("errors", [])
        resolution_warnings = resolution.pop("warnings", [])
        for item in resolution_errors + resolution_warnings:
            item["source_row"] = source.source_row
        row["errors"].extend(resolution_errors)
        row["warnings"].extend(resolution_warnings)
        row.update(resolution)
        row["approximation_markers"] = [warning["code"] for warning in row["warnings"]]
        row["valid"] = not row["errors"]
        output.append(row)
    _normalise_day_names(output)
    _mark_per_day_count_errors(output)
    _mark_conflicting_day_names(output)
    return output


def _normalise_row(
    source: ProgramImportRow,
    order_by_day: dict[tuple[str, str | int], int],
    *,
    text_import: bool = False,
) -> dict[str, Any]:
    errors: list[dict[str, Any]] = []
    warnings = [
        _row_warning(
            item.code,
            source.source_row,
            item.column,
            item.original if text_import and item.column == "load_kg" else "",
        )
        for item in source.approximation_markers
    ]
    week, week_order, week_error = _normalise_template_week(
        source.week_order, source.week, source.source_row
    )
    if week_error:
        errors.append(week_error)
    day = _parse_integer(source.day)
    day_valid = day is not None and 1 <= day <= MAX_PROGRAM_DAYS
    if not day_valid:
        errors.append(_row_error("program_import.invalid_day.v1", source.source_row, "day", "Day is outside the supported range."))
        if day is None or not text_import:
            day = 1
    if day_valid:
        day_key = (_week_group_key(week, week_order), day)
        order_by_day[day_key] = order_by_day.get(day_key, 0) + 1
    order = _parse_integer(source.order)
    if order is None:
        order = order_by_day.get((_week_group_key(week, week_order), day), 1)
    elif not 1 <= order <= MAX_PROGRAM_IMPORT_ROWS:
        errors.append(_row_error("program_import.invalid_order.v1", source.source_row, "order", "Exercise order is outside the supported range."))
        order = order_by_day.get((_week_group_key(week, week_order), day), 1)

    sets, set_error, set_warning = _normalise_sets(source.sets, source.source_row)
    if set_error:
        errors.append(set_error)
    if set_warning:
        warnings.append(set_warning)

    reps_min, reps_max, reps_error, reps_warning, reps_note = _normalise_reps(
        source.reps_min, source.reps_max, source.reps_original, source.source_row
    )
    if reps_error:
        errors.append(reps_error)
    if reps_warning:
        warnings.append(reps_warning)
    marker_notes = [f"reps: {item.original}" for item in source.approximation_markers if item.column == "reps" and item.original]
    if marker_notes:
        reps_note = "\n".join(marker_notes)

    target_rir, effort_errors, effort_warnings, effort_notes = _normalise_effort(
        source.rir, source.rpe, source.source_row
    )
    errors.extend(effort_errors)
    warnings.extend(effort_warnings)

    rest_seconds = _parse_integer(source.rest_seconds)
    if source.rest_seconds not in (None, "") and (
        rest_seconds is None or not 0 <= rest_seconds <= MAX_PROGRAM_IMPORT_REST_SECONDS
    ):
        errors.append(_row_error("program_import.invalid_rest.v1", source.source_row, "rest_seconds", "Rest is outside the supported range."))
        rest_seconds = DEFAULT_EXERCISE_REST_SECONDS
    elif rest_seconds is None:
        rest_seconds = DEFAULT_EXERCISE_REST_SECONDS

    exercise_name = _cell_text(source.exercise)
    if not exercise_name:
        errors.append(_row_error("program_import.exercise_required.v1", source.source_row, "exercise", "Enter an exercise name."))
    elif len(exercise_name) > MAX_PROGRAM_IMPORT_EXERCISE_NAME_LENGTH:
        errors.append(_row_error("program_import.exercise_name_too_long.v1", source.source_row, "exercise", "Exercise name is too long."))

    day_name = _cell_text(source.day_name) or None
    if day_name and len(day_name) > MAX_PROGRAM_IMPORT_DAY_NAME_LENGTH:
        errors.append(_row_error("program_import.day_name_too_long.v1", source.source_row, "day_name", "Day name is too long."))

    tempo = _cell_text(source.tempo) or None
    if tempo and len(tempo) > MAX_TEMPO_LENGTH:
        errors.append(_row_error("program_import.tempo_too_long.v1", source.source_row, "tempo", "Tempo is too long."))

    parsed_loads: dict[str, float] = {}
    load_markers: list[Approximation] = []
    for column, raw_load, kind in (
        ("load_kg", source.load_kg, "kg"),
        ("load_pct_e1rm", source.load_pct_e1rm, "pct_e1rm"),
    ):
        if _is_load_placeholder(raw_load):
            continue
        parsed_load = _parse_load_value(
            raw_load, allow_percent=kind == "pct_e1rm"
        )
        if parsed_load is None:
            marker = _load_approximation(_cell_text(raw_load), column=column)
            load_markers.append(marker)
            if not any(
                item.code == marker.code
                and item.column == marker.column
                and item.original == marker.original
                for item in source.approximation_markers
            ):
                warnings.append(
                    _row_warning(marker.code, source.source_row, column, marker.original)
                )
        elif not _load_in_range(parsed_load, kind):
            errors.append(
                _row_error(
                    "program_import.invalid_load.v1",
                    source.source_row,
                    column,
                    (
                        f"Load must be greater than {MIN_PROGRAM_LOAD_KG:g} and at most "
                        f"{MAX_PROGRAM_LOAD_KG:g} kg."
                        if kind == "kg"
                        else f"e1RM load must be from {MIN_PROGRAM_LOAD_PCT_E1RM:g} "
                        f"to {MAX_PROGRAM_LOAD_PCT_E1RM:g} percent."
                    ),
                )
            )
        else:
            parsed_loads[kind] = parsed_load

    if len(parsed_loads) > 1:
        errors.append(
            _row_error(
                "program_import.multiple_loads.v1",
                source.source_row,
                "load",
                "Choose either kg or percent of e1RM.",
            )
        )

    preserved_notes = list(dict.fromkeys(
        f"{item.column}: {item.original}"
        for item in (*source.approximation_markers, *load_markers)
        if item.original and item.column not in {"reps", "rir", "rpe"}
        and item.original.casefold() not in (source.notes or "").casefold()
    ))
    notes = _join_notes(source.notes, reps_note, *effort_notes, *preserved_notes)
    if len(notes) > MAX_PROGRAM_NOTES_LENGTH:
        errors.append(_row_error("program_import.notes_too_long.v1", source.source_row, "notes", "Notes are too long."))

    load = source.load or (
        {"kind": "kg", "value": parsed_loads["kg"]} if "kg" in parsed_loads else
        {"kind": "pct_e1rm", "value": parsed_loads["pct_e1rm"]}
        if "pct_e1rm" in parsed_loads else None
    )

    normalized = {
        "source_row": source.source_row,
        "week": week,
        "week_order": week_order,
        "original_text": source.original_text,
        "approximation_markers": list(dict.fromkeys(
            item.code for item in (*source.approximation_markers, *load_markers)
        )),
        "suggestion_search_terms": list(source.suggestion_terms),
        "day": day,
        "day_valid": day_valid,
        "day_name": day_name,
        "order": order,
        "exercise_name": exercise_name,
        "exercise_id": source.confirmed_exercise_id,
        "resolution": "unresolved",
        "suggestions": [],
        "sets": sets,
        "reps_min": reps_min,
        "reps_max": reps_max,
        "target_rir": target_rir,
        "rest_seconds": rest_seconds,
        "tempo": tempo,
        "notes": notes or None,
        "load": load,
        "errors": errors,
        "warnings": warnings,
    }
    if text_import:
        normalized["input_exercise_name"] = source.input_exercise_name
    return normalized


def _normalise_template_week(
    raw_week_order: ImportValue, source_week: str | None, source_row: int
) -> tuple[str | None, int | None, dict[str, Any] | None]:
    if raw_week_order is None:
        return source_week, None, None
    week_order = _parse_integer(raw_week_order)
    if week_order is None or not 1 <= week_order <= MAX_PROGRAM_WEEKS:
        error = _row_error(
            "program_import.week_invalid.v1",
            source_row,
            "week",
            f"Week must be a number from 1 to {MAX_PROGRAM_WEEKS}.",
        )
        return "Week 1", 1, error
    return f"Week {week_order}", week_order, None


def _is_load_placeholder(value: Any) -> bool:
    return _cell_text(value).casefold() in _LOAD_PLACEHOLDERS


def _load_in_range(value: float, kind: str) -> bool:
    if kind == "kg":
        return MIN_PROGRAM_LOAD_KG < value <= MAX_PROGRAM_LOAD_KG
    return MIN_PROGRAM_LOAD_PCT_E1RM <= value <= MAX_PROGRAM_LOAD_PCT_E1RM


def _parse_load_value(value: Any, *, allow_percent: bool = False) -> float | None:
    if _is_load_placeholder(value):
        return None
    try:
        text = _cell_text(value).casefold().strip()
        if "%" in text:
            if not allow_percent or not re.fullmatch(r"\d+(?:\.\d+)?\s*%", text):
                return None
            text = text.removesuffix("%").strip()
        elif allow_percent and re.search(r"\bkg\b|\blb\b", text):
            return None
        else:
            if re.search(r"\blb\b", text):
                return None
            text = re.sub(r"\s*kg\s*$", "", text)
        parsed = float(text)
    except (TypeError, ValueError, OverflowError):
        return None
    return parsed if math.isfinite(parsed) else None


def _week_group_key(week: str | None, week_order: int | None) -> tuple[str, str | int]:
    if week_order is not None:
        return ("order", week_order)
    if week:
        return ("label", _week_identity(week))
    return ("order", 1)


def _normalise_sets(value: Any, source_row: int) -> tuple[int, dict[str, Any] | None, dict[str, Any] | None]:
    sets = _parse_integer(value)
    if sets is not None and 1 <= sets <= MAX_PROGRAM_IMPORT_SETS:
        return sets, None, None
    error = _row_error("program_import.invalid_sets.v1", source_row, "sets", "Working sets are outside the supported range.")
    return DEFAULT_TARGET_SETS, error, None


def _normalise_reps(
    lower_value: Any, upper_value: Any, original_value: Any, source_row: int
) -> tuple[int, int, dict[str, Any] | None, dict[str, Any] | None, str]:
    if isinstance(lower_value, (dt.date, dt.datetime)):
        error = _row_error("program_import.reps_as_date.v1", source_row, "reps", "Format reps as text and type the prescription again.")
        return DEFAULT_TARGET_REPS_MIN, DEFAULT_TARGET_REPS_MAX, error, None, ""
    lower = _parse_integer(lower_value)
    upper = _parse_integer(upper_value)
    raw = _cell_text(lower_value)
    original = _cell_text(original_value)
    if upper is None and raw:
        match = _REP_RANGE.fullmatch(raw)
        if match:
            lower, upper = int(match.group(1)), int(match.group(2))
        elif list_match := _PER_SET_REP_LIST.fullmatch(raw):
            values = [int(part) for part in re.split(r"[,/]", list_match.group(1))]
            lower, upper = min(values), max(values)
            warning = _row_warning("program_import.reps_approximated.v1", source_row, "reps", "Per-set repetitions were converted to a range.")
            return _check_rep_bounds(lower, upper, source_row, warning, raw)
        elif re.fullmatch(r"\d+(?:\.0+)?", raw):
            lower = upper = int(float(raw))
        elif _is_unsupported_rep_prescription(raw or original):
            raw = raw or original
            warning = _row_warning("program_import.reps_approximated.v1", source_row, "reps", "Unsupported repetitions were approximated.")
            return DEFAULT_TARGET_REPS_MIN, DEFAULT_TARGET_REPS_MAX, None, warning, f"reps: {raw}"
        else:
            error = _row_error("program_import.invalid_reps.v1", source_row, "reps", "Enter a number, range, or supported prescription.")
            return DEFAULT_TARGET_REPS_MIN, DEFAULT_TARGET_REPS_MAX, error, None, ""
    if lower is None or upper is None:
        error = _row_error("program_import.reps_required.v1", source_row, "reps", "Enter repetitions.")
        return DEFAULT_TARGET_REPS_MIN, DEFAULT_TARGET_REPS_MAX, error, None, ""
    return _check_rep_bounds(lower, upper, source_row, None, "")


def _check_rep_bounds(
    lower: int,
    upper: int,
    source_row: int,
    warning: dict[str, Any] | None,
    original: str,
) -> tuple[int, int, dict[str, Any] | None, dict[str, Any] | None, str]:
    if lower > upper or lower < MIN_REPS or upper > MAX_REPS:
        error = _row_error("program_import.invalid_reps.v1", source_row, "reps", "Repetitions are outside the supported range.")
        return DEFAULT_TARGET_REPS_MIN, DEFAULT_TARGET_REPS_MAX, error, None, ""
    return lower, upper, None, warning, f"reps: {original}" if warning and original else ""


def _is_unsupported_rep_prescription(value: str) -> bool:
    clean = value.casefold()
    return any(token in clean for token in ("amrap", "failure", "%", "top set", "back-off", "back off", "superset", "+", "x"))


def _normalise_effort(
    rir_value: Any, rpe_value: Any, source_row: int
) -> tuple[float, list[dict[str, Any]], list[dict[str, Any]], list[str]]:
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    notes: list[str] = []
    raw_rir, raw_rpe = _cell_text(rir_value), _cell_text(rpe_value)
    if (raw_rir.startswith("@") or raw_rir.casefold().startswith("rpe")) and not raw_rpe:
        raw_rpe, raw_rir = raw_rir, ""
    rir, rir_approx, rir_garbage = _parse_effort_value(raw_rir, is_rpe=False)
    rpe, rpe_approx, rpe_garbage = _parse_effort_value(raw_rpe, is_rpe=True)
    if rir_garbage or rpe_garbage:
        column = "rir" if rir_garbage else "rpe"
        errors.append(_row_error("program_import.invalid_effort.v1", source_row, column, "Enter a numeric RIR or RPE value."))
    converted = None if rpe is None else 10.0 - rpe
    if rir is not None and converted is not None and abs(rir - converted) > 0.001:
        errors.append(_row_error("program_import.conflicting_effort.v1", source_row, "rpe", "RIR and RPE values do not match."))
    if converted is not None:
        warnings.append(_row_warning("program_import.rpe_converted.v1", source_row, "rpe", "RPE was converted to RIR."))
        notes.append(f"RPE: {raw_rpe}")
    if rir_approx or rpe_approx:
        original = raw_rir or raw_rpe
        warnings.append(_row_warning("program_import.effort_approximated.v1", source_row, "rir" if raw_rir else "rpe", "Effort was approximated using the midpoint, clamped to the supported range."))
        notes.append(f"effort: {original}")
    selected = rir if rir is not None else converted
    if selected is None:
        return DEFAULT_TARGET_RIR, errors, warnings, notes
    return min(MAX_TARGET_RIR, max(MIN_TARGET_RIR, selected)), errors, warnings, notes


def _parse_effort_value(raw: str, *, is_rpe: bool) -> tuple[float | None, bool, bool]:
    if not raw:
        return None, False, False
    value_match = _RPE_VALUE.fullmatch(raw) if is_rpe else _EFFORT_VALUE.fullmatch(raw)
    range_match = _RPE_RANGE.fullmatch(raw) if is_rpe else _EFFORT_RANGE.fullmatch(raw)
    approximate = False
    if value_match:
        value = float(value_match.group(1))
        approximate = raw != value_match.group(1)
    elif range_match:
        value = (float(range_match.group(1)) + float(range_match.group(2))) / 2
        approximate = True
    else:
        return None, False, True
    low, high = (5.0, 10.0) if is_rpe else (MIN_TARGET_RIR, MAX_TARGET_RIR)
    if not low <= value <= high:
        value = min(high, max(low, value))
        approximate = True
    return value, approximate, False


def _join_notes(*values: Any) -> str:
    return "\n".join(text for value in values if (text := _cell_text(value)))


def _normalise_day_names(rows: list[dict[str, Any]]) -> None:
    names: dict[tuple[tuple[str, str | int], int], str] = {}
    for row in rows:
        if row["day_valid"] and row["day_name"]:
            names.setdefault(
                (_week_group_key(row["week"], row["week_order"]), row["day"]),
                row["day_name"],
            )
    for row in rows:
        if not row["day_name"]:
            day_key = (_week_group_key(row["week"], row["week_order"]), row["day"])
            row["day_name"] = names.get(day_key, f"Day {row['day']}")


def _mark_per_day_count_errors(rows: list[dict[str, Any]]) -> None:
    counts: dict[tuple[tuple[str, str | int], int], list[dict[str, Any]]] = {}
    for row in rows:
        if row["valid"] and row["day_valid"] and row["exercise_name"]:
            day_key = (_week_group_key(row["week"], row["week_order"]), row["day"])
            counts.setdefault(day_key, []).append(row)
    for day_rows in counts.values():
        for row in day_rows[MAX_EXERCISES_PER_DAY:]:
            row["errors"].append(_row_error("program_import.too_many_exercises.v1", row["source_row"], "exercise", "A day has too many exercises."))
            row["valid"] = False


def _mark_conflicting_day_names(rows: list[dict[str, Any]]) -> None:
    names: dict[tuple[tuple[str, str | int], int], str] = {}
    for row in rows:
        if not row["day_valid"] or not row["day_name"]:
            continue
        day_key = (_week_group_key(row["week"], row["week_order"]), row["day"])
        current = names.setdefault(day_key, row["day_name"])
        if row["day_name"] != current:
            row["errors"].append(_row_error("program_import.conflicting_day_name.v1", row["source_row"], "day_name", "Use one name for each day number."))
            row["valid"] = False


def _resolve_exercise(db: Any, coach_account_id: str, name: str, confirmed_id: str | None) -> dict[str, Any]:
    if confirmed_id:
        entry = db.get_exercise_library_entry(confirmed_id)
        if entry is None:
            entry = db.get_coach_exercise(coach_account_id, confirmed_id)
        if entry is None:
            return {"exercise_id": None, "resolution": "invalid", "suggestions": [],
                    "errors": [_row_error("program_import.exercise_invalid.v1", 0, "exercise", "Choose an exercise owned by this coach or in the library.")]}
        return {"exercise_id": confirmed_id, "exercise_name": entry["name"], "resolution": "confirmed", "suggestions": [], "errors": []}
    exact_id = db.find_unique_exercise_id_by_exact_name(name) if name else None
    if exact_id:
        entry = db.get_exercise_library_entry(exact_id)
        return {"exercise_id": exact_id, "exercise_name": entry["name"], "resolution": "resolved", "suggestions": [], "errors": []}
    coach_matches = db.search_coach_exercises(coach_account_id, name, limit=None) if name else []
    own_exact = [entry for entry in coach_matches if entry["name"].casefold() == name.casefold()]
    if len(own_exact) == 1:
        entry = own_exact[0]
        return {"exercise_id": entry["id"], "exercise_name": entry["name"], "resolution": "resolved", "suggestions": [], "errors": []}
    library_matches = db.find_exercises_by_name(name, limit=10) if name else []
    suggestions = _unique_suggestions([*library_matches, *own_exact])[:3]
    ambiguous = len(library_matches) > 1 or len(own_exact) > 1
    code = "program_import.exercise_ambiguous.v1" if ambiguous else "program_import.exercise_unresolved.v1"
    column_message = "Several exercises match; choose one." if ambiguous else "No exact exercise match; choose a suggestion or create your exercise."
    warning = _row_warning(code, 0, "exercise", column_message)
    return {"exercise_id": None, "resolution": "ambiguous" if ambiguous else "unresolved", "suggestions": suggestions, "errors": [], "warnings": [warning]}


def _unique_suggestions(matches: list[dict[str, Any]]) -> list[dict[str, str]]:
    seen: set[str] = set()
    suggestions = []
    for match in matches:
        exercise_id = str(match["id"])
        if exercise_id not in seen:
            seen.add(exercise_id)
            suggestions.append({"exercise_id": exercise_id, "name": str(match["name"])})
    return suggestions


def _unresolved_names(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    unresolved: dict[str, dict[str, Any]] = {}
    for row in rows:
        if row["resolution"] in {"resolved", "confirmed"}:
            continue
        name = row["exercise_name"]
        item = unresolved.setdefault(name, {"name": name, "resolution": row["resolution"], "rows": [], "suggestions": row["suggestions"]})
        item["rows"].append(row["source_row"])
    return list(unresolved.values())


def _confirmed_row(row: dict[str, Any]) -> ProgramImportRow:
    return ProgramImportRow(
        source_row=row["source_row"],
        day=row["day"],
        day_name=row.get("day_name"),
        order=row.get("order"),
        exercise=row["exercise_name"],
        sets=row["sets"],
        reps_min=row["reps_min"],
        reps_max=row["reps_max"],
        reps_original=f"{row['reps_min']}-{row['reps_max']}",
        rir=row["target_rir"],
        rpe=None,
        rest_seconds=row["rest_seconds"],
        tempo=row.get("tempo"),
        notes=row.get("notes"),
        load=row.get("load"),
        original_text="",
        confirmed_exercise_id=row["exercise_id"],
        week_order=row.get("week_order", 1),
    )


def _confirmed_rows_by_week(
    db: Any, coach_account_id: str, rows: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    source_orders = sorted(
        {
            row.get("week_order", 1)
            for row in rows
            if row["valid"] and row["exercise_id"]
        }
    )
    if not source_orders:
        raise ProgramImportError(
            "program_import.no_rows.v1",
            "No valid confirmed rows are ready to create a draft.",
        )
    return [
        _confirmed_week_for_import(
            db, coach_account_id, rows, (week_order, source_order)
        )
        for week_order, source_order in enumerate(source_orders, start=1)
    ]


def _confirmed_week_for_import(
    db: Any,
    coach_account_id: str,
    rows: list[dict[str, Any]],
    orders: tuple[int, int],
) -> dict[str, Any]:
    week_order, source_order = orders
    week_rows = [row for row in rows if row.get("week_order", 1) == source_order]
    return {
        "week_name": f"Week {week_order}",
        "week_order": week_order,
        "days": _confirmed_rows_by_day(db, coach_account_id, week_rows, week_order),
    }


def _confirmed_rows_by_day(
    db: Any,
    coach_account_id: str,
    rows: list[dict[str, Any]],
    week_order: int,
) -> list[dict[str, Any]]:
    days: dict[int, dict[str, Any]] = {}
    for row in sorted(rows, key=lambda candidate: (candidate["day"], candidate["order"], candidate["source_row"])):
        if not row["valid"] or not row["exercise_id"]:
            continue
        if row["errors"]:
            continue
        exercise_id = row["exercise_id"]
        entry = db.get_exercise_library_entry(exercise_id)
        is_coach = False
        if entry is None:
            entry = db.get_coach_exercise(coach_account_id, exercise_id)
            is_coach = entry is not None
        if entry is None:
            continue
        day_number = row["day"]
        day = days.setdefault(day_number, {
            "day_name": row["day_name"], "day_order": day_number,
            "week_order": week_order,
            "warmup_exercises": [], "exercises": [], "cardio": None,
        })
        day["exercises"].append({
            "exercise_id": exercise_id,
            "exercise_name": entry["name"],
            "body_part": entry.get("body_part") if is_coach else None,
            "equipment": entry.get("equipment"),
            "note": entry.get("note") if is_coach else None,
            "video_url": entry.get("video_url") if is_coach else None,
            "is_coach_exercise": is_coach,
            "target_sets": row["sets"],
            "target_reps_min": row["reps_min"],
            "target_reps_max": row["reps_max"],
            "target_rir": row["target_rir"],
            "rest_seconds": row["rest_seconds"],
            "tempo": row["tempo"],
            "notes": row["notes"],
            "load": row.get("load"),
        })
    if not days:
        raise ProgramImportError("program_import.no_rows.v1", "No valid confirmed rows are ready to create a draft.")
    return list(days.values())


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _parse_integer(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return int(value)
    try:
        text = _cell_text(value)
        return int(text) if text else None
    except (ValueError, TypeError, OverflowError):
        return None


def _row_error(code: str, source_row: int, column: str, message: str) -> dict[str, Any]:
    return {"source_row": source_row, "column": column, "code": code, "message": message}


def _row_warning(code: str, source_row: int, column: str, message: str) -> dict[str, Any]:
    return {"source_row": source_row, "column": column, "code": code, "message": message}
