"""Program template import contract tests (#298)."""

import io
import csv
import datetime as dt
import json
import logging
import re
import sqlite3
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from service import coach as coach_service
from svc.app import create_app
from svc.dependencies import get_db

TEST_JWT_SECRET = "test-secret-key-0123456789abcdef"
HEADERS = [
    "day", "day_name", "order", "exercise", "sets", "reps", "rir", "rpe",
    "load_kg", "load_pct_e1rm", "rest_sec", "tempo", "notes",
]
CONFIRMED_ROW_FIELDS = (
    "source_row", "day", "day_name", "order", "exercise_name", "exercise_id",
    "sets", "reps_min", "reps_max", "target_rir", "rest_seconds", "tempo", "notes", "load",
            "group",
)


@pytest.fixture
def api(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("SKIP_LLM_LOAD", "true")
    monkeypatch.setenv("TESTING", "1")
    monkeypatch.setenv("JWT_SECRET", TEST_JWT_SECRET)
    # A developer .env may enable free-form import; tests opt in explicitly.
    monkeypatch.delenv("PROGRAM_IMPORT_AI_ENABLED", raising=False)
    monkeypatch.delenv("PROGRAM_IMPORT_AI_EVAL_REPORT", raising=False)
    monkeypatch.delenv("PROGRAM_IMPORT_TEXT_AI_ENABLED", raising=False)
    monkeypatch.delenv("PROGRAM_IMPORT_TEXT_AI_EVAL_REPORT", raising=False)
    from database.database_manager import DatabaseManager
    from svc.rate_limit import limiter

    limiter._storage.reset()
    catalog_path = tmp_path / "catalog.db"
    connection = sqlite3.connect(catalog_path)
    connection.execute(
        "CREATE TABLE exercises (id TEXT PRIMARY KEY, name TEXT, body_part TEXT, target_muscle TEXT, "
        "equipment TEXT, image_path TEXT, gif_path TEXT, instructions TEXT)"
    )
    connection.execute("CREATE TABLE exercise_secondary_muscles (exercise_id TEXT, muscle TEXT)")
    connection.execute(
        "INSERT INTO exercises (id, name, body_part, target_muscle, equipment) "
        "VALUES ('sq', 'Squat', 'Upper Legs', 'Quads', 'barbell'), "
        "('bp', 'Bench Press', 'Chest', 'Chest', 'barbell')"
    )
    connection.commit()
    connection.close()
    db = DatabaseManager(
        catalog_path=catalog_path,
        ledgers_dir=tmp_path / "users",
        backups_dir=tmp_path / "backups",
        default_ledger_id="bootstrap",
    )
    app = create_app()
    app.dependency_overrides[get_db] = lambda: db
    try:
        with TestClient(app) as client:
            yield client, db
    finally:
        if db.ledger_conn is not None:
            db.ledger_conn.close()
        db.catalog_conn.close()


def _register(client, username):
    response = client.post(
        "/auth/register",
        json={"trainee_id": username, "password": "correct-horse-1"},
    )
    assert response.status_code == 201, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _assignment(client, db, prefix="import"):
    coach_username = f"{prefix}-coach"
    player_username = f"{prefix}-player"
    coach_headers = _register(client, coach_username)
    invite = coach_service.issue_coach_invite(db, coach_username, actor="cli")
    assert invite["ok"]
    assert client.post(
        "/coach/invite/redeem", headers=coach_headers, json={"token": invite["token"]}
    ).status_code == 200
    assert client.put(
        "/coach/profile",
        headers=coach_headers,
        json={"display_name": "Coach", "bio": "", "specialization": "", "capacity": 5},
    ).status_code == 200
    player_headers = _register(client, player_username)
    assignment_invite = client.post("/coach/assignments/invites", headers=coach_headers).json()
    assignment = client.post(
        "/assignments/invites/redeem",
        headers=player_headers,
        json={"token": assignment_invite["token"], "consent": True},
    ).json()["assignment"]["assignment_id"]
    return coach_headers, player_headers, assignment


def _csv(rows):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(HEADERS)
    writer.writerows(rows)
    return output.getvalue().encode()


def _confirmed_rows(rows):
    return [
        {
            **{field: row.get(field) for field in CONFIRMED_ROW_FIELDS},
            "week_order": row.get("week_order") or 1,
        }
        for row in rows
    ]


def _week_csv(rows):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["week", *HEADERS])
    writer.writerows(rows)
    return output.getvalue().encode()


def _week_xlsx(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Program"
    sheet.append(["week", *HEADERS])
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _enable_program_import_ai(monkeypatch, tmp_path, mode="live"):
    from service import program_import_ai
    from service.program_import_evaluation import fixture_entries_hash
    from tests.eval import run_program_import_evaluation as evaluation_runner

    model, backend = program_import_ai.coach_model_identity()
    fixture = {
        "id": "test-sample",
        "description": "One synthetic row used to exercise the production enable gate.",
        "format": "csv",
        "tags": ["test"],
        "expected_tabs": ["CSV"],
        "expected_requires_tab_choice": False,
        "expected_selected_tab": "CSV",
        "expected_weeks": [],
        "expected_chosen_week": None,
        "expected_confirm_layout": False,
        "expected_rows": [{
            "source_row": 2, "week": None, "day": 1, "day_name": "Push",
            "input_exercise_name": "Squat", "exercise_name": "Squat",
            "expected_exercise_id": "sq", "expected_resolution": "library",
            "sets": 3, "reps_min": 6, "reps_max": 8, "approximation_codes": [],
        }],
    }
    entries = [{"filename": "01_sample.json", "fixture": fixture}]
    dataset_hash = fixture_entries_hash(entries)
    review = {
        "reviewed": True,
        "reviewed_by": "test owner",
        "reviewed_on": "2026-10-08",
        "reviewed_dataset_hash": dataset_hash,
    }
    actual = {
        "source_row": 2, "week": None, "day": 1, "day_name": "Push",
        "exercise_name": "Squat", "exercise_id": "sq", "sets": 3,
        "reps_min": 6, "reps_max": 8, "warnings": [],
    }
    import_result = {
        "rows": [actual], "errors": [], "detected_tabs": ["CSV"],
        "selected_tab": "CSV", "detected_weeks": [], "selected_week": None,
        "weeks_not_imported": [], "confirm_layout": False,
        "import_mode": "freeform",
    }
    scored = evaluation_runner.score_import_result(
        fixture, import_result, http_status=200, preflight=import_result
    )
    report = evaluation_runner.build_report(
        [scored], mode="live", privacy_pass=True, dataset_review=review,
        dataset_entries=entries, model_instance=None,
    )
    report["mode"] = mode
    report["model_run"] = "hosted_coach"
    report["model"] = model
    report["backend"] = backend
    report["pass"] = True
    assert report["model"] == model and report["backend"] == backend
    report_path = tmp_path / "program-import-report.json"
    report_path.write_text(json.dumps(report), encoding="utf-8")
    monkeypatch.setenv("PROGRAM_IMPORT_AI_ENABLED", "true")
    monkeypatch.setenv("PROGRAM_IMPORT_AI_EVAL_REPORT", str(report_path))
    return program_import_ai


def test_csv_import_creates_only_program_draft_and_replacement_requires_confirmation(api):
    client, db = api
    coach_headers, player_headers, assignment_id = _assignment(client, db)
    active_before = client.get("/programs/active", headers=player_headers).json()
    upload = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("plan.csv", _csv([["1", "Push", "1", "Squat", "3", "6-8", "2", "", "80", "", "90", "301", "Keep form"]]), "text/csv")},
    )
    assert upload.status_code == 200, upload.text
    result = upload.json()
    assert result["draft_exists"] is False
    assert result["rows"][0]["exercise_id"] == "sq"
    assert result["rows"][0]["target_rir"] == 2
    assert result["rows"][0]["load"] == {"kind": "kg", "value": 80}
    assert "rest_seconds" in result["rows"][0]
    assert "rest_sec" not in result["rows"][0]

    reversed_reps = {**result["rows"][0], "reps_min": 9, "reps_max": 8}
    invalid = client.post(
        f"/coach/assignments/{assignment_id}/program-draft/import",
        headers=coach_headers,
        json={"program_name": "plan", "rows": _confirmed_rows([reversed_reps])},
    )
    assert invalid.status_code == 400
    invalid_exercise = {**result["rows"][0], "exercise_id": "not-owned"}
    rejected_exercise = client.post(
        f"/coach/assignments/{assignment_id}/program-draft/import",
        headers=coach_headers,
        json={"program_name": "plan", "rows": _confirmed_rows([invalid_exercise])},
    )
    assert rejected_exercise.status_code == 400
    assert rejected_exercise.json()["message_code"] == "program_import.no_rows.v1"

    created = client.post(
        f"/coach/assignments/{assignment_id}/program-draft/import",
        headers=coach_headers,
        json={"program_name": "plan", "rows": _confirmed_rows(result["rows"])},
    )
    assert created.status_code == 200, created.text
    assert created.json()["draft"]["days"][0]["exercises"][0]["exercise_id"] == "sq"
    assert created.json()["draft"]["days"][0]["exercises"][0]["load"] == {"kind": "kg", "value": 80}
    assert len(created.json()["draft"]["weeks"]) == 1
    assert client.get("/programs/active", headers=player_headers).json() == active_before
    assert client.post(
        f"/coach/assignments/{assignment_id}/program-draft/import",
        headers=coach_headers,
        json={"program_name": "again", "rows": _confirmed_rows(result["rows"])},
    ).status_code == 409


def test_template_percent_load_becomes_program_load(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="pct-load")
    upload = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("plan.csv", _csv([["1", "Push", "1", "Squat", "3", "6-8", "2", "", "", "75", "90", "", ""]]), "text/csv")},
    )
    assert upload.status_code == 200, upload.text
    row = upload.json()["rows"][0]
    assert row["load"] == {"kind": "pct_e1rm", "value": 75}
    created = client.post(
        f"/coach/assignments/{assignment_id}/program-draft/import",
        headers=coach_headers,
        json={"program_name": "plan", "rows": _confirmed_rows([row])},
    )
    assert created.status_code == 200, created.text
    assert created.json()["draft"]["days"][0]["exercises"][0]["load"] == {
        "kind": "pct_e1rm",
        "value": 75,
    }
    minimum_upload = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={
            "file": (
                "minimum.csv",
                _csv([["1", "Push", "1", "Squat", "3", "6-8", "2", "", "", "1", "90", "", ""]]),
                "text/csv",
            )
        },
    )
    assert minimum_upload.status_code == 200, minimum_upload.text
    assert minimum_upload.json()["rows"][0]["load"] == {
        "kind": "pct_e1rm",
        "value": 1,
    }


def test_template_consecutive_group_rows_build_one_ordered_exercise_prescription(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="groups")
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["week", "day", "day_name", "order", "exercise", "group",
                     "sets", "reps", "rir", "rpe", "load_kg", "load_pct_e1rm",
                     "rest_sec", "tempo", "notes"])
    writer.writerows([
        [1, 1, "Push", 1, "Squat", "", 3, "5-8", 2, "", 100, "", 180, "", ""],
        [1, 1, "Push", 1, "Squat", 3, 1, "15", 2, "", "", "", 120, "", ""],
        [1, 1, "Push", 1, "Squat", 2, 2, "10-12", 1, "", "", "", 120, "", ""],
    ])
    upload = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("groups.csv", output.getvalue().encode(), "text/csv")},
    )
    assert upload.status_code == 200, upload.text
    confirmed = _confirmed_rows(upload.json()["rows"])
    created = client.post(
        f"/coach/assignments/{assignment_id}/program-draft/import",
        headers=coach_headers,
        json={"program_name": "grouped", "rows": confirmed},
    )
    assert created.status_code == 200, created.text
    exercise = created.json()["draft"]["days"][0]["exercises"][0]
    assert len(created.json()["draft"]["days"][0]["exercises"]) == 1
    assert [group["sets"] for group in exercise["set_groups"]] == [3, 2, 1]
    assert exercise["target_sets"] == 3
    assert exercise["target_reps_min"] == 5
    bad_rows = [dict(row) for row in confirmed]
    bad_rows[0]["group"] = "not a whole number"
    invalid = client.post(
        f"/coach/assignments/{assignment_id}/program-draft/import",
        headers=coach_headers,
        json={"program_name": "bad group", "rows": bad_rows},
    )
    assert invalid.status_code == 400, invalid.text
    assert invalid.json()["message_code"] == "program_import.group_invalid.v1"
    assert "column group" in invalid.json()["message_fallback"]


@pytest.mark.parametrize(
    ("raw_load", "expected_note"),
    [
        ("", None),
        ("—", None),
        ("-", None),
        ("n/a", None),
        ("60-70", "load_kg: 60-70"),
        ("BW", "load_kg: BW"),
        ("80 lb", "load_kg: 80 lb"),
        ("3,5", "load_kg: 3,5"),
        ("75%", "load_kg: 75%"),
        ("nan", "load_kg: nan"),
        ("inf", "load_kg: inf"),
    ],
)
def test_template_unsupported_load_is_preserved_without_rejecting_row(
    api, raw_load, expected_note
):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="load-note")
    upload = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={
            "file": (
                "plan.csv",
                _csv([["1", "Push", "1", "Squat", "3", "6-8", "2", "", raw_load, "", "90", "", ""]]),
                "text/csv",
            )
        },
    )

    assert upload.status_code == 200, upload.text
    row = upload.json()["rows"][0]
    assert row["valid"] is True
    assert row["load"] is None
    if expected_note is None:
        assert row["notes"] is None
        assert not any(
            warning["code"] == "program_import.load_preserved_as_note.v1"
            for warning in row["warnings"]
        )
    else:
        assert expected_note in row["notes"]
        marker = next(
            warning for warning in row["warnings"]
            if warning["code"] == "program_import.load_preserved_as_note.v1"
        )
        assert marker["message"] == raw_load


def test_csv_bare_rpe_values_convert_to_rir(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="bare-rpe")
    upload = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("plan.csv", _csv([
            ["1", "Push", "1", "Squat", "3", "6-8", "", "8", "", "", "90", "", ""],
            ["1", "Push", "2", "Squat", "3", "6-8", "", "7.5", "", "", "90", "", ""],
        ]), "text/csv")},
    )
    assert upload.status_code == 200, upload.text
    rows = upload.json()["rows"]
    assert [row["errors"] for row in rows] == [[], []]
    assert [row["target_rir"] for row in rows] == [2, 2.5]


def test_xlsx_requires_tab_choice_and_keeps_rpe_load_and_notes(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db)
    workbook = Workbook()
    program = workbook.active
    program.title = "Week A"
    program.append(HEADERS)
    program.append(["1", "يوم الدفع", "1", "Bench Press", "3", "8", "", "8", "60", "", "90", "", "ملاحظة أصلية"])
    notes = workbook.create_sheet("Notes")
    notes.append(["This tab is not a program"])
    file_bytes = io.BytesIO()
    workbook.save(file_bytes)
    first = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("plan.xlsx", file_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert first.status_code == 200, first.text
    assert first.json()["requires_tab_choice"] is True
    assert first.json()["rows"] == []

    second = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        data={"sheet": "Week A"},
        files={"file": ("plan.xlsx", file_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert second.status_code == 200, second.text
    row = second.json()["rows"][0]
    assert row["target_rir"] == 2
    assert row["load"] == {"kind": "kg", "value": 60}
    assert "load: 60" not in row["notes"]
    assert "ملاحظة أصلية" in row["notes"]
    assert not any("load" in warning["code"] for warning in row["warnings"])


def test_unknown_exercise_is_reported_and_invalid_reps_are_row_scoped(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db)
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("plan.csv", _csv([
            ["1", "Day A", "1", "Not in library", "3", "8", "2", "", "", "", "90", "", ""],
            ["1", "Day A", "2", "Squat", "3", "99", "2", "", "", "", "90", "", ""],
            ["1", "Day A", "3", "Bench Press", "3", "8", "hard", "", "", "", "90", "", ""],
        ]), "text/csv")},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["rows"][0]["exercise_id"] is None
    assert body["unresolved_names"][0]["name"] == "Not in library"
    assert body["rows"][1]["valid"] is False
    assert body["rows"][1]["errors"][0]["column"] == "reps"
    assert body["rows"][1]["errors"][0]["source_row"] == 3
    assert body["rows"][2]["valid"] is False
    assert body["rows"][2]["errors"][0]["code"] == "program_import.invalid_effort.v1"
    coach_exercise = client.post(
        "/coach/exercises", headers=coach_headers, json={"name": "Not in library"}
    )
    assert coach_exercise.status_code == 201
    body["rows"][0]["exercise_id"] = coach_exercise.json()["id"]
    confirmed_rows = [dict(body["rows"][0]), dict(body["rows"][1])]
    confirmed_rows[1]["reps_min"] = 99
    imported = client.post(
        f"/coach/assignments/{assignment_id}/program-draft/import",
        headers=coach_headers,
        json={"program_name": "plan", "rows": _confirmed_rows(confirmed_rows)},
    )
    assert imported.status_code == 200, imported.text
    imported_exercises = imported.json()["draft"]["days"][0]["exercises"]
    assert len(imported_exercises) == 1
    assert imported_exercises[0]["is_coach_exercise"] is True


def test_ambiguous_alias_is_offered_as_unconfirmed_suggestions(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db)
    db.catalog_conn.executemany(
        "INSERT INTO exercise_aliases (exercise_id, alias, normalized_alias) VALUES (?, ?, ?)",
        [("sq", "Smith incline", "smith incline"), ("bp", "Smith incline", "smith incline")],
    )
    db.catalog_conn.commit()
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("plan.csv", _csv([["1", "Day A", "1", "Smith incline", "3", "8", "2", "", "", "", "90", "", ""]]), "text/csv")},
    )
    assert response.status_code == 200, response.text
    row = response.json()["rows"][0]
    assert row["resolution"] == "ambiguous"
    assert row["exercise_id"] is None
    assert len(row["suggestions"]) == 2


def test_unsupported_amrap_is_kept_as_a_warning_and_more_than_500_rows_is_refused(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db)
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("plan.csv", _csv([["1", "Day A", "1", "Squat", "3", "AMRAP", "2", "", "", "", "90", "", ""]]), "text/csv")},
    )
    assert response.status_code == 200, response.text
    row = response.json()["rows"][0]
    assert "reps: AMRAP" in row["notes"]
    assert any(warning["code"] == "program_import.reps_approximated.v1" for warning in row["warnings"])
    too_many = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={
            "file": (
                "plan.csv",
                _csv([["1", "Day A", str(index), "Squat", "3", "8", "2", "", "", "", "90", "", ""] for index in range(501)]),
                "text/csv",
            )
        },
    )
    assert too_many.status_code == 413


def test_upload_limits_and_assignment_gate(api, monkeypatch, tmp_path, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db)
    _enable_program_import_ai(monkeypatch, tmp_path)
    freeform_file = ("own-layout.csv", b"Lift,Sets,Reps\nSquat,3,8", "text/csv")
    too_large = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("plan.csv", b"x" * (1024 * 1024 + 1), "text/csv")},
    )
    assert too_large.status_code == 413
    _, _, foreign_assignment_id = _assignment(client, db, prefix="foreign")
    foreign = client.post(
        f"/coach/assignments/{foreign_assignment_id}/program-import",
        headers=coach_headers,
        files={"file": freeform_file},
    )
    draft_action = client.post(
        f"/coach/assignments/{foreign_assignment_id}/program-draft/import",
        headers=coach_headers,
        json={
            "program_name": "plan",
            "rows": [{
                "source_row": 2,
                "day": 1,
                "order": 1,
                "exercise_name": "Squat",
                "exercise_id": "sq",
                "sets": 3,
                "reps_min": 8,
                "reps_max": 8,
                "target_rir": 2,
                "rest_seconds": 90,
            }],
        },
    )
    assert foreign.status_code == 403
    assert draft_action.status_code == 403
    assert foreign.json() == draft_action.json()
    unknown = client.post(
        "/coach/assignments/not-an-assignment/program-import",
        headers=coach_headers,
        files={"file": freeform_file},
    )
    assert unknown.status_code == 403
    assert foreign.json()["message_code"] == unknown.json()["message_code"] == "assignment.none_active.v1"
    assert client.post(
        f"/coach/assignments/{assignment_id}/revoke", headers=coach_headers
    ).status_code == 200
    ended = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": freeform_file},
    )
    assert ended.status_code == 403
    assert ended.json()["message_code"] == "assignment.none_active.v1"
    assert scripted_chat_model.calls == []


def test_template_downloads_are_authenticated(api):
    client, db = api
    coach_headers, _, _ = _assignment(client, db)
    csv_response = client.get("/coach/program-import/template.csv", headers=coach_headers)
    xlsx_response = client.get("/coach/program-import/template.xlsx", headers=coach_headers)
    assert csv_response.status_code == 200
    assert next(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))[0] == "week"
    assert xlsx_response.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True, data_only=True)
    assert next(workbook["Program"].values)[0] == "week"
    workbook.close()
    assert client.get("/coach/program-import/template.csv").status_code == 401


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_template_week_column_builds_ordered_weeks_in_draft(api, extension):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix=f"weeks-{extension}")
    rows = [
        [week, "1", "Full body", "1", "Squat", "3", "6-8", "2", "", "", "", "90", "", ""]
        for week in range(1, 5)
    ]
    payload = _week_csv(rows) if extension == "csv" else _week_xlsx(rows)
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": (f"four-weeks.{extension}", payload, "text/csv" if extension == "csv" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["detected_weeks"] == [f"Week {week}" for week in range(1, 5)]
    assert [row["week_order"] for row in result["rows"]] == [1, 2, 3, 4]

    created = client.post(
        f"/coach/assignments/{assignment_id}/program-draft/import",
        headers=coach_headers,
        json={"program_name": "four weeks", "rows": _confirmed_rows(result["rows"])},
    )
    assert created.status_code == 200, created.text
    draft = created.json()["draft"]
    assert [week["week_order"] for week in draft["weeks"]] == [1, 2, 3, 4]
    assert [week["week_name"] for week in draft["weeks"]] == [f"Week {week}" for week in range(1, 5)]
    assert [[day["day_name"] for day in week["days"]] for week in draft["weeks"]] == [
        ["Full body"], ["Full body"], ["Full body"], ["Full body"]
    ]
    assert [week["days"][0]["exercises"][0]["exercise_id"] for week in draft["weeks"]] == ["sq"] * 4


def test_template_blank_week_defaults_to_one_and_bad_weeks_are_reported_by_row(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="bad-week")
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={
            "file": (
                "bad-weeks.csv",
                _week_csv([
                    ["", "1", "Push", "1", "Bench Press", "3", "6-8", "2", "", "", "", "90", "", ""],
                    ["four", "1", "Push", "1", "Bench Press", "3", "6-8", "2", "", "", "", "90", "", ""],
                    [0, "1", "Push", "1", "Bench Press", "3", "6-8", "2", "", "", "", "90", "", ""],
                    [13, "1", "Push", "1", "Bench Press", "3", "6-8", "2", "", "", "", "90", "", ""],
                ]),
                "text/csv",
            )
        },
    )
    assert response.status_code == 200, response.text
    rows = response.json()["rows"]
    assert rows[0]["week_order"] == 1 and rows[0]["valid"] is True
    assert [row["valid"] for row in rows[1:]] == [False, False, False]
    assert [row["errors"][0]["code"] for row in rows[1:]] == [
        "program_import.week_invalid.v1",
        "program_import.week_invalid.v1",
        "program_import.week_invalid.v1",
    ]
    assert [row["errors"][0]["column"] for row in rows[1:]] == ["week"] * 3
    assert [row["errors"][0]["source_row"] for row in rows[1:]] == [3, 4, 5]


def test_effort_and_per_set_reps_are_approximated_with_original_text(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db)
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("effort.csv", _csv([
            ["1", "Day A", "1", "Squat", "3", "8,8,6", "7-8", "", "", "", "90", "", ""],
            ["1", "Day A", "2", "Bench Press", "3", "8/8/6", "2-3", "", "", "", "90", "", ""],
            ["1", "Day A", "3", "Squat", "3", "8", "", "RPE 8", "", "", "90", "", ""],
            ["1", "Day A", "4", "Bench Press", "3", "8", "@8", "", "", "", "90", "", ""],
            ["1", "Day A", "5", "Squat", "3", "8", "hard", "", "", "", "90", "", ""],
            ["1", "Day A", "6", "Bench Press", "3", "3x8,6,4", "2", "", "", "", "90", "", ""],
            ["1", "Day A", "7", "Squat", "3", "8", "", "8", "", "", "90", "", ""],
            ["1", "Day A", "8", "Bench Press", "3", "8", "7.5", "", "", "", "90", "", ""],
            ["1", "Day A", "9", "Squat", "3", "8", "8", "", "", "", "90", "", ""],
        ]), "text/csv")},
    )
    assert response.status_code == 200, response.text
    rows = response.json()["rows"]
    assert [row["target_rir"] for row in rows[:4]] == [5, 2.5, 2, 2]
    assert "reps: 8,8,6" in rows[0]["notes"]
    assert "effort: 7-8" in rows[0]["notes"]
    assert "reps: 8/8/6" in rows[1]["notes"]
    assert "[]" not in rows[2]["notes"]
    assert "effort: RPE 8" in rows[2]["notes"]
    assert any(warning["code"] == "program_import.rpe_converted.v1" for warning in rows[2]["warnings"])
    assert any(warning["code"] == "program_import.effort_approximated.v1" for warning in rows[3]["warnings"])
    assert rows[4]["valid"] is False
    assert rows[4]["errors"][0]["code"] == "program_import.invalid_effort.v1"
    assert rows[5]["reps_min"] == 4
    assert rows[5]["reps_max"] == 8
    assert any(warning["code"] == "program_import.reps_approximated.v1" for warning in rows[5]["warnings"])
    assert rows[6]["valid"] is True
    assert rows[6]["target_rir"] == 2
    assert rows[7]["target_rir"] == 5
    assert any(warning["code"] == "program_import.effort_approximated.v1" for warning in rows[7]["warnings"])
    assert rows[8]["target_rir"] == 5
    assert any(warning["code"] == "program_import.effort_approximated.v1" for warning in rows[8]["warnings"])


def test_freeform_import_retries_malformed_reply_and_returns_week_review_without_player_data(
    api, monkeypatch, tmp_path, scripted_chat_model
):
    client, db = api
    coach_headers, player_headers, assignment_id = _assignment(client, db, prefix="freeform")
    profile_response = client.put(
        "/profile",
        headers=player_headers,
        json={"weekly_frequency": 4, "current_goal": "PRIVATE_PROFILE_MARKER"},
    )
    assert profile_response.status_code == 200
    with db.open_ledger("freeform-player") as player_ledger:
        player_ledger.log_workout_session(
            "private-import-session",
            "2026-10-08",
            "PRIVATE_SPLIT_MARKER",
            "2026-10-08T08:00:00+00:00",
            "2026-10-08T08:30:00+00:00",
            notes="PRIVATE_LEDGER_MARKER",
        )
    _enable_program_import_ai(monkeypatch, tmp_path)
    from service.model_limits import reset_model_limits

    monkeypatch.setenv("MODEL_RATE_LIMIT_REQUESTS", "1")
    reset_model_limits()
    scripted_chat_model.script(
        {"rows": "malformed"},
        {
            "rows": [
                {
                    "source_row": 2,
                    "week": "Week 1",
                    "day": 1,
                    "day_name": "يوم الدفع",
                    "order": 1,
                    "exercise": "Bench Press",
                    "sets": 3,
                    "reps_min": 8,
                    "reps_max": 8,
                    "reps_original": "8",
                    "rir": None,
                    "rpe": 8,
                    "rest_seconds": 90,
                    "tempo": None,
                    "notes": "ملاحظة أصلية",
                    "original_text": "Bench Press · 3 × 8 · RPE 8",
                    "approximation_markers": [],
                },
                {
                    "source_row": 3,
                    "week": "Week 1",
                    "day": 1,
                    "day_name": "يوم الدفع",
                    "order": 2,
                    "exercise": "Squat",
                    "sets": 0,
                    "reps_min": 8,
                    "reps_max": 8,
                    "reps_original": "8",
                    "rir": 2,
                    "rpe": None,
                    "rest_seconds": 90,
                    "tempo": None,
                    "notes": None,
                    "original_text": "Squat · 0 × 8",
                    "approximation_markers": [],
                },
            ],
            "layout": {
                "days": [{"day": 1, "name": "يوم الدفع"}],
                "weeks": ["Week 1", "Week 2"],
                "confidence": 0.98,
                "confirm_layout": False,
            },
        },
    )
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={
            "file": (
                "freeform.csv",
                b"Exercise,Sets,Reps,Effort,Notes\nBench Press,3,8,RPE 8,\xd9\x85\xd9\x84\xd8\xa7\xd8\xad\xd8\xb8\xd8\xa9 \xd8\xa3\xd8\xb5\xd9\x84\xd9\x8a\xd8\xa9\nSquat,0,8,,",
                "text/csv",
            )
        },
    )

    assert response.status_code == 200, response.text
    result = response.json()
    assert result["detected_weeks"] == ["Week 1", "Week 2"]
    assert result["selected_week"] == "Week 1"
    assert result["weeks_not_imported"] == ["Week 2"]
    assert result["confirm_layout"] is False
    assert result["freeform_available"] is True
    assert result["layout_days"] == [{"day": 1, "name": "يوم الدفع"}]
    assert result["layout_confidence"] == 0.98
    row = result["rows"][0]
    assert row["day_name"] == "يوم الدفع"
    assert row["target_rir"] == 2
    assert "ملاحظة أصلية" in row["notes"]
    assert any(warning["code"] == "program_import.rpe_converted.v1" for warning in row["warnings"])
    invalid_row = result["rows"][1]
    assert invalid_row["valid"] is False
    assert invalid_row["errors"][0]["code"] == "program_import.invalid_sets.v1"
    assert invalid_row["errors"][0]["column"] == "sets"
    model_payloads = [
        str(message.content)
        for call in scripted_chat_model.calls
        for message in call["messages"]
    ]
    assert any("Bench Press" in payload for payload in model_payloads)
    assert all("freeform-player" not in payload for payload in model_payloads)
    assert all("PRIVATE_PROFILE_MARKER" not in payload for payload in model_payloads)
    assert all("PRIVATE_LEDGER_MARKER" not in payload for payload in model_payloads)
    assert all("PRIVATE_SPLIT_MARKER" not in payload for payload in model_payloads)
    limited = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={
            "file": (
                "freeform.csv",
                b"Exercise,Sets,Reps\nBench Press,3,8",
                "text/csv",
            )
        },
    )
    reset_model_limits()
    assert limited.status_code == 429
    assert limited.json()["message_code"] == "ai_limit.request_rate.v1"


def test_program_import_gate_off_keeps_template_working_and_cell_cap_refuses_freeform(
    api, monkeypatch, tmp_path, scripted_chat_model
):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="gate")
    monkeypatch.setenv("PROGRAM_IMPORT_AI_ENABLED", "false")
    monkeypatch.delenv("PROGRAM_IMPORT_AI_EVAL_REPORT", raising=False)
    template = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("plan.csv", _csv([["1", "Push", "1", "Squat", "3", "8", "2", "", "", "", "90", "", ""]]), "text/csv")},
    )
    assert template.status_code == 200, template.text

    disabled = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("own-layout.csv", b"Lift,Sets,Reps\nBench Press,3,8", "text/csv")},
    )
    assert disabled.status_code == 503
    assert disabled.json()["message_code"] == "program_import.freeform_disabled.v1"
    _enable_program_import_ai(monkeypatch, tmp_path, mode="mock")
    mock_refused = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("own-layout.csv", b"Lift,Sets,Reps\nBench Press,3,8", "text/csv")},
    )
    template_while_mocked = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("plan.csv", _csv([["1", "Push", "1", "Squat", "3", "8", "2", "", "", "", "90", "", ""]]), "text/csv")},
    )
    assert mock_refused.status_code == 503
    assert mock_refused.json()["message_code"] == "program_import.freeform_disabled.v1"
    assert template_while_mocked.status_code == 200, template_while_mocked.text
    assert template_while_mocked.json()["freeform_available"] is False
    assert scripted_chat_model.calls == []

    _enable_program_import_ai(monkeypatch, tmp_path)
    rows = [["Bench Press", "3", "8", "notes"] for _ in range(501)]
    freeform_csv = io.StringIO(newline="")
    writer = csv.writer(freeform_csv)
    writer.writerow(["Exercise", "Sets", "Reps", "Notes"])
    writer.writerows(rows)
    capped = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("large.csv", freeform_csv.getvalue().encode(), "text/csv")},
    )
    assert capped.status_code == 413
    assert capped.json()["message_code"] == "program_import.too_many_cells.v1"
    assert capped.json()["message_params"] == {"max_cells": 2000}
    assert scripted_chat_model.calls == []


def test_partial_template_header_uses_freeform_reader(api, monkeypatch, tmp_path, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="partial-template")
    _enable_program_import_ai(monkeypatch, tmp_path)
    scripted_chat_model.script(
        {
            "rows": [{
                "source_row": 602,
                "week": "Week 1",
                "day": 1,
                "day_name": "Push",
                "order": 1,
                "exercise": "Bench Press",
                "sets": 3,
                "reps_min": 8,
                "reps_max": 8,
                "reps_original": "8",
                "rir": None,
                "rpe": 8,
                "rest_seconds": 90,
                "tempo": None,
                "notes": None,
                "original_text": "Bench Press,3,8 @8",
                "approximation_markers": [],
            }],
            "layout": {
                "days": [{"day": 1, "name": "Push"}],
                "weeks": ["Week 1"],
                "confidence": 0.98,
                "confirm_layout": False,
            },
        }
    )

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("partial.csv", b"\n" * 600 + b"day,exercise,sets,Reps x RPE\n1,Bench Press,3,8 @8", "text/csv")},
    )

    assert response.status_code == 200, response.text
    assert response.json()["import_mode"] == "freeform"
    assert response.json()["rows"][0]["target_rir"] == 2
    assert response.json()["freeform_available"] is True


def test_freeform_import_runs_metered_at_coach_resource_cutoff(
    api, monkeypatch, tmp_path
):
    from tests.fakes.chat_model import ScriptedChatModel
    from utils import model_downloader
    from utils.model_metering import MeteringCallback

    client, db = api
    coach_headers, _player_headers, assignment_id = _assignment(
        client, db, prefix="import-cutoff"
    )
    coach_account_id = db.get_active_account_by_username("import-cutoff-coach")[
        "account_id"
    ]
    monkeypatch.setenv("AI_LAUNCH_ALLOWANCES_ENABLED", "true")
    monkeypatch.setenv("MODEL_DAILY_TOKEN_LIMIT", "1")
    monkeypatch.setenv("MODEL_RATE_LIMIT_REQUESTS", "0")
    for capability in ("COACH", "LIFTER"):
        for plan in ("FREE", "PRO"):
            monkeypatch.setenv(f"{capability}_AI_{plan}_DAILY_REQUESTS", "2")
            monkeypatch.setenv(
                f"{capability}_AI_{plan}_DAILY_RESOURCE_TOKENS", "10"
            )
    db.record_model_usage(
        coach_account_id,
        "coach",
        "deepseek-ai/DeepSeek-V4-Flash",
        10,
        0,
        0.0,
        False,
        purpose="coach_assistant",
        created_at=dt.datetime.now(dt.timezone.utc).isoformat(),
    )
    _enable_program_import_ai(monkeypatch, tmp_path)
    import_model = ScriptedChatModel(
        turns=[{
            "rows": [{
                "source_row": 2,
                "day": 1,
                "day_name": "Push",
                "order": 1,
                "exercise": "Squat",
                "sets": 3,
                "reps_min": 6,
                "reps_max": 8,
                "reps_original": "6-8",
                "rir": 2,
                "original_text": "Squat,3,6-8",
            }],
            "layout": {
                "days": [{"day": 1, "name": "Push"}],
                "weeks": [],
                "confidence": 0.98,
                "confirm_layout": False,
            },
        }],
        model_name="deepseek-ai/DeepSeek-V4-Flash",
        usage_metadata={
            "input_tokens": 40,
            "output_tokens": 20,
            "total_tokens": 60,
            "input_token_details": {},
        },
        callbacks=[MeteringCallback("deepseek-ai/DeepSeek-V4-Flash")],
    )
    monkeypatch.setattr(
        model_downloader, "get_coach_import_llm", lambda: import_model
    )

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={
            "file": (
                "own-layout.csv",
                b"Exercise,Sets,Reps\nSquat,3,6-8",
                "text/csv",
            )
        },
    )

    assert response.status_code == 200, response.text
    allowance = client.get("/auth/me", headers=coach_headers).json()["coach_ai_allowance"]
    assert allowance["status"] == "resource_cutoff"
    assert allowance["requests_used"] == 0
    assert allowance["resource_used"] == 70
    usage = db.catalog_conn.execute(
        "SELECT input_tokens, output_tokens FROM model_usage"
        " WHERE account_id = ? AND purpose = 'coach_program_import'",
        (coach_account_id,),
    ).fetchone()
    assert tuple(usage) == (40, 20)


def test_malformed_freeform_reply_retries_then_returns_a_clean_failure(
    api, monkeypatch, tmp_path, scripted_chat_model
):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="bad-model")
    _enable_program_import_ai(monkeypatch, tmp_path)
    scripted_chat_model.script({"rows": "bad"}, {"rows": "still bad"})

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("own-layout.csv", b"Lift,Sets,Reps\nBench Press,3,8", "text/csv")},
    )

    assert response.status_code == 422
    assert response.json()["message_code"] == "program_import.interpretation_failed.v1"
    assert "malformed" not in response.text.lower()
    assert len(scripted_chat_model.calls) == 2


@pytest.mark.parametrize("provider_failure", ["overload", "server_error", "connection_error", "timeout"])
def test_provider_unavailability_returns_busy_error_through_import_endpoint(
    provider_failure, api, monkeypatch, tmp_path, scripted_chat_model
):
    import httpx
    from langchain_openai.chat_models.base import OpenAIRateLimitError
    from openai import APIConnectionError, APITimeoutError, InternalServerError

    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix=f"busy-{provider_failure}")
    _enable_program_import_ai(monkeypatch, tmp_path)
    error_body = {
        "error": {
            "message": "Model busy, retry later",
            "type": "invalid_request_error",
            "code": "engine_overloaded",
        }
    }
    request = httpx.Request("POST", "https://api.deepinfra.com/v1/openai/chat/completions")
    if provider_failure in {"overload", "server_error"}:
        status_code = 429 if provider_failure == "overload" else 503
        response = httpx.Response(status_code, request=request, json=error_body)
        error_type = OpenAIRateLimitError if provider_failure == "overload" else InternalServerError
        provider_error = error_type("provider unavailable", response=response, body=error_body)
    elif provider_failure == "connection_error":
        provider_error = APIConnectionError(message="connection failed", request=request)
    else:
        provider_error = APITimeoutError(request=request)
    scripted_chat_model.script(provider_error)

    import_response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("own-layout.csv", b"Lift,Sets,Reps\nBench Press,3,8", "text/csv")},
    )

    assert import_response.status_code == 503
    assert import_response.json()["message_code"] == "program_import.model_busy.v1"
    assert import_response.json()["detail"] == (
        "The import assistant is busy with another request. Try again in a minute."
    )


def test_freeform_import_returns_all_weeks_and_defaults_to_week_one(
    api, monkeypatch, tmp_path, scripted_chat_model
):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="week-choice")
    _enable_program_import_ai(monkeypatch, tmp_path)
    reply = {
        "rows": [
            {
                "source_row": 2,
                "week": "Week 1",
                "day": 1,
                "day_name": "Push",
                "order": 1,
                "exercise": "Bench Press",
                "sets": 3,
                "reps_min": 8,
                "reps_max": 8,
                "reps_original": "8",
                "rir": 2,
                "rpe": None,
                "rest_seconds": 90,
                "tempo": None,
                "notes": None,
                "original_text": "Bench Press, 3, 8",
                "approximation_markers": [],
            },
            {
                "source_row": 3,
                "week": "Week 2",
                "day": 1,
                "day_name": "Push",
                "order": 1,
                "exercise": "Squat",
                "sets": 3,
                "reps_min": 8,
                "reps_max": 8,
                "reps_original": "8",
                "rir": 2,
                "rpe": None,
                "rest_seconds": 90,
                "tempo": None,
                "notes": None,
                "original_text": "Squat, 3, 8",
                "approximation_markers": [],
            },
        ],
        "layout": {
            "days": [{"day": 1, "name": "Push"}],
            "weeks": ["Week 1", "Week 2"],
            "confidence": 0.99,
            "confirm_layout": False,
        },
    }
    scripted_chat_model.script(reply)
    upload = {
        "file": (
            "own-layout.csv",
            b"Exercise,Sets,Reps\nBench Press,3,8\nSquat,3,8",
            "text/csv",
        )
    }

    default_week = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files=upload,
    )
    assert default_week.status_code == 200, default_week.text
    assert default_week.json()["selected_week"] == "Week 1"
    assert [row["exercise_name"] for row in default_week.json()["rows"]] == ["Bench Press", "Squat"]
    assert [row["week"] for row in default_week.json()["rows"]] == ["Week 1", "Week 2"]
    assert [row["exercise_name"] for row in default_week.json()["rows_by_week"]["Week 1"]] == ["Bench Press"]
    assert [row["exercise_name"] for row in default_week.json()["rows_by_week"]["Week 2"]] == ["Squat"]
    assert default_week.json()["weeks_not_imported"] == ["Week 2"]


def test_freeform_source_cells_recover_day_names_and_canonical_week_labels():
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    source_grid = {
        "sheet": "CSV",
        "cells": program_import._compact_cell_grid([
            ["Day", "Session", "Exercise", "Week 1 sets", "Week 1 reps", "Week 1 effort", "Week 2 sets", "Week 2 reps", "Week 2 effort"],
            [1, "Lower A", "Squat", 3, 8, "2 RIR", 4, 6, "@8"],
            [2, "Push A", "Bench Press", 3, 10, "2 RIR", 4, 8, "@8"],
        ]),
    }
    rows = []
    for source_row, day, day_name, exercise in (
        (2, 1, "Day 1", "Squat"),
        (2, 1, None, "Squat"),
        (3, "Push A", "Day 2", "Bench Press"),
        (3, "Push A", "Day 2", "Bench Press"),
    ):
        week = "1" if len(rows) % 2 == 0 else "2"
        rows.append({
            "source_row": source_row, "week": week, "day": day, "day_name": day_name,
            "exercise": exercise, "sets": 3, "reps_min": 8, "reps_max": 8,
            "rir": 2, "rpe": None, "original_text": exercise,
        })
    reply = ProgramImportReply.model_validate({
        "rows": rows,
        "layout": {
            "days": [{"day": 1, "name": "Day 1"}, {"day": 2, "name": "Day 2"}],
            "weeks": ["1", "2"], "confidence": 0.99,
        },
    })

    source = program_import._freeform_reader_output(reply, ["CSV"], "CSV", source_grid)
    assert source.detected_weeks == ["Week 1", "Week 2"]
    assert [row.week for row in source.rows] == ["Week 1", "Week 2"] * 2
    assert [(row.day, row.day_name) for row in source.rows] == [
        (1, "Lower A"), (1, "Lower A"), (2, "Push A"), (2, "Push A"),
    ]
    normalized = [program_import._normalise_row(row, {}) for row in source.rows]
    assert any(warning["code"] == "program_import.rpe_converted.v1" for warning in normalized[1]["warnings"])


def test_single_week_title_does_not_turn_day_numbers_into_week_selection():
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    source_grid = {
        "sheet": "CSV",
        "cells": program_import._compact_cell_grid([
            ["Week A / primary work"],
            ["Day", "Session", "Exercise", "Sets", "Reps", "Effort"],
            [1, "Lower A", "Squat", 3, 8, "2 RIR"],
            [2, "Push A", "Bench Press", 3, 8, "2 RIR"],
            [3, "Pull A", "Row", 3, 8, "2 RIR"],
        ]),
    }
    rows = [
        {
            "source_row": row_number, "week": str(day), "day": day,
            "day_name": day_name, "exercise": exercise, "sets": 3,
            "reps_min": 8, "reps_max": 8, "rir": 2, "original_text": exercise,
        }
        for row_number, day, day_name, exercise in (
            (3, 1, "Lower A", "Squat"),
            (4, 2, "Push A", "Bench Press"),
            (5, 3, "Pull A", "Row"),
        )
    ]
    reply = ProgramImportReply.model_validate({
        "rows": rows,
        "layout": {
            "days": [{"day": 1, "name": "Lower A"}, {"day": 2, "name": "Push A"}, {"day": 3, "name": "Pull A"}],
            "weeks": ["1", "2", "3"], "confidence": 0.99,
        },
    })

    source = program_import._freeform_reader_output(reply, ["CSV"], "CSV", source_grid)
    assert source.detected_weeks == []
    assert all(row.week is None for row in source.rows)


# #400 found note cells being misread as day/week layout titles.
def test_grid_titles_ignore_day_notes_and_preserve_model_day_names():
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    source_grid = {
        "sheet": "CSV",
        "cells": program_import._compact_cell_grid([
            ["Day", "Session", "Exercise", "Sets", "Reps", "Notes"],
            [1, "Sheet Push", "Bench Press", 3, 8, ""],
            ["Day 9 - deload", "", "", "", "", ""],
            [2, "Pull", "Squat", 3, 8, ""],
        ]),
    }
    reply = ProgramImportReply.model_validate({
        "rows": [
            {
                "source_row": 2, "day": 1, "day_name": "Model Push", "exercise": "Bench Press",
                "sets": 3, "reps_min": 8, "reps_max": 8, "original_text": "Bench Press",
            },
            {
                "source_row": 4, "day": 2, "day_name": "Pull", "exercise": "Squat",
                "sets": 3, "reps_min": 8, "reps_max": 8, "original_text": "Squat",
            },
        ],
        "layout": {
            "days": [{"day": 1, "name": "Model Push"}, {"day": 2, "name": "Pull"}],
            "weeks": [], "confidence": 0.99,
        },
    })

    source = program_import._freeform_reader_output(reply, ["CSV"], "CSV", source_grid)

    assert [(row.day, row.day_name) for row in source.rows] == [(1, "Model Push"), (2, "Pull")]
    context = program_import._freeform_grid_context(source_grid, reply.layout.days)
    assert 9 not in context.day_names


def test_week_note_next_to_exercise_data_does_not_add_a_week():
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    source_grid = {
        "sheet": "CSV",
        "cells": program_import._compact_cell_grid([
            ["Week 3 - accumulation"],
            ["Day", "Session", "Exercise", "Sets", "Reps", "Notes"],
            [1, "Push", "Bench Press", 3, 8, "Wk 4 - test 1RM"],
            [2, "Pull", "Squat", 3, 8, ""],
        ]),
    }
    reply = ProgramImportReply.model_validate({
        "rows": [
            {
                "source_row": source_row, "week": "Week 3", "day": day, "day_name": name,
                "exercise": exercise, "sets": 3, "reps_min": 8, "reps_max": 8,
                "original_text": exercise,
            }
            for source_row, day, name, exercise in (
                (3, 1, "Push", "Bench Press"), (4, 2, "Pull", "Squat"),
            )
        ],
        "layout": {
            "days": [{"day": 1, "name": "Push"}, {"day": 2, "name": "Pull"}],
            "weeks": ["Week 3"], "confidence": 0.99,
        },
    })

    source = program_import._freeform_reader_output(reply, ["CSV"], "CSV", source_grid)

    assert source.detected_weeks == []
    assert source.selected_week is None
    assert source.weeks_not_imported == []
    assert all(row.week is None for row in source.rows)


@pytest.mark.parametrize("model_weeks", [["Week 3"], []])
def test_week_note_between_exercise_rows_does_not_override_model_layout(model_weeks):
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    source_grid = {
        "sheet": "Week 3",
        "cells": program_import._compact_cell_grid([
            ["Day", "Session", "Exercise", "Sets", "Reps", "Notes"],
            [1, "Push", "Bench Press", 3, 8, ""],
            ["", "", "", "", "", "Wk 4 - test 1RM"],
            [2, "Pull", "Squat", 3, 8, ""],
        ]),
    }
    reply = ProgramImportReply.model_validate({
        "rows": [
            {
                "source_row": source_row,
                "week": model_weeks[0] if model_weeks else None,
                "day": day,
                "day_name": name, "exercise": exercise, "sets": 3,
                "reps_min": 8, "reps_max": 8, "original_text": exercise,
            }
            for source_row, day, name, exercise in (
                (2, 1, "Push", "Bench Press"), (4, 2, "Pull", "Squat"),
            )
        ],
        "layout": {
            "days": [{"day": 1, "name": "Push"}, {"day": 2, "name": "Pull"}],
            "weeks": model_weeks, "confidence": 0.99,
        },
    })

    source = program_import._freeform_reader_output(reply, ["Week 3"], "Week 3", source_grid)

    assert source.detected_weeks == []
    assert source.selected_week is None
    assert source.weeks_not_imported == []


def test_day_title_in_later_exercise_column_cannot_rename_generic_model_day():
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    source_grid = {
        "sheet": "CSV",
        "cells": program_import._compact_cell_grid([
            ["Day", "Session", "Exercise", "Sets", "Reps"],
            [1, "Push", "Bench Press", 3, 8],
            ["", "", "", "", "", "Day", "Session", "Movement", "Sets", "Reps"],
            ["", "", "", "", "", "", "", "Day 2 - optional cardio"],
            ["", "", "", "", "", 2, "Day 2", "Squat", 3, 8],
        ]),
    }
    reply = ProgramImportReply.model_validate({
        "rows": [
            {
                "source_row": 2, "day": 1, "day_name": "Push", "exercise": "Bench Press",
                "sets": 3, "reps_min": 8, "reps_max": 8, "original_text": "Bench Press",
            },
            {
                "source_row": 5, "day": 2, "day_name": "Day 2", "exercise": "Squat",
                "sets": 3, "reps_min": 8, "reps_max": 8, "original_text": "Squat",
            },
        ],
        "layout": {
            "days": [{"day": 1, "name": "Push"}, {"day": 2, "name": "Day 2"}],
            "weeks": [], "confidence": 0.99,
        },
    })

    source = program_import._freeform_reader_output(reply, ["CSV"], "CSV", source_grid)
    context = program_import._freeform_grid_context(source_grid, reply.layout.days)

    assert context.day_names == {1: "Push"}
    assert all(row.day_name != "optional cardio" for row in source.rows)
    assert program_import._has_unplaced_exercise_rows(context, reply.rows[:1]) is True


def test_layout_day_rename_keeps_column_derived_name_to_number_mapping():
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    source_grid = {
        "cells": program_import._compact_cell_grid([
            ["Day", "Session", "Exercise", "Sets", "Reps"],
            [1, "Column Push", "Bench Press", 3, 8],
        ])
    }
    reply = ProgramImportReply.model_validate({
        "rows": [{
            "source_row": 2, "day": 1, "day_name": "Model Push", "exercise": "Bench Press",
            "sets": 3, "reps_min": 8, "reps_max": 8, "original_text": "Bench Press",
        }],
        "layout": {
            "days": [{"day": 1, "name": "Model Push"}],
            "weeks": [], "confidence": 0.99,
        },
    })

    context = program_import._freeform_grid_context(source_grid, reply.layout.days)

    assert context.day_names[1] == "Model Push"
    assert context.day_numbers["column push"] == 1
    assert context.day_numbers["model push"] == 1


@pytest.mark.parametrize("note_column", [0, 5])
def test_week_note_below_table_does_not_create_a_week(note_column):
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    rows = [
        ["Week 3 - accumulation"],
        ["Day", "Session", "Exercise", "Sets", "Reps", "Notes"],
        [1, "Push", "Bench Press", 3, 8, ""],
        [2, "Pull", "Squat", 3, 8, ""],
    ]
    note = [""] * 6
    note[note_column] = "Wk 4 - test 1RM"
    rows.append(note)
    source_grid = {"sheet": "Week 3", "cells": program_import._compact_cell_grid(rows)}
    reply = ProgramImportReply.model_validate({
        "rows": [
            {
                "source_row": source_row, "week": "Week 3", "day": day,
                "day_name": name, "exercise": exercise, "sets": 3,
                "reps_min": 8, "reps_max": 8, "original_text": exercise,
            }
            for source_row, day, name, exercise in (
                (3, 1, "Push", "Bench Press"), (4, 2, "Pull", "Squat"),
            )
        ],
        "layout": {
            "days": [{"day": 1, "name": "Push"}, {"day": 2, "name": "Pull"}],
            "weeks": ["Week 3"], "confidence": 0.99,
        },
    })

    source = program_import._freeform_reader_output(reply, ["Week 3"], "Week 3", source_grid)

    assert source.detected_weeks == []
    assert source.selected_week is None
    assert source.weeks_not_imported == []
    assert all(row.week is None for row in source.rows)


@pytest.mark.parametrize(
    ("title", "model_week"),
    [
        ("Week 3 — Coach shorthand retained from notebook", "Week 3"),
        ("Week 3", "Week 3"),
        ("Wk 2 heavy", "Week 2"),
        ("الأسبوع 3", "Week 3"),
    ],
)
def test_single_week_sheet_title_does_not_tag_rows_with_that_week(title, model_week):
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    source_grid = {
        "sheet": "CSV",
        "cells": program_import._compact_cell_grid([
            [title],
            ["Day", "Session", "Exercise", "Sets", "Reps"],
            [1, "Lower A", "Back Squat", 3, 8],
            [2, "Push A", "Bench Press", 3, 8],
        ]),
    }
    reply = ProgramImportReply.model_validate({
        "rows": [
            {
                "source_row": source_row, "week": model_week, "day": day,
                "day_name": day_name, "exercise": exercise, "sets": 3,
                "reps_min": 8, "reps_max": 8, "original_text": exercise,
            }
            for source_row, day, day_name, exercise in (
                (3, 1, "Lower A", "Back Squat"), (4, 2, "Push A", "Bench Press"),
            )
        ],
        "layout": {
            "days": [{"day": 1, "name": "Lower A"}, {"day": 2, "name": "Push A"}],
            "weeks": [model_week], "confidence": 0.99,
        },
    })

    source = program_import._freeform_reader_output(reply, ["CSV"], "CSV", source_grid)
    context = program_import._freeform_grid_context(source_grid, reply.layout.days)

    assert context.week_labels == [model_week]
    assert source.detected_weeks == []
    assert source.selected_week is None
    assert source.weeks_not_imported == []
    assert all(row.week is None for row in source.rows)


def test_unlabeled_week_blocks_expand_model_rows_from_source_values():
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    fixture_path = Path(__file__).parent / "eval/datasets/program_import/21_unlabeled_week_blocks.json"
    case = json.loads(fixture_path.read_text(encoding="utf-8"))
    tab_name, fixture_rows = next(iter(case["sheets"].items()))
    source_grid = {
        "sheet": tab_name,
        "cells": program_import._compact_cell_grid(fixture_rows),
    }
    first_week_rows = [row for row in case["expected_rows"] if row["week"] == "Week 1"]
    day_orders: dict[int, int] = {}
    model_rows = []
    for expected in first_week_rows:
        day = expected["day"]
        order = day_orders.get(day, 0) + 1
        day_orders[day] = order
        model_rows.append({
            "source_row": expected["source_row"], "week": None, "day": day,
            "day_name": expected["day_name"], "order": order,
            "exercise": expected["input_exercise_name"], "sets": expected["sets"],
            "reps_min": expected["reps_min"], "reps_max": expected["reps_max"],
            "rir": expected["rir"], "rpe": expected["rpe"],
            "original_text": expected["original_text"],
        })
    reply = ProgramImportReply.model_validate({
        "rows": model_rows,
        "layout": {
            "days": [{"day": day, "name": name} for day, name in (
                (1, "Lower A"), (2, "Push A"), (3, "Pull A"),
            )],
            "weeks": [], "confidence": 0.99,
        },
    })

    source = program_import._freeform_reader_output(reply, [tab_name], tab_name, source_grid)
    normalized = [program_import._normalise_row(row, {}) for row in source.rows]

    assert source.detected_weeks == ["Week 1", "Week 2"]
    assert source.confirm_layout is True
    assert [
        (row.source_row, row.week, int(row.sets), int(row.reps_min))
        for row in source.rows
    ] == [
        (row["source_row"], row["week"], row["sets"], row["reps_min"])
        for row in case["expected_rows"]
    ]
    assert not any(
        error["code"] == "program_import.invalid_order.v1"
        for row in normalized
        for error in row["errors"]
    )
    assert [row["target_rir"] for row in normalized] == [
        float(row["rir"]) for row in case["expected_rows"]
    ]


@pytest.mark.parametrize(
    ("order_base", "order_step", "expected_step"),
    [(0, 1, 1), (1, 1, 1), (1, 2, 2)],
    ids=["zero-based-renumbered", "one-based-kept", "gapped-in-range-kept"],
)
def test_issue_400_unlabeled_week_blocks_normalize_orders_and_preserve_notes(
    api, monkeypatch, tmp_path, scripted_chat_model, order_base, order_step, expected_step
):
    client, coach_headers, assignment_id, _case, expected_week_one, workbook = (
        _issue_400_fixture21_setup(api, monkeypatch, tmp_path, "week-order")
    )

    def model_reply():
        day_orders = {}
        rows = []
        for expected in expected_week_one:
            day = expected["day"]
            order = day_orders[day] + order_step if day in day_orders else order_base
            day_orders[day] = order
            rows.append({
                "source_row": expected["source_row"],
                "week": "",
                "day": day,
                "day_name": expected["day_name"],
                "order": order,
                "exercise": expected["input_exercise_name"],
                "sets": expected["sets"],
                "reps_min": expected["reps_min"],
                "reps_max": expected["reps_max"],
                "rir": expected["rir"],
                "rpe": expected["rpe"],
                "notes": expected["notes"],
                "original_text": expected["original_text"],
            })
        return {
            "rows": rows,
            "layout": {
                "days": [
                    {"day": 1, "name": "Lower A"},
                    {"day": 2, "name": "Push A"},
                    {"day": 3, "name": "Pull A"},
                ],
                "weeks": [],
                "confidence": 0.99,
                "confirm_layout": True,
            },
        }

    scripted_chat_model.script(model_reply())

    response = _post_issue_400_fixture21(client, coach_headers, assignment_id, workbook)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["detected_weeks"] == ["Week 1", "Week 2"]
    assert result["selected_week"] == "Week 1"
    assert result["confirm_layout"] is True
    assert len(result["rows"]) == 24
    assert {row["week"] for row in result["rows"]} == {"Week 1", "Week 2"}
    assert all(row["valid"] for row in result["rows"])
    assert not any("[]" in (row["notes"] or "") for row in result["rows"])

    expected_notes = {row["source_row"]: row["notes"] for row in expected_week_one}
    grouped_orders = {}
    for row in result["rows"]:
        grouped_orders.setdefault((row["week"], row["day"]), []).append(
            (row["source_row"], row["order"])
        )
        note = expected_notes[row["source_row"]]
        assert row["notes"] == note
    for rows in grouped_orders.values():
        ordered = sorted(rows)
        assert [order for _source_row, order in ordered] == list(
            range(1, expected_step * len(ordered) + 1, expected_step)
        )


def _issue_400_fixture21_setup(api, monkeypatch, tmp_path, prefix):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix=prefix)
    fixture_path = (
        Path(__file__).parent
        / "eval/datasets/program_import/21_unlabeled_week_blocks.json"
    )
    case = json.loads(fixture_path.read_text(encoding="utf-8"))
    tab_name, fixture_rows = next(iter(case["sheets"].items()))
    expected_week_one = [row for row in case["expected_rows"] if row["week"] == "Week 1"]
    db.catalog_conn.executemany(
        "INSERT INTO exercises (id, name, body_part, target_muscle, equipment) "
        "VALUES (?, ?, 'Other', 'Other', 'barbell')",
        sorted({
            (row["expected_exercise_id"], row["exercise_name"])
            for row in expected_week_one
        }),
    )
    db.catalog_conn.commit()
    _enable_program_import_ai(monkeypatch, tmp_path)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = tab_name
    for row in fixture_rows:
        sheet.append(row)
    workbook_bytes = io.BytesIO()
    workbook.save(workbook_bytes)
    workbook.close()
    return client, coach_headers, assignment_id, case, expected_week_one, workbook_bytes.getvalue()


def _issue_400_fixture21_reply(expected_rows):
    model_rows = []
    order_by_source = {}
    for expected in expected_rows:
        order_key = (expected["day"], expected["source_row"])
        if order_key not in order_by_source:
            order_by_source[order_key] = 1 + sum(
                day == expected["day"] for day, _source_row in order_by_source
            )
        model_rows.append({
            "source_row": expected["source_row"],
            "week": "",
            "day": expected["day"],
            "day_name": expected["day_name"],
            "order": order_by_source[order_key],
            "exercise": expected["input_exercise_name"],
            "sets": expected["sets"],
            "reps_min": expected["reps_min"],
            "reps_max": expected["reps_max"],
            "rir": expected["rir"],
            "rpe": expected["rpe"],
            "notes": expected["notes"],
            "original_text": expected["original_text"],
        })
    return {
        "rows": model_rows,
        "layout": {
            "days": [
                {"day": day, "name": name}
                for day, name in dict.fromkeys(
                    (row["day"], row["day_name"]) for row in expected_rows
                )
            ],
            "weeks": [],
            "confidence": 0.99,
            "confirm_layout": True,
        },
    }


def _post_issue_400_fixture21(client, coach_headers, assignment_id, workbook_bytes):
    return client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": (
            "blocks.xlsx",
            workbook_bytes,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )},
    )


def test_issue_400_repeated_blank_week_rows_map_to_implicit_weeks(
    api, monkeypatch, tmp_path, scripted_chat_model
):
    client, coach_headers, assignment_id, case, _week_one, workbook = (
        _issue_400_fixture21_setup(api, monkeypatch, tmp_path, "duplicate-week-blocks")
    )
    scripted_chat_model.script(_issue_400_fixture21_reply(case["expected_rows"]))

    response = _post_issue_400_fixture21(client, coach_headers, assignment_id, workbook)

    assert response.status_code == 200, response.text
    result = response.json()
    assert result["detected_weeks"] == ["Week 1", "Week 2"]
    assert result["confirm_layout"] is True
    assert len(result["rows"]) == 24
    assert all(row["valid"] for row in result["rows"])
    expected = {
        (row["source_row"], row["week"]): row for row in case["expected_rows"]
    }
    for row in result["rows"]:
        source = expected[(row["source_row"], row["week"])]
        assert (row["sets"], row["reps_min"], row["reps_max"]) == (
            source["sets"], source["reps_min"], source["reps_max"],
        )
        assert row["target_rir"] == float(source["rir"])


def test_issue_400_partial_duplicate_rows_fill_only_missing_implicit_weeks(
    api, monkeypatch, tmp_path, scripted_chat_model
):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="partial-duplicate-weeks")
    _enable_program_import_ai(monkeypatch, tmp_path)
    first_row = {
        "source_row": 2, "week": "", "day": 1, "day_name": "Push", "order": 1,
        "exercise": "Bench Press", "sets": 2, "reps_min": 8, "reps_max": 8,
        "rir": 3, "original_text": "Model week one",
    }
    second_row = {
        **first_row,
        "sets": 3,
        "reps_min": 7,
        "reps_max": 7,
        "rir": 2,
        "original_text": "Model week two",
    }
    scripted_chat_model.script({
        "rows": [first_row, second_row],
        "layout": {
            "days": [{"day": 1, "name": "Push"}],
            "weeks": [],
            "confidence": 0.99,
            "confirm_layout": True,
        },
    })
    csv_content = (
        b"Training day,Session name,Movement,Set,Rep,Effort,Set,Rep,Effort,Set,Rep,Effort\n"
        b"1,Push,Bench Press,2,8,2 RIR,3,6,1 RIR,4,5,1 RIR"
    )

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("three-weeks.csv", csv_content, "text/csv")},
    )

    assert response.status_code == 200, response.text
    result = response.json()
    rows_by_week = {row["week"]: row for row in result["rows"]}
    assert len(rows_by_week) == 3
    assert all(row["valid"] for row in rows_by_week.values()), [
        (row["week"], row["errors"]) for row in rows_by_week.values()
    ]
    assert (
        rows_by_week["Week 1"]["sets"],
        rows_by_week["Week 1"]["reps_min"],
        rows_by_week["Week 1"]["target_rir"],
    ) == (2, 8, 3)
    assert (
        rows_by_week["Week 2"]["sets"],
        rows_by_week["Week 2"]["reps_min"],
        rows_by_week["Week 2"]["target_rir"],
    ) == (3, 7, 2)
    assert (
        rows_by_week["Week 3"]["sets"],
        rows_by_week["Week 3"]["reps_min"],
        rows_by_week["Week 3"]["target_rir"],
    ) == (4, 5, 1)


def test_issue_400_repeated_blank_week_rows_over_implicit_count_are_refused(
    api, monkeypatch, tmp_path, scripted_chat_model
):
    client, coach_headers, assignment_id, _case, expected_week_one, workbook = (
        _issue_400_fixture21_setup(api, monkeypatch, tmp_path, "too-many-duplicate-weeks")
    )
    repeated_source = expected_week_one[0]["source_row"]
    model_rows = [
        expected
        for row in expected_week_one
        for expected in [row] * (3 if row["source_row"] == repeated_source else 1)
    ]
    scripted_chat_model.script(_issue_400_fixture21_reply(model_rows))

    response = _post_issue_400_fixture21(client, coach_headers, assignment_id, workbook)

    assert response.status_code == 422
    assert response.json()["message_code"] == "program_import.interpretation_failed.v1"
    assert response.json()["detail"] == "The sheet could not be interpreted. Try the MYOS template."


@pytest.mark.parametrize("second_reply_repeats", [True, False], ids=["refused", "retry-recovers"])
def test_issue_400_duplicate_blank_weeks_without_implicit_blocks_get_one_repair(
    api, monkeypatch, tmp_path, scripted_chat_model, second_reply_repeats
):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="duplicate-no-weeks")
    _enable_program_import_ai(monkeypatch, tmp_path)
    repeated_row = {
        "source_row": 2,
        "week": "",
        "day": 1,
        "day_name": "Push",
        "order": 1,
        "exercise": "Bench Press",
        "sets": 3,
        "reps_min": 8,
        "reps_max": 8,
        "rir": 2,
        "original_text": "Bench Press, 3, 8",
    }
    layout = {
        "days": [{"day": 1, "name": "Push"}],
        "weeks": [],
        "confidence": 0.99,
        "confirm_layout": False,
    }
    duplicate_reply = {"rows": [repeated_row, repeated_row.copy()], "layout": layout}
    single_reply = {"rows": [repeated_row], "layout": layout}
    scripted_chat_model.script(
        duplicate_reply, duplicate_reply if second_reply_repeats else single_reply
    )

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("own-layout.csv", b"Exercise,Sets,Reps\nBench Press,3,8", "text/csv")},
    )

    if second_reply_repeats:
        assert response.status_code == 422
        assert response.json()["message_code"] == "program_import.interpretation_failed.v1"
        assert response.json()["detail"] == "The sheet could not be interpreted. Try the MYOS template."
    else:
        assert response.status_code == 200, response.text
        assert [row["source_row"] for row in response.json()["rows"]] == [2]


def test_issue_400_explicit_same_week_duplicates_remain_rejected(
    api, monkeypatch, tmp_path, scripted_chat_model
):
    client, coach_headers, assignment_id, _case, expected_week_one, workbook = (
        _issue_400_fixture21_setup(api, monkeypatch, tmp_path, "explicit-duplicate-week")
    )

    def duplicate_week_reply():
        reply = _issue_400_fixture21_reply(expected_week_one + [expected_week_one[0]])
        for row in reply["rows"]:
            row["week"] = "Week 1"
        reply["layout"]["weeks"] = ["Week 1", "Week 2"]
        return reply

    scripted_chat_model.script(duplicate_week_reply(), duplicate_week_reply())

    response = _post_issue_400_fixture21(client, coach_headers, assignment_id, workbook)

    assert response.status_code == 422
    assert response.json()["message_code"] == "program_import.interpretation_failed.v1"
    assert response.json()["detail"] == "The sheet could not be interpreted. Try the MYOS template."


def test_implicit_week_expansion_preserves_model_values_and_normalizes_missing_source_values():
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    source_grid = {
        "sheet": "CSV",
        "cells": program_import._compact_cell_grid([
            ["Day", "Session", "Exercise", "Sets", "Reps", "Effort", "Sets", "Reps", "Effort"],
            [1, "Push", "Bench Press", 5, 7, "3 RIR", 4, "8-10", "2 RIR"],
        ]),
    }
    reply = ProgramImportReply.model_validate({
        "rows": [{
            "source_row": 2, "week": None, "day": 1, "day_name": "Push",
            "exercise": "Bench Press", "sets": 3, "reps_min": 6, "reps_max": 6,
            "rir": 4, "original_text": "Bench Press",
        }],
        "layout": {
            "days": [{"day": 1, "name": "Push"}], "weeks": [], "confidence": 0.99,
        },
    })

    source = program_import._freeform_reader_output(reply, ["CSV"], "CSV", source_grid)
    first_week, second_week = source.rows
    normalized_second_week = program_import._normalise_row(second_week, {})

    assert (first_week.week, first_week.sets, first_week.reps_min, first_week.reps_max, first_week.rir) == (
        "Week 1", 3, 6, 6, 4,
    )
    assert (
        second_week.week, second_week.sets, second_week.reps_min,
        second_week.reps_max, normalized_second_week["target_rir"],
    ) == ("Week 2", 4, 8, 10, 2)


def test_two_distinct_grid_and_model_week_labels_are_reported():
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    source_grid = {
        "sheet": "CSV",
        "cells": program_import._compact_cell_grid([
            ["Week 3"],
            ["Day", "Session", "Exercise", "Sets", "Reps"],
            [1, "Push", "Bench Press", 3, 8],
            ["Week 4"],
            [2, "Pull", "Squat", 3, 8],
        ]),
    }
    reply = ProgramImportReply.model_validate({
        "rows": [{
            "source_row": 3, "week": "Week 3", "day": 1, "day_name": "Push",
            "exercise": "Bench Press", "sets": 3, "reps_min": 8, "reps_max": 8,
            "original_text": "Bench Press",
        }],
        "layout": {
            "days": [{"day": 1, "name": "Push"}],
            "weeks": ["Week 3", "Week 4"], "confidence": 0.99,
        },
    })

    context = program_import._freeform_grid_context(source_grid, reply.layout.days)
    source = program_import._freeform_reader_output(reply, ["CSV"], "CSV", source_grid)

    assert context.week_labels == ["Week 3"]
    assert source.detected_weeks == ["Week 3", "Week 4"]
    assert source.selected_week == "Week 3"
    assert source.weeks_not_imported == ["Week 4"]


def test_stacked_day_title_rows_preserve_names_without_day_column():
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    source_grid = {
        "sheet": "CSV",
        "cells": program_import._compact_cell_grid([
            ["Day 1 - Upper"],
            ["Exercise", "Sets", "Reps"],
            ["Bench Press", 3, 8],
            ["Day 2 - Lower"],
            ["Squat", 3, 8],
        ]),
    }
    reply = ProgramImportReply.model_validate({
        "rows": [
            {
                "source_row": 3, "day": 1, "day_name": "Day 1", "exercise": "Bench Press",
                "sets": 3, "reps_min": 8, "reps_max": 8, "original_text": "Bench Press",
            },
            {
                "source_row": 5, "day": 2, "day_name": "Day 2", "exercise": "Squat",
                "sets": 3, "reps_min": 8, "reps_max": 8, "original_text": "Squat",
            },
        ],
        "layout": {
            "days": [{"day": 1, "name": "Day 1"}, {"day": 2, "name": "Day 2"}],
            "weeks": [], "confidence": 0.99,
        },
    })

    source = program_import._freeform_reader_output(reply, ["CSV"], "CSV", source_grid)

    assert [(row.day, row.day_name) for row in source.rows] == [(1, "Upper"), (2, "Lower")]


def test_single_implicit_day_without_day_column_does_not_need_confirmation():
    from service import program_import
    from types import SimpleNamespace

    grid = {
        "cells": program_import._compact_cell_grid([
            ["Exercise", "Sets", "Reps"],
            ["Bench Press", 3, 8],
            ["Squat", 3, 8],
        ])
    }
    context = program_import._freeform_grid_context(grid, [])
    rows = [
        SimpleNamespace(source_row=2, day=None, day_name=None),
        SimpleNamespace(source_row=3, day=None, day_name=None),
    ]

    assert context.header_row is None
    assert context.columns == {}
    assert program_import._needs_layout_confirmation(context, rows) is False


def test_multiple_day_titles_without_axis_need_confirmation_when_unplaced():
    from service import program_import
    from types import SimpleNamespace

    grid = {
        "cells": program_import._compact_cell_grid([
            ["Day 1 - Upper"],
            ["Exercise", "Sets", "Reps"],
            ["Bench Press", 3, 8],
            ["Day 2 - Lower"],
            ["Squat", 3, 8],
        ])
    }
    context = program_import._freeform_grid_context(grid, [])
    rows = [
        SimpleNamespace(source_row=3, day=None, day_name=None),
        SimpleNamespace(source_row=5, day=None, day_name=None),
    ]

    assert program_import._needs_layout_confirmation(context, rows) is True


@pytest.mark.parametrize(
    ("candidate_exercise", "candidate_reps"),
    [("Warm-up: Band Pull Apart", 10), ("Bench Press", "AMRAP")],
)
def test_reported_or_rejected_candidate_row_does_not_need_confirmation(
    candidate_exercise, candidate_reps
):
    from service import program_import
    from types import SimpleNamespace

    grid = {
        "cells": program_import._compact_cell_grid([
            ["Day", "Exercise", "Sets", "Reps", "Notes"],
            [1, candidate_exercise, 1, candidate_reps, "reported as a notice"],
        ])
    }
    context = program_import._freeform_grid_context(grid, [])
    reported_row = SimpleNamespace(source_row=2, day=1, day_name="Day 1", exercise=None)

    assert program_import._has_unplaced_exercise_rows(context, [reported_row]) is False


def test_unreported_exercise_row_needs_confirmation():
    from service import program_import

    grid = {
        "cells": program_import._compact_cell_grid([
            ["Day", "Exercise", "Sets", "Reps"],
            [1, "Bench Press", 3, 8],
        ])
    }
    context = program_import._freeform_grid_context(grid, [])

    assert program_import._has_unplaced_exercise_rows(context, []) is True


def test_clear_recognized_layout_ignores_model_confirmation_hint_and_low_confidence():
    from service import program_import
    from types import SimpleNamespace

    grid = {
        "cells": program_import._compact_cell_grid([
            ["Day", "Session", "Exercise", "Sets", "Reps"],
            [1, "Push", "Bench Press", 3, 8],
        ])
    }
    context = program_import._freeform_grid_context(grid, [])
    model_rows = [SimpleNamespace(source_row=2, day=1, day_name="Push")]

    assert program_import._needs_layout_confirmation(context, model_rows, True, 0.3) is False


def test_recognized_header_without_day_axis_ignores_model_confirmation_hint():
    from service import program_import
    from types import SimpleNamespace

    grid = {
        "cells": program_import._compact_cell_grid([
            ["Exercise", "Sets", "Reps"],
            ["Bench Press", 3, 8],
            ["Squat", 3, 8],
        ])
    }
    context = program_import._freeform_grid_context(grid, [])
    model_rows = [
        SimpleNamespace(source_row=2, day=1, day_name="Day 1"),
        SimpleNamespace(source_row=3, day=1, day_name="Day 1"),
    ]

    assert context.header_row is None
    assert context.row_facts.header_rows == frozenset({1})
    assert program_import._needs_layout_confirmation(context, model_rows, True, 0.3) is False


@pytest.mark.parametrize(
    ("model_confirm_layout", "layout_confidence"),
    [(True, 0.99), (False, 0.74)],
)
def test_unrecognized_header_honors_model_confirmation_hint_or_low_confidence(
    model_confirm_layout, layout_confidence
):
    from service import program_import

    grid = {
        "cells": program_import._compact_cell_grid([
            ["Tag", "Übung", "Sätze", "Wdh"],
            [1, "Bench Press", 3, 8],
        ])
    }
    context = program_import._freeform_grid_context(grid, [])

    assert context.header_row is None
    assert program_import._needs_layout_confirmation(
        context, [], model_confirm_layout, layout_confidence
    ) is True


def test_structural_grid_ambiguity_requires_confirmation_even_with_confident_model():
    from service import program_import
    from types import SimpleNamespace

    grid = {
        "cells": program_import._compact_cell_grid([
            ["Day", "Session", "Exercise", "Sets", "Reps", "Sets", "Reps"],
            [1, "Push", "Bench Press", 3, 8, 4, 6],
        ])
    }
    context = program_import._freeform_grid_context(grid, [])
    model_rows = [SimpleNamespace(source_row=2, day=1, day_name="Push")]

    assert program_import._needs_layout_confirmation(context, model_rows, False, 0.99) is True


def test_arabic_headers_map_to_repaired_source_fields():
    from service import program_import

    grid = {
        "cells": program_import._compact_cell_grid([
            ["برنامج القوة"],
            ["اليوم", "اسم اليوم", "التمرين", "سيتات", "تكرارات", "RIR/RPE", "راحة", "الوزن", "ملاحظات"],
            [1, "دفع", "Bench Press", 3, 8, "2 RIR", 90, 60, "ثبت الكتفين"],
        ])
    }
    context = program_import._freeform_grid_context(grid, [])
    source = program_import._freeform_source_row(context, 3, None)

    assert source.fields == {
        "day": "1", "day_name": "دفع", "exercise": "Bench Press", "sets": "3",
        "reps": "8", "effort": "2 RIR", "rest_seconds": "90", "load_kg": "60",
        "notes": "ثبت الكتفين",
    }


@pytest.mark.parametrize(
    ("header", "expected_field"),
    [
        ("تمرين", "exercise"), ("التمرين", "exercise"),
        ("مجموعات", "sets"), ("جولات", "sets"), ("سيتات", "sets"),
        ("تكرارات", "reps"), ("عدات", "reps"),
        ("راحة", "rest_seconds"), ("RIR / RPE", "effort"),
        ("وزن / كجم", "load_kg"), ("الوزن", "load_kg"), ("ملاحظات", "notes"),
        ("Exercise / التمرين", "exercise"), ("Day / اليوم", "day"),
        ("Session / الحصة", "day_name"), ("ملاحظات المدرب", "notes"),
    ],
)
def test_arabic_coach_header_aliases_map_to_import_fields(header, expected_field):
    from service import program_import

    mapped_header = program_import._source_header_field(header)

    assert mapped_header is not None
    assert mapped_header[0] == expected_field


@pytest.mark.parametrize(
    ("header", "expected_week"),
    [("Week 2 sets", 2), ("الأسبوع 2 sets", 2), ("اسبوع 3 مجموعات", 3)],
)
def test_english_and_arabic_week_prefixes_map_to_source_headers(header, expected_week):
    from service import program_import

    assert program_import._source_header_field(header) == ("sets", expected_week)


def test_freeform_source_preserves_kg_per_set_and_superset_approximations():
    from service import program_import
    from service.program_import_ai import ProgramImportReply

    source_grid = {
        "sheet": "CSV",
        "cells": program_import._compact_cell_grid([
            ["Day", "Session", "Exercise", "Reps by set", "Load (kg)", "Effort", "Slot", "Load progression"],
            [1, "Lower A", "Squat", "10/8/6", 65, "RPE 8", "A1", "+2.5 kg/week"],
        ]),
    }
    reply = ProgramImportReply.model_validate({
        "rows": [{
            "source_row": 2, "week": None, "day": 1, "day_name": "Lower A",
            "exercise": "Squat", "sets": 3, "reps_min": 6, "reps_max": 10,
            "rir": 2, "rpe": None, "load_kg": 65, "original_text": "Squat",
        }],
        "layout": {"days": [{"day": 1, "name": "Lower A"}], "weeks": [], "confidence": 0.99},
    })

    source = program_import._freeform_reader_output(reply, ["CSV"], "CSV", source_grid)
    row = program_import._normalise_row(source.rows[0], {})
    warning_codes = {warning["code"] for warning in row["warnings"]}
    assert warning_codes >= {
        "program_import.reps_approximated.v1",
        "program_import.rpe_converted.v1",
        "program_import.unsupported_preserved_as_note.v1",
    }
    assert row["load"] == {"kind": "kg", "value": 65}
    assert "load_kg: 65 kg" not in row["notes"]
    assert "reps: 10/8/6" in row["notes"]
    assert "+2.5 kg/week" in row["notes"]


def test_freeform_model_suggestion_is_confirm_only_and_uses_library_candidates(
    api, monkeypatch, tmp_path, scripted_chat_model
):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="suggestion")
    _enable_program_import_ai(monkeypatch, tmp_path)
    scripted_chat_model.script(
        {
            "rows": [
                {
                    "source_row": 2,
                    "week": None,
                    "day": 1,
                    "day_name": "Push",
                    "order": 1,
                    "exercise": "بنش برس",
                    "exercise_search_terms": ["bench press"],
                    "sets": 3,
                    "reps_min": 8,
                    "reps_max": 8,
                    "reps_original": "8",
                    "rir": 2,
                    "rpe": None,
                    "rest_seconds": 90,
                    "tempo": None,
                    "notes": None,
                    "original_text": "بنش برس 3 × 8",
                    "approximation_markers": [],
                }
            ],
            "layout": {
                "days": [{"day": 1, "name": "Push"}],
                "weeks": [],
                "confidence": 0.95,
                "confirm_layout": False,
            },
        },
        {"suggestions": [{"exercise_name": "بنش برس", "exercise_ids": ["bp"]}]},
    )

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("own-layout.csv", b"Lift,Sets,Reps\n\xd8\xa8\xd9\x86\xd8\xb4 \xd8\xa8\xd8\xb1\xd8\xb3,3,8", "text/csv")},
    )

    assert response.status_code == 200, response.text
    row = response.json()["rows"][0]
    assert row["exercise_id"] is None
    assert row["exercise_name"] == "بنش برس"
    assert row["suggestions"] == [{"exercise_id": "bp", "name": "Bench Press"}]


def test_xlsx_date_reps_and_downloaded_template_sheet_behavior(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db)
    downloaded = client.get("/coach/program-import/template.xlsx", headers=coach_headers)
    assert downloaded.status_code == 200
    template = load_workbook(io.BytesIO(downloaded.content), data_only=True)
    assert template["Program"]["G1"].number_format == "@"
    assert template["Program"]["H1"].number_format == "@"
    assert template["Program"]["I1"].number_format == "@"
    assert template["Program"]["G500"].number_format == "@"
    template.close()
    upload_template = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("template.xlsx", downloaded.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert upload_template.status_code == 200, upload_template.text
    assert upload_template.json()["requires_tab_choice"] is False
    assert upload_template.json()["selected_tab"] == "Program"
    explicit_example = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        data={"sheet": "Example"},
        files={"file": ("template.xlsx", downloaded.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert explicit_example.status_code == 200, explicit_example.text
    assert explicit_example.json()["selected_tab"] == "Example"
    assert len(explicit_example.json()["rows"]) == 1

    rpe_without_reps = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={
            "file": (
                "missing-reps.csv",
                b"day,exercise,sets,rpe\n1,Squat,3,8\n",
                "text/csv",
            )
        },
    )
    assert rpe_without_reps.status_code == 503
    assert rpe_without_reps.json()["message_code"] == "program_import.freeform_disabled.v1"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Program"
    sheet.append(HEADERS)
    sheet.append([1, "Day", 1, "Squat", 3, dt.date(2026, 10, 8), 2, "", "", "", 90, "", ""])
    payload = io.BytesIO()
    workbook.save(payload)
    workbook.close()
    date_response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("date.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert date_response.status_code == 200, date_response.text
    error = date_response.json()["rows"][0]["errors"][0]
    assert error["code"] == "program_import.reps_as_date.v1"
    assert error["column"] == "reps"


def test_xlsx_file_and_row_limits_and_corrupt_workbooks_are_clean_errors(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db)
    too_large = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("large.xlsx", b"x" * (1024 * 1024 + 1), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert too_large.status_code == 413
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for index in range(501):
        sheet.append([1, "Day", index + 1, "Squat", 3, 8, 2, "", "", "", 90, "", ""])
    payload = io.BytesIO()
    workbook.save(payload)
    workbook.close()
    too_many = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("many.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert too_many.status_code == 413
    malformed_zip = io.BytesIO()
    with zipfile.ZipFile(malformed_zip, "w") as archive:
        archive.writestr("[Content_Types].xml", "<broken")
    corrupt = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("broken.xlsx", malformed_zip.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert corrupt.status_code == 400
    assert corrupt.json()["message_code"] == "program_import.file_invalid.v1"

    many_tabs = Workbook()
    many_tabs.active.title = "Program"
    for index in range(20):
        many_tabs.create_sheet(f"Sheet {index}")
    many_tab_bytes = io.BytesIO()
    many_tabs.save(many_tab_bytes)
    many_tabs.close()
    too_many_tabs = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("many-tabs.xlsx", many_tab_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert too_many_tabs.status_code == 413
    assert too_many_tabs.json()["message_code"] == "program_import.too_many_tabs.v1"


def test_invalid_day_does_not_conflict_or_rename_other_rows(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db)
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("days.csv", _csv([
            ["0", "Wrong", "1", "Squat", "3", "8", "2", "", "", "", "90", "", ""],
            ["1", "", "1", "Bench Press", "3", "8", "2", "", "", "", "90", "", ""],
            ["1", "Push", "2", "Squat", "3", "8", "2", "", "", "", "90", "", ""],
        ]), "text/csv")},
    )
    assert response.status_code == 200, response.text
    rows = response.json()["rows"]
    assert rows[0]["valid"] is False
    assert all(error["code"] != "program_import.conflicting_day_name.v1" for error in rows[0]["errors"])
    assert rows[1]["valid"] is True
    assert rows[1]["day_name"] == "Push"


def test_invalid_day_rows_do_not_count_toward_a_real_day_limit(api):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db)
    rows = [
        ["1", "Push", str(index), "Squat", "3", "8", "2", "", "", "", "90", "", ""]
        for index in range(1, 15)
    ]
    rows.append(["bad", "Push", "15", "Squat", "3", "8", "2", "", "", "", "90", "", ""])
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("day-limit.csv", _csv(rows), "text/csv")},
    )
    assert response.status_code == 200, response.text
    imported = response.json()["rows"]
    assert all(
        error["code"] != "program_import.too_many_exercises.v1"
        for row in imported
        for error in row["errors"]
    )
    assert imported[-1]["valid"] is False


def test_weekly_progression_is_flagged_even_when_preserved_as_a_load_note():
    from types import SimpleNamespace

    from service.program_import import Approximation, _detect_unsupported_prescriptions

    row = SimpleNamespace(original_text="Back Squat (Barbell)", notes="load_progression: +2.5 kg/week")
    existing = [Approximation("program_import.load_preserved_as_note.v1", "notes", "+2.5 kg/week")]

    found = _detect_unsupported_prescriptions(row, existing)

    assert [marker.code for marker in found] == ["program_import.unsupported_preserved_as_note.v1"]


def test_sheet_title_rule_reaches_every_exercise_row():
    from service import program_import

    grid = {
        "cells": program_import._compact_cell_grid([
            ["Compounds: one top set followed by back-off work"],
            ["Day", "Exercise", "Sets", "Reps"],
            [1, "Back Squat (Barbell)", 1, 5],
            [1, "Romanian Deadlift (Barbell)", 3, 10],
        ])
    }
    context = program_import._freeform_grid_context(grid, [])
    sources = [program_import._freeform_source_row(context, row_number, None) for row_number in (3, 4)]

    assert all("top set" in source.text for source in sources)


def test_full_inference_queue_reports_busy_not_uninterpretable(monkeypatch):
    import pytest

    from service import program_import, program_import_ai

    def busy(*_args, **_kwargs):
        raise TimeoutError("Inference queue is full; retry shortly.")

    monkeypatch.setattr(program_import_ai, "translate_sheet", busy)
    grid = program_import.FreeFormCellGrid({"sheet": "CSV", "cells": [{"address": "A1", "value": "Squat"}]}, ["CSV"], "CSV")

    with pytest.raises(program_import.ProgramImportError) as raised:
        program_import.FreeFormProgramImportReader(None).read(grid)

    assert raised.value.code == "program_import.model_busy.v1"
    assert raised.value.status_code == 503


def _enable_text_import_for_route(monkeypatch):
    from service import program_import_ai
    from service.evaluation_report_gate import GateStatus

    monkeypatch.setattr(program_import_ai, "resolve_text_enable_gate", lambda: GateStatus(True, True, "test report accepted"))


def _text_reply(line=2, *, week=None, exercise="Squat", sets=3, notes=None, markers=None, search_terms=None,
                original_text=None, day=1, day_name="Push"):
    return {
        "rows": [{"source_row": line, "week": week, "day": day, "day_name": day_name, "order": 1,
                  "exercise": exercise, "sets": sets, "reps_min": 8, "reps_max": 10, "rir": 2,
                  "notes": notes, "exercise_search_terms": search_terms or [],
                  "approximation_markers": markers or [], "original_text": original_text or exercise}],
        "layout": {"days": [{"day": 1, "name": "Push"}], "weeks": [week] if week else [],
                   "confidence": 0.98, "confirm_layout": False},
    }


def _valid_text_evaluation_report():
    import copy

    from service import program_import_ai
    from service.program_import_evaluation import fixture_entries_hash

    report = json.loads((Path(__file__).parents[1] / "reports/program_import_eval.json").read_text(encoding="utf-8"))
    report = copy.deepcopy(report)
    report["suite"] = "program_import_text"
    report["report_version"] = program_import_ai.TEXT_REPORT_VERSION
    report["prompt_hash"] = program_import_ai.text_prompt_version_hash()
    entries = report["dataset_fixtures"]
    for entry in entries:
        for row in entry["fixture"]["expected_rows"]:
            row["source_line"] = row.pop("source_row")
    report["dataset_hash"] = fixture_entries_hash(entries)
    report["dataset_review"]["reviewed_dataset_hash"] = report["dataset_hash"]
    for run in report["runs"]:
        run["checks"].pop("tab_choice", None)
        run["metrics"].pop("tab_choice_errors", None)
        for sheet in run["sheets"]:
            sheet.pop("tab_choice_correct", None)
            sheet["checks"].pop("tab_choice_correct", None)
            sheet["errors"] = [{**error, "source_line": error.pop("source_row")} for error in sheet["errors"]]
            for row in sheet["row_results"]:
                row["source_line"] = row.pop("source_row")
                row["reported_errors"] = [
                    {**error, "source_line": error.pop("source_row")}
                    for error in row["reported_errors"]
                ]
    report["gates"]["evaluation"]["metrics"].pop("tab_choice_errors", None)
    return report


def test_text_upload_and_paste_share_the_same_http_pipeline(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-equivalence")
    _enable_text_import_for_route(monkeypatch)
    scripted_chat_model.script(_text_reply(line=2), _text_reply(line=2))
    url = f"/coach/assignments/{assignment_id}/program-import"
    uploaded = client.post(url, headers=coach_headers, files={"file": ("plan.txt", b"Day 1 - Push\nSquat 3x8-10", "text/plain")})
    pasted = client.post(url, headers=coach_headers, data={"text": "Day 1 - Push\nSquat 3x8-10"})
    assert uploaded.status_code == pasted.status_code == 200
    assert uploaded.json() == pasted.json()
    assert uploaded.json()["import_mode"] == "text"
    assert uploaded.json()["text_available"] is True
    assert uploaded.json()["rows"][0]["source_row"] == 2


def test_text_route_requires_exactly_one_input(api, monkeypatch):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-input-choice")
    _enable_text_import_for_route(monkeypatch)
    url = f"/coach/assignments/{assignment_id}/program-import"
    neither = client.post(url, headers=coach_headers)
    empty = client.post(url, headers=coach_headers, data={"text": ""})
    both = client.post(url, headers=coach_headers, data={"text": "program"}, files={"file": ("x.txt", b"program", "text/plain")})
    assert neither.status_code == empty.status_code == both.status_code == 422
    assert neither.json().get("message_code") == empty.json().get("message_code") == both.json().get("message_code") == "program_import.input_required.v1", (neither.text, empty.text, both.text)


def test_text_empty_cap_invalid_utf8_and_bom(api, monkeypatch, scripted_chat_model):
    from service.program_import_constants import MAX_PROGRAM_IMPORT_TEXT_CHARS

    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-bounds")
    _enable_text_import_for_route(monkeypatch)
    url = f"/coach/assignments/{assignment_id}/program-import"
    for text_value, code in ((" \r\n\t", "program_import.text_empty.v1"),
                             ("x" * (MAX_PROGRAM_IMPORT_TEXT_CHARS + 1), "program_import.text_too_large.v1")):
        response = client.post(url, headers=coach_headers, data={"text": text_value})
        assert response.status_code == 422
        assert response.json().get("message_code") == code, response.text
    invalid = client.post(url, headers=coach_headers, files={"file": ("bad.txt", b"\xff", "text/plain")})
    assert invalid.status_code == 422
    assert invalid.json()["message_code"] == "program_import.text_encoding.v1"
    assert scripted_chat_model.calls == []
    scripted_chat_model.script(_text_reply(line=1))
    bom = client.post(url, headers=coach_headers, files={"file": ("bom.txt", "\ufeffSquat 3x8".encode(), "text/plain")})
    assert bom.status_code == 200, bom.text
    assert len(scripted_chat_model.calls) == 1


def test_text_week_selection_and_line_source_positions(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-weeks")
    _enable_text_import_for_route(monkeypatch)
    marker = {"code": "program_import.load_preserved_as_note.v1", "column": "notes", "original": "70% 1RM"}
    first = _text_reply(line=2, week="Week 1", notes="70% 1RM", markers=[marker])
    second = _text_reply(line=5, week="Week 2", exercise="Odd Lift", sets=0)["rows"][0]
    scripted_chat_model.script({"rows": first["rows"] + [second], "layout": {
        "days": [{"day": 1, "name": "Push"}], "weeks": ["Week 1", "Week 2"],
        "confidence": 0.7, "confirm_layout": True,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": "Squat 3x8\nSquat 3x8\n\nWeek 2\nOdd Lift"})
    assert response.status_code == 200, response.text
    import_result = response.json()
    assert import_result["selected_week"] == "Week 1"
    assert import_result["weeks_not_imported"] == ["Week 2"]
    assert import_result["confirm_layout"] is True
    assert import_result["rows"][1]["source_row"] == 5
    assert import_result["rows"][1]["errors"][0]["source_row"] == 5
    assert import_result["rows"][0]["notes"].find("70% 1RM") >= 0
    assert import_result["rows"][0]["warnings"][0]["source_row"] == 2


def test_text_line_number_validation_repairs_an_invalid_source_position(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-line-repair")
    _enable_text_import_for_route(monkeypatch)
    scripted_chat_model.script(_text_reply(line=99, original_text="Squat 3x8"),
                               _text_reply(line=2, original_text="Squat 3x8"))

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        data={"text": "Day 1 - Push\nSquat 3x8\nSquat 3x8"},
    )

    assert response.status_code == 200, response.text
    assert response.json()["rows"][0]["source_row"] == 2
    assert len(scripted_chat_model.calls) == 2


def test_text_live_failures_recover_day_line_name_and_markers(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-live-failures")
    _enable_text_import_for_route(monkeypatch)
    text = "Day 1 - Push\nbench 3x8\n\nDay 2 - Pull\nسحب أمامي 3x10 AMRAP"
    first = _text_reply(
        line=1, day=None, day_name="Push", exercise="Bench Press (Barbell)",
        original_text="bench 3x8",
    )["rows"][0]
    second = _text_reply(
        line=5, day=99, day_name="Pull", exercise="Lat Pulldown",
        original_text="سحب أمامي 3x10 AMRAP",
    )["rows"][0]
    second["approximation_markers"] = []
    scripted_chat_model.script({"rows": [first, second], "layout": {
        "days": [{"day": None, "name": "Push"}, {"day": 99, "name": "Pull"}],
        "weeks": [], "confidence": 0.4, "confirm_layout": True,
    }})

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        data={"text": text},
    )

    assert response.status_code == 200, response.text
    result = response.json()
    assert [row["day"] for row in result["rows"]] == [1, 2]
    assert all("program_import.invalid_day.v1" not in [error["code"] for error in row["errors"]]
               for row in result["rows"])
    assert [row["exercise_name"] for row in result["rows"]] == ["bench", "سحب أمامي"]
    assert "program_import.unsupported_preserved_as_note.v1" in [
        warning["code"] for warning in result["rows"][1]["warnings"]
    ]
    assert result["confirm_layout"] is False


def test_text_source_line_remaps_shifted_number_when_original_text_is_unique(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-shifted-source")
    _enable_text_import_for_route(monkeypatch)
    scripted_chat_model.script(_text_reply(line=1, original_text="bench 3x8"))

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        data={"text": "Day 1 - Push\n\nbench 3x8"},
    )

    assert response.status_code == 200, response.text
    assert response.json()["rows"][0]["source_row"] == 3


def test_text_source_line_with_duplicate_original_requires_repair(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-ambiguous-source")
    _enable_text_import_for_route(monkeypatch)
    bad = _text_reply(line=1, original_text="Squat 3x8")
    scripted_chat_model.script(bad, bad)

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        data={"text": "Day 1 - Push\nSquat 3x8\nSquat 3x8"},
    )

    assert response.status_code == 422
    assert len(scripted_chat_model.calls) == 2


def test_text_source_markers_cover_rep_lists_ranges_holds_and_failure(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-source-markers")
    _enable_text_import_for_route(monkeypatch)
    lines = [
        "bench 3 مجموعات × 12/10/8 عدة",
        "push ups 3x10 AMRAP",
        "side plank 3 x 30 sec hold",
        "side plank 3 x 30 sec hold — حتى قرب الفشل",
        "cable row 3x10 RIR 1-2",
        "hip thrust 3x8 حتى قرب الفشل",
    ]
    rows = []
    for index, line in enumerate(lines, start=1):
        rows.append(_text_reply(line=index, exercise="rewritten name", original_text=line)["rows"][0])
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Push"}], "weeks": [],
        "confidence": 0.99, "confirm_layout": False,
    }})

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        data={"text": "\n".join(lines)},
    )

    assert response.status_code == 200, response.text
    warnings = [[warning["code"] for warning in row["warnings"]] for row in response.json()["rows"]]
    assert "program_import.reps_approximated.v1" in warnings[0]
    assert "program_import.reps_approximated.v1" in warnings[1]
    assert "program_import.effort_approximated.v1" in warnings[2]
    assert "program_import.effort_approximated.v1" in warnings[4]
    assert "program_import.effort_approximated.v1" in warnings[5]
    assert response.json()["rows"][3]["reps_min"] == 30
    assert response.json()["rows"][3]["reps_max"] == 30
    assert "program_import.effort_approximated.v1" in warnings[3]


def test_text_week_labels_follow_model_and_repeated_day_names_alone_are_not(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-repeat-weeks")
    _enable_text_import_for_route(monkeypatch)
    case = json.loads((Path(__file__).parent / "eval/datasets/program_import_text/07_repeated_blocks.json").read_text())
    lines = case["text"].splitlines()
    second_block = lines.index("Full Body A", lines.index("Full Body A") + 1) + 1
    rows = []
    for line_number, line in enumerate(lines, 1):
        if not re.search(r"\d+\s*[x×]\s*\d+", line):
            continue
        row = _text_reply(line=line_number, exercise="rewritten", original_text=line)["rows"][0]
        row["day_name"] = "Full Body A"
        row["week"] = "Week 1" if line_number < second_block else "Week 2"
        rows.append(row)
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Full Body A"}], "weeks": ["Week 1", "Week 2"],
        "confidence": 0.98, "confirm_layout": False,
    }})

    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": case["text"]})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["detected_weeks"] == ["Week 1", "Week 2"]
    assert [row["week"] for row in result["rows"]].count("Week 1") == 8
    assert [row["week"] for row in result["rows"]].count("Week 2") == 8
    assert result["confirm_layout"] is False

    # Same six day labels in both halves do not establish a repeated week when
    # each repeated day has a different exercise.
    split_text = "\n".join(
        f"{day}\n{exercise} 3x8"
        for day, exercise in zip(
            ["Push", "Pull", "Legs", "Push", "Pull", "Legs"],
            ["Bench", "Row", "Squat", "Fly", "Curl", "Lunge"],
        )
    )
    split_rows = []
    for line_number, line in enumerate(split_text.splitlines(), 1):
        if "3x8" in line:
            day_name = split_text.splitlines()[line_number - 2]
            split_rows.append(_text_reply(line=line_number, day_name=day_name, exercise="rewritten", original_text=line)["rows"][0])
    scripted_chat_model.script({"rows": split_rows, "layout": {
        "days": [{"day": None, "name": name} for name in ("Push", "Pull", "Legs")], "weeks": [],
        "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": split_text})
    assert response.status_code == 200, response.text
    assert response.json()["detected_weeks"] == []
    assert [row["day"] for row in response.json()["rows"]] == [1, 2, 3, 4, 5, 6]
    assert "program_import.invalid_day.v1" in [
        error["code"] for error in response.json()["rows"][-1]["errors"]
    ]


def test_text_repeated_blocks_override_conflicting_model_weeks_and_keep_rows(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-repeat-weeks-conflict")
    _enable_text_import_for_route(monkeypatch)
    case = json.loads((Path(__file__).parent / "eval/datasets/program_import_text/07_repeated_blocks.json").read_text())
    rows = []
    for line_number, line in enumerate(case["text"].splitlines(), 1):
        if re.search(r"\d+\s*[x×]\s*\d+", line):
            rows.append(_text_reply(line=line_number, exercise="rewritten", day=1, day_name="Full Body A",
                                    week="Week 2", original_text=line)["rows"][0])
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Full Body A"}], "weeks": ["Week 2"],
        "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": case["text"]})
    assert response.status_code == 200, response.text
    result = response.json()
    assert len(result["rows"]) == 16
    assert result["detected_weeks"] == ["Week 1", "Week 2"]
    assert [row["week"] for row in result["rows"]].count("Week 1") == 8
    assert [row["week"] for row in result["rows"]].count("Week 2") == 8
    assert result["confirm_layout"] is False


def test_text_live_eval_6_repeated_progression_reply_confirms_false(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-live6-repeat-progression")
    _enable_text_import_for_route(monkeypatch)
    case = json.loads((Path(__file__).parent / "eval/datasets/program_import_text/07_repeated_blocks.json").read_text())
    replay = json.loads((Path(__file__).parent / "fixtures/program_import_text_live_eval_6_07.json").read_text())
    rows = []
    for recorded in replay["rows"]:
        row = _text_reply(
            line=recorded["source_row"], week=recorded["week"], day=recorded["day"],
            day_name=recorded["day_name"], exercise=recorded["input_exercise_name"],
            sets=recorded["sets"], notes=recorded["notes"], original_text=recorded["original_text"],
        )["rows"][0]
        row.update(
            order=recorded["order"], reps_min=recorded["reps_min"], reps_max=recorded["reps_max"],
            rir=recorded["target_rir"], rest_seconds=recorded["rest_seconds"], tempo=recorded["tempo"],
        )
        rows.append(row)
    scripted_chat_model.script({"rows": rows, "layout": replay["layout"]})

    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": case["text"]})
    assert response.status_code == 200, response.text
    result = response.json()
    assert len(result["rows"]) == 16
    assert [(row["week"], row["day"], row["day_name"]) for row in result["rows"]] == [
        ("Week 1", 1, "Full Body A") for _ in range(8)
    ] + [("Week 2", 1, "Full Body A") for _ in range(8)]
    assert result["confirm_layout"] is False

    scripted_chat_model.script({"rows": rows, "layout": {
        **replay["layout"], "confidence": 0.5, "confirm_layout": True,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": case["text"]})
    assert response.status_code == 200, response.text
    assert response.json()["confirm_layout"] is False


def test_text_raw_live_07_repeated_same_name_progression_replies_confirm_false(
    api, monkeypatch, scripted_chat_model
):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-raw-live07-replays")
    _enable_text_import_for_route(monkeypatch)
    case = json.loads((Path(__file__).parent / "eval/datasets/program_import_text/07_repeated_blocks.json").read_text())
    replies = json.loads((Path(__file__).parent / "fixtures/program_import_text_raw_live_07.json").read_text())

    for reply in replies:
        scripted_chat_model.script(reply)
        response = client.post(
            f"/coach/assignments/{assignment_id}/program-import",
            headers=coach_headers,
            data={"text": case["text"]},
        )
        assert response.status_code == 200, response.text
        result = response.json()
        assert len(result["rows"]) == 16
        assert [(row["week"], row["day"], row["day_name"]) for row in result["rows"]] == [
            ("Week 1", 1, "Full Body A") for _ in range(8)
        ] + [("Week 2", 1, "Full Body A") for _ in range(8)]
        assert result["confirm_layout"] is False


def test_text_progression_with_distinct_repeated_day_names_still_confirms(
    api, monkeypatch, scripted_chat_model
):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-distinct-repeat-day-name")
    _enable_text_import_for_route(monkeypatch)
    text = "Full Body A\nBench 3x8\nRow 3x8\nFull Body A\nBench 4x6\nRow 4x6"
    rows = [
        _text_reply(line=2, day=1, day_name="Full Body A", exercise="Bench", original_text="Bench 3x8")["rows"][0],
        _text_reply(line=3, day=1, day_name="Full Body A", exercise="Row", original_text="Row 3x8")["rows"][0],
        _text_reply(line=5, day=2, day_name="Full Body B", exercise="Bench", original_text="Bench 4x6")["rows"][0],
        _text_reply(line=6, day=2, day_name="Full Body B", exercise="Row", original_text="Row 4x6")["rows"][0],
    ]
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Full Body A"}, {"day": 2, "name": "Full Body B"}],
        "weeks": ["Week 1"], "confidence": 1.0, "confirm_layout": False,
    }})
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        data={"text": text},
    )
    assert response.status_code == 200, response.text
    assert response.json()["confirm_layout"] is True


def test_text_partial_repeated_block_match_still_requests_confirmation(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-partial-repeat-weeks")
    _enable_text_import_for_route(monkeypatch)
    first = ["Bench", "Row", "Squat", "Fly", "Curl"]
    second = ["Bench", "Row", "Squat", "Lunge", "Press"]
    text = "\n".join(
        ["Full Body A", *[f"{name} 3x8" for name in first], "Full Body A", *[f"{name} 4x8" for name in second]]
    )
    rows = []
    for line_number, line in enumerate(text.splitlines(), 1):
        if "x" in line:
            rows.append(_text_reply(line=line_number, week="Week 2", day=1, day_name="Full Body A",
                                    exercise="rewritten", original_text=line)["rows"][0])
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Full Body A"}], "weeks": ["Week 2"],
        "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": text})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["detected_weeks"] == ["Week 1", "Week 2"]
    assert result["confirm_layout"] is True


def test_text_live_eval_5_headingless_alternatives_request_confirmation(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-live5-headingless-layout")
    _enable_text_import_for_route(monkeypatch)
    case = json.loads((Path(__file__).parent / "eval/datasets/program_import_text/10_ambiguous_layout.json").read_text())
    replay = json.loads((Path(__file__).parent / "fixtures/program_import_text_live_eval_5_confirm.json").read_text())
    rows = []
    for recorded in replay["rows"]:
        row = _text_reply(
            line=recorded["source_row"], week=recorded["week"], day=recorded["day"],
            day_name=recorded["day_name"], exercise=recorded["exercise"],
            original_text=recorded["original_text"],
        )["rows"][0]
        row["order"] = recorded["order"]
        rows.append(row)
    scripted_chat_model.script({"rows": rows, "layout": replay["layout"]})

    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": case["text"]})
    assert response.status_code == 200, response.text
    assert len(response.json()["rows"]) == 15
    assert {row["day"] for row in response.json()["rows"]} == {1}
    assert response.json()["confirm_layout"] is True

    for text, expected_confirmation in (
        ("Bench 3x8\n--- maybe alternate split ---\nRow 3x8", True),
        ("Bench 3x8\nRow 3x8", False),
    ):
        rows = [
            _text_reply(line=line_number, day=1, day_name="Day 1", exercise="rewritten",
                        original_text=line)["rows"][0]
            for line_number, line in enumerate(text.splitlines(), 1)
            if "3x8" in line
        ]
        scripted_chat_model.script({"rows": rows, "layout": {
            "days": [{"day": 1, "name": "Day 1"}], "weeks": [],
            "confidence": 0.98, "confirm_layout": False,
        }})
        response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                               data={"text": text})
        assert response.status_code == 200, response.text
        assert response.json()["confirm_layout"] is expected_confirmation


def test_text_identical_ppl_conflict_keeps_one_consistent_model_reading(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-identical-ppl-weeks")
    _enable_text_import_for_route(monkeypatch)
    split_text = "\n".join(
        f"{day}\n{exercise} 3x8"
        for day, exercise in zip(
            ["Push", "Pull", "Legs", "Push", "Pull", "Legs"],
            ["Bench", "Row", "Squat", "Bench", "Row", "Squat"],
        )
    )
    lines = split_text.splitlines()
    rows = []
    for line_number, line in enumerate(lines, 1):
        if "3x8" in line:
            rows.append(_text_reply(line=line_number, day=line_number // 2, day_name=lines[line_number - 2],
                                    exercise="rewritten", original_text=line)["rows"][0])
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": day, "name": name} for day, name in enumerate(("Push", "Pull", "Legs", "Push", "Pull", "Legs"), 1)],
        "weeks": [], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": split_text})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["detected_weeks"] == []
    assert [row["day"] for row in result["rows"]] == [1, 2, 3, 4, 5, 6]
    assert [row["week"] for row in result["rows"]] == [None] * 6
    assert [item["day"] for item in result["layout_days"]] == [1, 2, 3, 4, 5, 6]
    assert "program_import.invalid_day.v1" in [
        error["code"] for error in result["rows"][-1]["errors"]
    ]
    assert result["confirm_layout"] is True


def test_text_sixth_day_is_preview_error_and_only_five_days_enter_draft(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-sixth-day-domain-cap")
    _enable_text_import_for_route(monkeypatch)
    names = ["Push", "Pull", "Legs", "Upper", "Lower", "Arms"]
    text_lines = []
    rows = []
    for day, name in enumerate(names, 1):
        text_lines.extend([name, f"Squat 3x8 day{day}"])
        rows.append(_text_reply(line=2 * day, week="Week 1", day=day, day_name=name, exercise="Squat",
                                original_text=f"Squat 3x8 day{day}")["rows"][0])
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": day, "name": name, "week": "Week 1"} for day, name in enumerate(names, 1)],
        "weeks": ["Week 1"], "confidence": 0.98, "confirm_layout": False,
    }})
    preview = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                          data={"text": "\n".join(text_lines)})
    assert preview.status_code == 200, preview.text
    body = preview.json()
    sixth = body["rows"][5]
    assert sixth["valid"] is False
    assert "program_import.invalid_day.v1" in [error["code"] for error in sixth["errors"]]
    created = client.post(
        f"/coach/assignments/{assignment_id}/program-draft/import", headers=coach_headers,
        json={"program_name": "five day cap", "rows": _confirmed_rows(body["rows"][:5])},
    )
    assert created.status_code == 200, created.text
    assert [day["day_order"] for day in created.json()["draft"]["weeks"][0]["days"]] == [1, 2, 3, 4, 5]


def test_text_long_holds_preview_and_confirm_with_representable_prescriptions(api, monkeypatch, scripted_chat_model):
    client, db = api
    db.catalog_conn.execute(
        "INSERT INTO exercises (id, name, body_part, target_muscle, equipment) "
        "VALUES ('plank', 'Plank', 'Core', 'Abs', 'bodyweight')"
    )
    db.catalog_conn.commit()
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-hold-confirm-roundtrip")
    _enable_text_import_for_route(monkeypatch)
    lines = ["Plank 3x30 sec hold", "Plank 3x45s hold", "Plank 3x60s hold"]
    rows = [_text_reply(line=index, exercise="Plank", original_text=line)["rows"][0]
            for index, line in enumerate(lines, 1)]
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Core"}], "weeks": [],
        "confidence": 0.98, "confirm_layout": False,
    }})
    preview = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                          data={"text": "\n".join(lines)})
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert [(row["reps_min"], row["reps_max"]) for row in body["rows"]] == [
        (30, 30), (8, 12), (8, 12),
    ]
    for row, duration in zip(body["rows"][1:], (45, 60), strict=True):
        assert row["valid"] is True
        assert "program_import.reps_approximated.v1" in [warning["code"] for warning in row["warnings"]]
        assert f"reps: {duration} sec hold" in row["notes"]
    created = client.post(
        f"/coach/assignments/{assignment_id}/program-draft/import", headers=coach_headers,
        json={"program_name": "hold round trip", "rows": _confirmed_rows(body["rows"])},
    )
    assert created.status_code == 200, created.text
    exercises = created.json()["draft"]["days"][0]["exercises"]
    assert [(exercise["target_reps_min"], exercise["target_reps_max"]) for exercise in exercises] == [
        (30, 30), (8, 12), (8, 12),
    ]
    assert "45 sec hold" in exercises[1]["notes"]
    assert "60 sec hold" in exercises[2]["notes"]


def test_text_repeated_identical_prescriptions_remain_ambiguous(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-identical-repeat-ambiguous")
    _enable_text_import_for_route(monkeypatch)
    text = "Full Body A\nBench 3x8\nRow 3x8\nFull Body A\nBench 3x8\nRow 3x8"
    rows = [
        _text_reply(line=2, week="Week 1", day=1, day_name="Full Body A", exercise="Bench",
                    original_text="Bench 3x8")["rows"][0],
        _text_reply(line=3, week="Week 1", day=1, day_name="Full Body A", exercise="Row",
                    original_text="Row 3x8")["rows"][0],
        _text_reply(line=5, week="Week 2", day=1, day_name="Full Body A", exercise="Bench",
                    original_text="Bench 3x8")["rows"][0],
        _text_reply(line=6, week="Week 2", day=1, day_name="Full Body A", exercise="Row",
                    original_text="Row 3x8")["rows"][0],
    ]
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Full Body A"}], "weeks": ["Week 1", "Week 2"],
        "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": text})
    assert response.status_code == 200, response.text
    assert [row["week"] for row in response.json()["rows"]] == ["Week 1", "Week 1", "Week 2", "Week 2"]
    assert response.json()["confirm_layout"] is True


def test_text_source_hold_and_per_set_reps_override_model_values(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-prescription-override")
    _enable_text_import_for_route(monkeypatch)
    lines = ["side plank 3x30s", "plank 3 x 30-45 sec", "bench 3x6/5/4", "Hold the bar at the top, 3x8"]
    rows = []
    for index, line in enumerate(lines, 1):
        row = _text_reply(line=index, exercise="rewritten", original_text=line)["rows"][0]
        row.update(reps_min=8, reps_max=12)
        rows.append(row)
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Push"}], "weeks": [], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": "\n".join(lines)})
    assert response.status_code == 200, response.text
    imported = response.json()["rows"]
    assert [(row["source_row"], row["reps_min"], row["reps_max"]) for row in imported] == [
        (1, 30, 30), (2, 8, 12), (3, 4, 6), (4, 8, 8),
    ]
    assert "program_import.reps_approximated.v1" in [warning["code"] for warning in imported[1]["warnings"]]
    assert "reps: 30-45 sec hold" in imported[1]["notes"]
    assert "program_import.effort_approximated.v1" not in [warning["code"] for warning in imported[3]["warnings"]]


def test_text_names_headings_multiline_source_and_layout_confirmation(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-extraction-and-layout")
    _enable_text_import_for_route(monkeypatch)
    examples = [
        ("1. Squat 4x5", "Squat"), ("Squat: 4x5", "Squat"), ("Push-up - 3x15", "Push-up"),
        ("A1. Bench 3x8", "Bench"), ("Bench press 3 sets of 8-10", "Bench press"),
        ("Pull-ups 3xAMRAP", "Pull-ups"), ("Bench 5/5/3 reps", "Bench"),
        ("3x5 squat", "squat"), ("Leg press x 12, 10, 8", "Leg press"),
    ]
    rows = []
    source_lines = []
    for index, (line, _expected) in enumerate(examples, 1):
        source_lines.append(line)
        row = _text_reply(line=index, exercise="rewritten model name", original_text=line)["rows"][0]
        rows.append(row)
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Push"}], "weeks": [], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": "\n".join(source_lines)})
    assert response.status_code == 200, response.text
    assert [row["input_exercise_name"] for row in response.json()["rows"]] == [name for _, name in examples]

    heading_text = "### Push\n**Push day**\nPush (Week 2)\nSquat 3x8"
    heading_rows = []
    for line_number, line in enumerate(heading_text.splitlines(), 1):
        heading_rows.append(_text_reply(line=line_number, day_name="Push", exercise="Push" if line_number < 4 else "Squat",
                                        original_text=line if line_number < 4 else "Squat 3x8")["rows"][0])
    scripted_chat_model.script({"rows": heading_rows, "layout": {
        "days": [{"day": 1, "name": "Push"}], "weeks": [], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": heading_text})
    assert response.status_code == 200, response.text
    assert [row["exercise_name"] for row in response.json()["rows"]] == ["Squat"]

    multiline = "Squat 3x8\nKeep the brace tight"
    row = _text_reply(line=1, exercise="Squat", original_text=multiline)["rows"][0]
    scripted_chat_model.script({"rows": [row], "layout": {
        "days": [{"day": 1, "name": "Push"}], "weeks": [], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": multiline})
    assert response.status_code == 200, response.text
    assert response.json()["rows"][0]["source_row"] == 1

    arabic = "اليوم الأول – Full Body A\nSquat 3x8\nاليوم التاني – رجل\nSquat 3x8"
    arabic_rows = [
        _text_reply(line=2, day=None, day_name="Full Body A", original_text="Squat 3x8")["rows"][0],
        _text_reply(line=4, day=None, day_name="اليوم التاني – رجل", original_text="Squat 3x8")["rows"][0],
    ]
    scripted_chat_model.script({"rows": arabic_rows, "layout": {
        "days": [{"day": None, "name": "اليوم الأول – Full Body A"}, {"day": None, "name": "اليوم التاني – رجل"}],
        "weeks": [], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": arabic})
    assert response.status_code == 200, response.text
    assert [row["day_name"] for row in response.json()["rows"]] == ["Full Body A", "رجل"]

    by_week = "Week 1\nPush\nBench 3x8\nPull\nRow 3x8\nWeek 2\nLegs\nSquat 3x8"
    week_rows = [
        _text_reply(line=3, week="Week 1", day_name="Push", exercise="Bench", original_text="Bench 3x8")["rows"][0],
        _text_reply(line=5, week="Week 1", day_name="Pull", exercise="Row", original_text="Row 3x8")["rows"][0],
        _text_reply(line=8, week="Week 2", day_name="Legs", exercise="Squat", original_text="Squat 3x8")["rows"][0],
    ]
    scripted_chat_model.script({"rows": week_rows, "layout": {
        "days": [{"day": 99, "name": "Legs", "week": "Week 2"}],
        "weeks": ["Week 1", "Week 2"], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": by_week})
    assert response.status_code == 200, response.text
    assert response.json()["layout_days"] == [{"day": 1, "name": "Legs"}]


def test_text_week_inference_requires_standalone_week_titles_and_confirms_exclusion(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-week-evidence")
    _enable_text_import_for_route(monkeypatch)

    six_day_lines = [
        "Push", "Bench 4x8", "Pull", "Row 4x8", "Legs", "Squat 4x6",
        "Push", "Incline press 4x10", "Pull", "Deadlift 3x5", "Legs", "Lunge 3x10",
    ]
    rows = []
    for line_number, line in enumerate(six_day_lines, 1):
        if re.search(r"\d+\s*[x×]\s*\d+", line):
            day_name = six_day_lines[line_number - 2]
            rows.append(_text_reply(line=line_number, week="Week 1", day=len(rows) + 1, day_name=day_name,
                                    exercise="rewritten", original_text=line)["rows"][0])
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [
            {"day": index, "name": name}
            for index, name in enumerate(("Push", "Pull", "Legs", "Push", "Pull", "Legs"), start=1)
        ],
        "weeks": ["Week 1"], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": "\n".join(six_day_lines)})
    assert response.status_code == 200, response.text
    assert response.json()["detected_weeks"] == []
    assert response.json()["weeks_not_imported"] == []
    assert [(row["day"], row["day_name"]) for row in response.json()["rows"]] == [
        (index, name) for index, name in enumerate(("Push", "Pull", "Legs", "Push", "Pull", "Legs"), start=1)
    ], [(row["source_row"], row["day"], row["day_name"]) for row in response.json()["rows"]]
    assert response.json()["layout_days"] == [
        {"day": index, "name": name}
        for index, name in enumerate(("Push", "Pull", "Legs", "Push", "Pull", "Legs"), start=1)
    ]
    assert response.json()["confirm_layout"] is False

    text = "Push\nBench 3x8\nPull\nRow 3x8"
    rows = [
        _text_reply(line=2, week="Week 2", day_name="Push", exercise="Bench", original_text="Bench 3x8")["rows"][0],
        _text_reply(line=4, week="Week 2", day_name="Pull", exercise="Row", original_text="Row 3x8")["rows"][0],
    ]
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Push"}, {"day": 2, "name": "Pull"}],
        "weeks": ["Week 2"], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": text})
    assert response.status_code == 200, response.text
    assert response.json()["detected_weeks"] == ["Week 2"]
    assert [row["week"] for row in response.json()["rows"]] == ["Week 2", "Week 2"]

    text = "Dips 3x12 superset with pushdowns, repeat\nBench 3x8"
    rows = [
        _text_reply(line=1, week="Week 1", day_name="Push", exercise="Dips", original_text="Dips 3x12 superset with pushdowns, repeat")["rows"][0],
        _text_reply(line=2, week="Week 1", day_name="Push", exercise="Bench", original_text="Bench 3x8")["rows"][0],
    ]
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Push"}], "weeks": ["Week 1"],
        "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": text})
    assert response.status_code == 200, response.text
    assert response.json()["detected_weeks"] == []
    assert response.json()["weeks_not_imported"] == []

    text = "Week 1\nPush\nBench 3x8\nWeek 2\nPull\nRow 3x8"
    rows = [
        _text_reply(line=3, week=None, day_name="Push", exercise="Bench", original_text="Bench 3x8")["rows"][0],
        _text_reply(line=6, week=None, day_name="Pull", exercise="Row", original_text="Row 3x8")["rows"][0],
    ]
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": None, "name": "Push"}, {"day": None, "name": "Pull"}],
        "weeks": [], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": text})
    assert response.status_code == 200, response.text
    assert response.json()["detected_weeks"] == ["Week 1", "Week 2"]
    assert response.json()["weeks_not_imported"] == ["Week 2"]
    assert response.json()["confirm_layout"] is True

    text = "الأسبوع الأول\nPush\nBench 3x8\nالأسبوع التاني\nPull\nRow 3x8"
    rows = [
        _text_reply(line=3, week=None, day_name="Push", exercise="Bench", original_text="Bench 3x8")["rows"][0],
        _text_reply(line=6, week=None, day_name="Pull", exercise="Row", original_text="Row 3x8")["rows"][0],
    ]
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": None, "name": "Push"}, {"day": None, "name": "Pull"}],
        "weeks": [], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": text})
    assert response.status_code == 200, response.text
    assert response.json()["detected_weeks"] == ["Week 1", "Week 2"]
    assert response.json()["confirm_layout"] is True


def test_text_rep_list_is_anchored_and_exact_effort_is_not_approximate(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-list-bounds")
    _enable_text_import_for_route(monkeypatch)
    cases = [
        ("Bench 3x12/10/8", 8, 12, 90),
        ("Bench 3 x 12/10/8", 8, 12, 90),
        ("12/10/8 reps", 8, 12, 90),
        ("Bench 3x5, 90s rest", 5, 5, 90),
        ("OHP 3x8, 2 min rest", 8, 8, 120),
        ("Bench 1x5, 3x5", 8, 10, 90),
        ("Bench 3x5 RPE 8", 5, 5, 90),
    ]
    rows = []
    for line_number, (line, _low, _high, rest) in enumerate(cases, 1):
        row = _text_reply(line=line_number, exercise="rewritten", original_text=line)["rows"][0]
        row["rest_seconds"] = rest
        rows.append(row)
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Push"}], "weeks": [], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": "\n".join(case[0] for case in cases)})
    assert response.status_code == 200, response.text
    imported = response.json()["rows"]
    assert [(row["reps_min"], row["reps_max"], row["rest_seconds"]) for row in imported] == [
        (low, high, rest) for _, low, high, rest in cases
    ]
    assert "program_import.effort_approximated.v1" not in [warning["code"] for warning in imported[-1]["warnings"]]


def test_text_loads_reuse_load_and_unsupported_markers(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-load-markers")
    _enable_text_import_for_route(monkeypatch)
    lines = ["Dips 3x8 +10kg", "Squat 3x5 at 70% 1RM", "Pull-up 3x8 bodyweight + 10kg"]
    rows = [_text_reply(line=i, exercise="rewritten", original_text=line)["rows"][0]
            for i, line in enumerate(lines, 1)]
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Pull"}], "weeks": [], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": "\n".join(lines)})
    assert response.status_code == 200, response.text
    expected_loads = ["+10kg", "70% 1RM", "bodyweight + 10kg"]
    for row, expected_load in zip(response.json()["rows"], expected_loads, strict=True):
        codes = [warning["code"] for warning in row["warnings"]]
        assert "program_import.unsupported_preserved_as_note.v1" in codes
        assert "program_import.load_preserved_as_note.v1" in codes
        assert row["notes"] == expected_load
        load_marker = next(
            warning for warning in row["warnings"]
            if warning["code"] == "program_import.load_preserved_as_note.v1"
        )
        assert load_marker["message"] == expected_load


def test_text_ai_load_kg_maps_to_load_without_prompt_schema_change(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-kg-load")
    _enable_text_import_for_route(monkeypatch)
    translated = _text_reply(
        line=1,
        exercise="Dips",
        original_text="Dips 3x8 +10kg",
    )["rows"][0]
    translated["load_kg"] = 10
    scripted_chat_model.script({"rows": [translated], "layout": {
        "days": [{"day": 1, "name": "Push"}],
        "weeks": [],
        "confidence": 0.98,
        "confirm_layout": False,
    }})
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        data={"text": "Dips 3x8 +10kg"},
    )
    assert response.status_code == 200, response.text
    imported = response.json()["rows"][0]
    assert imported["load"] == {"kind": "kg", "value": 10}
    assert imported["notes"] is None


def test_text_ai_unsupported_load_is_preserved_as_note(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-unsupported-load")
    _enable_text_import_for_route(monkeypatch)
    translated = _text_reply(
        line=1,
        exercise="Dips",
        original_text="Dips 3x8 BW+10",
    )["rows"][0]
    translated["load_kg"] = "BW+10"
    scripted_chat_model.script({"rows": [translated], "layout": {
        "days": [{"day": 1, "name": "Push"}],
        "weeks": [],
        "confidence": 0.98,
        "confirm_layout": False,
    }})
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        data={"text": "Dips 3x8 BW+10"},
    )

    assert response.status_code == 200, response.text
    imported = response.json()["rows"][0]
    assert imported["valid"] is True
    assert imported["load"] is None
    assert "BW+10" in imported["notes"]
    assert "program_import.invalid_load.v1" not in [
        warning["code"] for warning in imported["warnings"]
    ]
    assert "program_import.load_preserved_as_note.v1" in [
        warning["code"] for warning in imported["warnings"]
    ]


def test_text_hold_requires_hold_evidence_and_numbered_names_are_clean(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-hold-guard")
    _enable_text_import_for_route(monkeypatch)
    lines = ["1) Squat 4x5", "Bike 3x10 min"]
    rows = []
    for index, line in enumerate(lines, 1):
        row = _text_reply(line=index, exercise="rewritten", original_text=line)["rows"][0]
        row.update(reps_min=8, reps_max=12)
        rows.append(row)
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Push"}], "weeks": [], "confidence": 0.98, "confirm_layout": False,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": "\n".join(lines)})
    assert response.status_code == 200, response.text
    assert [row["input_exercise_name"] for row in response.json()["rows"]] == ["Squat", "Bike"]
    assert [(row["reps_min"], row["reps_max"]) for row in response.json()["rows"]] == [(5, 5), (8, 12)]
    assert "program_import.unsupported_preserved_as_note.v1" in [
        warning["code"] for warning in response.json()["rows"][1]["warnings"]
    ]
    assert "program_import.effort_approximated.v1" not in [
        warning["code"] for warning in response.json()["rows"][1]["warnings"]
    ]


def test_text_ambiguous_day_labels_keep_model_layout_hint(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-ambiguous-day-labels")
    _enable_text_import_for_route(monkeypatch)
    text = "Day 1\nSquat 3x8\nDay 2\nBench 3x8"
    rows = [
        _text_reply(line=2, day=1, day_name="Day 1", exercise="Squat", original_text="Squat 3x8")["rows"][0],
        _text_reply(line=4, day=2, day_name="Day 2", exercise="Bench", original_text="Bench 3x8")["rows"][0],
    ]
    scripted_chat_model.script({"rows": rows, "layout": {
        "days": [{"day": 1, "name": "Day 1"}, {"day": 2, "name": "Day 2"}],
        "weeks": [], "confidence": 0.8, "confirm_layout": True,
    }})
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": text})
    assert response.status_code == 200, response.text
    assert response.json()["confirm_layout"] is True


def test_text_unplaced_prescription_line_requires_layout_confirmation(api, monkeypatch, scripted_chat_model):
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-unplaced")
    _enable_text_import_for_route(monkeypatch)
    scripted_chat_model.script(_text_reply(line=2))

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        data={"text": "Day 1 – Push\nSquat 3x8\nBench 4×10"},
    )

    assert response.status_code == 200, response.text
    assert response.json()["confirm_layout"] is True


def test_text_suggestions_are_confirm_only_and_catalog_is_not_sent_to_model(api, monkeypatch, caplog, tmp_path):
    from tests.fakes.chat_model import ScriptedChatModel
    from utils import model_downloader
    from utils.model_metering import MeteringCallback

    caplog.set_level(logging.DEBUG)
    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-suggestions")
    _enable_text_import_for_route(monkeypatch)
    monkeypatch.setenv("AI_LAUNCH_ALLOWANCES_ENABLED", "true")
    monkeypatch.setenv("MODEL_RATE_LIMIT_REQUESTS", "0")
    for capability in ("COACH", "LIFTER"):
        for plan in ("FREE", "PRO"):
            monkeypatch.setenv(f"{capability}_AI_{plan}_DAILY_REQUESTS", "1")
            monkeypatch.setenv(f"{capability}_AI_{plan}_DAILY_RESOURCE_TOKENS", "10000")
    coach_text = "Mystery Press 3x8"
    import_model = ScriptedChatModel(
        turns=[
            _text_reply(line=1, exercise="Mystery Press", search_terms=["Bench Press"]),
            {"suggestions": [{"exercise_name": "Mystery Press", "exercise_ids": ["bp"]}]},
        ],
        model_name="deepseek-ai/DeepSeek-V4-Flash",
        usage_metadata={"input_tokens": 40, "output_tokens": 20, "total_tokens": 60, "input_token_details": {}},
        callbacks=[MeteringCallback("deepseek-ai/DeepSeek-V4-Flash")],
    )
    monkeypatch.setattr(model_downloader, "get_coach_import_llm", lambda: import_model)
    response = client.post(f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
                           data={"text": coach_text})
    assert response.status_code == 200, response.text
    row = response.json()["rows"][0]
    assert row["exercise_id"] is None
    assert row["suggestions"][0]["name"] == "Bench Press"
    assert len(import_model.calls) == 2
    model_content = " ".join(str(message.content) for call in import_model.calls for message in call["messages"])
    assert coach_text in model_content
    assert "Bench Press" in model_content
    assert coach_text not in caplog.text
    account_id = db.get_active_account_by_username("text-suggestions-coach")["account_id"]
    usage_count = db.catalog_conn.execute(
        "SELECT COUNT(*) FROM model_usage WHERE account_id = ? AND purpose = 'coach_program_import'",
        (account_id,),
    ).fetchone()[0]
    # The translation and suggestion calls share one admitted coach turn.
    assert usage_count == 2
    assert all(coach_text.encode() not in path.read_bytes() for path in tmp_path.rglob("*") if path.is_file())


def test_text_gate_requires_its_own_prompt_bound_report(api, monkeypatch, tmp_path):
    from service import program_import_ai

    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-gate")
    monkeypatch.setenv("PROGRAM_IMPORT_TEXT_AI_ENABLED", "true")
    monkeypatch.delenv("PROGRAM_IMPORT_TEXT_AI_EVAL_REPORT", raising=False)
    capability = client.get("/coach/program-import/capabilities", headers=coach_headers)
    assert capability.status_code == 200
    assert capability.json()["text_available"] is False
    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import", headers=coach_headers,
        data={"text": "Squat 3x8"},
    )
    assert response.status_code == 503
    assert response.json()["message_code"] == "program_import.text_disabled.v1"
    assert program_import_ai.text_prompt_version_hash() != program_import_ai.prompt_version_hash()


def test_committed_sheet_report_and_prompt_hash_remain_valid():
    from service import program_import_ai

    report = json.loads((Path(__file__).parents[1] / "reports/program_import_eval.json").read_text(encoding="utf-8"))
    assert program_import_ai.prompt_version_hash() == "a5c879aa7e28699a58ef47234425587fadabbe720b2b47471c1512788d1f3e7c"
    assert program_import_ai.text_prompt_version_hash() == "fdd73d190b388369fabe9f0270d985bc038725e2aacf342209ee2603136a2120"
    assert program_import_ai.validate_report(report) == (True, [])
    text_report = json.loads((Path(__file__).parents[1] / "reports/program_import_text_eval.json").read_text(encoding="utf-8"))
    assert program_import_ai.validate_report(text_report) == (True, [])


def test_text_report_validator_uses_lines_and_skips_tab_selection(monkeypatch, tmp_path):
    from service import program_import_ai
    report = _valid_text_evaluation_report()

    assert program_import_ai.validate_report(report) == (True, [])
    report_path = tmp_path / "text-report.json"
    report_path.write_text(json.dumps(report), encoding="utf-8")
    monkeypatch.setenv("PROGRAM_IMPORT_TEXT_AI_ENABLED", "true")
    monkeypatch.setenv("PROGRAM_IMPORT_TEXT_AI_EVAL_REPORT", str(report_path))
    monkeypatch.setenv("PROGRAM_IMPORT_AI_ENABLED", "true")
    monkeypatch.delenv("PROGRAM_IMPORT_AI_EVAL_REPORT", raising=False)
    assert program_import_ai.resolve_text_enable_gate().enabled is True
    assert program_import_ai.resolve_enable_gate().enabled is False
    monkeypatch.setenv("PROGRAM_IMPORT_TEXT_AI_ENABLED", "false")
    _enable_program_import_ai(monkeypatch, tmp_path)
    assert program_import_ai.resolve_enable_gate().enabled is True
    assert program_import_ai.resolve_text_enable_gate().enabled is False


def test_sheet_report_cannot_enable_text_gate_through_http_route(api, monkeypatch, tmp_path):
    from service import program_import_ai

    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="text-wrong-report")
    _enable_program_import_ai(monkeypatch, tmp_path)
    monkeypatch.setenv("PROGRAM_IMPORT_TEXT_AI_ENABLED", "true")
    monkeypatch.setenv(
        "PROGRAM_IMPORT_TEXT_AI_EVAL_REPORT", __import__("os").environ["PROGRAM_IMPORT_AI_EVAL_REPORT"]
    )

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        data={"text": "Squat 3x8"},
    )

    assert response.status_code == 503
    assert response.json()["message_code"] == "program_import.text_disabled.v1"
    assert program_import_ai.resolve_enable_gate().enabled is True


def test_text_report_cannot_enable_sheet_gate_through_http_route(api, monkeypatch, tmp_path):
    from service import program_import_ai

    client, db = api
    coach_headers, _, assignment_id = _assignment(client, db, prefix="sheet-wrong-report")
    report_path = tmp_path / "text-report.json"
    report_path.write_text(json.dumps(_valid_text_evaluation_report()), encoding="utf-8")
    monkeypatch.setenv("PROGRAM_IMPORT_AI_ENABLED", "true")
    monkeypatch.setenv("PROGRAM_IMPORT_AI_EVAL_REPORT", str(report_path))

    response = client.post(
        f"/coach/assignments/{assignment_id}/program-import",
        headers=coach_headers,
        files={"file": ("program.csv", b"exercise_name,sets,reps\nSquat,3,8\n", "text/csv")},
    )

    assert response.status_code == 503
    assert response.json()["message_code"] == "program_import.freeform_disabled.v1"
    assert program_import_ai.resolve_text_enable_gate().enabled is False


def test_startup_logs_text_and_sheet_gates_separately(monkeypatch, caplog):
    from service import program_import_ai
    from service.evaluation_report_gate import GateStatus

    caplog.set_level(logging.DEBUG)
    monkeypatch.setattr(
        program_import_ai,
        "resolve_enable_gate",
        lambda: GateStatus(True, False, "sheet report refused"),
    )
    monkeypatch.setattr(
        program_import_ai,
        "resolve_text_enable_gate",
        lambda: GateStatus(True, False, "text report refused"),
    )

    program_import_ai.log_enable_gate_at_startup()
    program_import_ai.log_text_enable_gate_at_startup()

    assert "Free-form Program import requested" in caplog.text
    assert "Plain-text Program import requested" in caplog.text
    assert "sheet report refused" in caplog.text
    assert "text report refused" in caplog.text
